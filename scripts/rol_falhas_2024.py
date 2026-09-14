"""
O rol de falhas de 2024, no formato da aba «Falha Equipamentos» da planilha base.

Pedido do gestor (14/09): «aqui é taxa de falha faça até 2024 — não conserte 2025 e
2026, só preciso das quantidades, 2025 e 2026 está correto, só quero 2024».

DE ONDE SAI: `data/analise_ia/leitura_ss_os/shard*.json` — a leitura das SS e OS
que os agentes fizeram, com peça, data, citação do texto e revisão de cada falha.
Ela é a FONTE do rol que o gestor já tem, e isso se prova sozinho:

                       leitura        rol do gestor
    2026 RL                 28                   28   ← bate exato
    2026 RT                  9                    9   ← bate exato
    2025 RL                 36                   35
    2025 RT                 19                   18
    2024 RL                 10            não existe
    2024 RT                  5            não existe

2026 bate no número, 2025 erra por um em cada tipo. É a mesma régua, então 2024 sai
na mesma medida que os anos que ele considera certos — que é exatamente o pedido.

**2024 = 10 religadores + 5 reguladores**, 17 falhas (dois ativos falharam duas
vezes: 7908708116 com tanque e controle, 5820790038 com célula e controle). Todas
com confiança alta, todas revisadas.

A RESSALVA QUE PRECISA IR JUNTO: a leitura cobriu os 129 ativos da carteira, que é a
foto do que está PENDENTE. Quem falhou e foi resolvido saiu da carteira, e quanto
mais para trás no tempo, mais gente saiu. É por isso que 2024 dá 15 e 2025 dá 55 —
não é que 2024 teve menos falha, é que 2024 teve mais tempo para ser resolvido. O
AIC mostra o avesso disso: obra de substituição concluída em 2024 são 53 no
religador contra 27 em 2025, justamente porque as de 2024 fecharam.

Então 10 e 5 são PISO na mesma régua do rol, não censo do parque.

Grava dist/ROL_FALHAS_2024.xlsx, com as linhas prontas para colar na aba
«Falha Equipamentos».

Rodar: python3 scripts/rol_falhas_2024.py
"""

import glob
import json
import os
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LEITURA = os.path.join(RAIZ, "data", "analise_ia", "leitura_ss_os", "shard*.json")
AJUSTES = os.path.join(RAIZ, "data", "raw", "GESTAO_DE_EQUIPAMENTOS.xlsx")
SAIDA = os.path.join(RAIZ, "dist", "ROL_FALHAS_2024.xlsx")

TINTA, PAPEL, SINAL = "FF211D15", "FFF2EFE6", "FFBC4B0E"
ANO = 2024
MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]
# as colunas da aba «Falha Equipamentos», na ordem exata da planilha do gestor
COLUNAS = ["Ano/TIPO", "Equipamento", "Mensal", "Parque", "Ativo", "SS", "Tensão",
           "Data", "Mês", "Ano", "Concat", "Peça (modo)", "Troca feita",
           "Causa raiz", "Citação do texto da SS", "Nota do analista", "Revisão"]
# o que o gestor já tem, para a aba de conferência provar que a régua é a mesma
ROL_GESTOR = {2025: {"RL": 35, "RT": 18}, 2026: {"RL": 28, "RT": 9}}


def ler_leitura():
    det = []
    for caminho in sorted(glob.glob(LEITURA)):
        with open(caminho, encoding="utf-8") as fh:
            det += json.load(fh).get("detalhe", [])
    return det


def tensao_por_ativo():
    """A tensão de cadastro, dos ajustes da proteção — mesma fonte da taxa por peça."""
    wb = load_workbook(AJUSTES, read_only=True, data_only=True)
    ten = {}
    for linha in list(wb["Ajustes RL Poste"].iter_rows(values_only=True))[1:]:
        if linha[0] and linha[12]:
            ten[str(linha[0]).strip()] = str(linha[12]).strip()
    for linha in list(wb["Ajustes Reguladores de Tensão"].iter_rows(values_only=True))[1:]:
        if linha[0] and linha[8]:
            ten[str(linha[0]).strip()] = str(linha[8]).strip()
    wb.close()
    return ten


def montar():
    det = ler_leitura()
    ten = tensao_por_ativo()
    linhas = []
    for x in sorted([d for d in det if d["ano"] == ANO],
                    key=lambda y: (y["familia"], y["data"][3:5], y["data"][:2])):
        tipo = "RL" if x["familia"] == "religador" else "RT"
        data = x.get("data", "")
        mes = MESES[int(data[3:5]) - 1] if len(data) >= 5 and data[3:5].isdigit() else ""
        linhas.append([
            f"{tipo} {ANO}", tipo, "", "", x["ativo"], x.get("ss", ""),
            ten.get(x["ativo"], ""), data, mes, ANO, f"{ANO}{mes}",
            x.get("peca", ""), "sim" if x.get("executada") else "",
            (x.get("revisao_motivo") or "")[:120],
            (x.get("evidencia") or "")[:900],
            f"confiança {x.get('confianca', '')}"
            + ("; revisado" if x.get("revisado") else ""),
            "leitura das SS/OS (mesma fonte do rol de 2025 e 2026)",
        ])
    return det, linhas


# ------------------------------------------------------------------ a planilha
def cabeca(ws, linha, titulo, cols, larguras=None):
    ws.cell(row=linha, column=1, value=titulo).font = Font(bold=True, size=11,
                                                           color=SINAL)
    ws.append(cols)
    for i in range(1, len(cols) + 1):
        c = ws.cell(row=linha + 1, column=i)
        c.font = Font(bold=True, color=PAPEL, size=10)
        c.fill = PatternFill("solid", fgColor=TINTA)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    if larguras:
        for i, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    return linha + 2


def fecha(ws, linha, n):
    for i in range(1, n + 1):
        ws.cell(row=linha, column=i).font = Font(bold=True)
        ws.cell(row=linha, column=i).border = Border(
            top=Side(style="medium", color=TINTA))


def planilha(det, linhas):
    wb = Workbook()

    # 1 — as linhas prontas para colar
    ws = wb.active
    ws.title = "Falha Equipamentos 2024"
    r0 = cabeca(ws, 1,
                f"As {len(linhas)} falhas de 2024 — colunas na ordem da aba "
                f"«Falha Equipamentos»",
                COLUNAS,
                [11, 12, 8, 8, 13, 22, 10, 11, 11, 7, 15, 12, 11, 40, 70, 20, 30])
    for linha in linhas:
        ws.append(linha)
        ws.cell(row=ws.max_row, column=15).alignment = Alignment(wrap_text=True,
                                                                 vertical="top")
        ws.cell(row=ws.max_row, column=14).alignment = Alignment(wrap_text=True,
                                                                 vertical="top")
        ws.row_dimensions[ws.max_row].height = 42
    ws.freeze_panes = f"A{r0}"
    ws.auto_filter.ref = f"A{r0 - 1}:Q{ws.max_row}"

    # 2 — a quantidade, que é o que foi pedido
    ws2 = wb.create_sheet("Quantidade")
    r0 = cabeca(ws2, 1, "Equipamentos que falharam em cada ano — mesma régua",
                ["Ano", "Religador", "Regulador", "Total", "Origem"],
                [10, 12, 12, 10, 46])
    rl24 = len({x["ativo"] for x in det if x["ano"] == 2024
                and x["familia"] == "religador"})
    rt24 = len({x["ativo"] for x in det if x["ano"] == 2024
                and x["familia"] == "regulador"})
    ws2.append([2024, rl24, rt24, rl24 + rt24, "leitura das SS/OS — NOVO"])
    for c in range(1, 6):
        ws2.cell(row=ws2.max_row, column=c).font = Font(bold=True, color=SINAL)
    for ano in (2025, 2026):
        g = ROL_GESTOR[ano]
        ws2.append([ano, g["RL"], g["RT"], g["RL"] + g["RT"],
                    "rol do gestor — não mexido"])
    fecha(ws2, ws2.max_row, 5)

    ws2.append([])
    r1 = cabeca(ws2, ws2.max_row + 1,
                "A prova de que é a mesma régua: a leitura reproduz o rol do gestor",
                ["Ano", "Tipo", "Leitura", "Rol do gestor", "Diferença"],
                None)
    for ano in (2025, 2026):
        for t, fam in (("RL", "religador"), ("RT", "regulador")):
            n = len({x["ativo"] for x in det if x["ano"] == ano
                     and x["familia"] == fam})
            ws2.append([ano, t, n, ROL_GESTOR[ano][t], n - ROL_GESTOR[ano][t]])
    fecha(ws2, ws2.max_row, 5)

    ws2.append([])
    for t in [
        "A ressalva que precisa ir junto do 10 e do 5:",
        "",
        "A leitura cobriu os 129 ativos da CARTEIRA, que é a foto do que está pendente.",
        "Quem falhou e foi resolvido saiu da carteira — e quanto mais para trás no",
        "tempo, mais gente saiu. Por isso 2024 dá 15 e 2025 dá 55: não é que 2024 teve",
        "menos falha, é que 2024 teve mais tempo para ser resolvido.",
        "",
        "O AIC mostra o avesso disso, e confirma:",
        "obra de substituição concluída em 2024 são 53 no religador contra 27 em 2025 —",
        "justamente porque as de 2024 já fecharam e as de 2025 ainda estão abertas.",
        "",
        "Então 10 e 5 são PISO na mesma régua do rol, não censo do parque. Para o censo",
        "de 2024 seria preciso um export da base de SS/OS com a coluna de descrição",
        "cobrindo o ano — aí a leitura roda sobre o parque inteiro, não sobre a carteira.",
    ]:
        ws2.append([t])
        if t.endswith(":"):
            ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True, size=11)

    # 3 — o detalhe por peça
    ws3 = wb.create_sheet("Por peça")
    cabeca(ws3, 1, "As 17 falhas de 2024 por peça",
           ["Peça", "Religador", "Regulador", "Total"], [16, 12, 12, 10])
    d24 = [x for x in det if x["ano"] == 2024]
    pecas = Counter((x["peca"], x["familia"]) for x in d24)
    for peca in sorted({p for p, _ in pecas}):
        rl = pecas.get((peca, "religador"), 0)
        rt = pecas.get((peca, "regulador"), 0)
        ws3.append([peca, rl or "—", rt or "—", rl + rt])
    ws3.append(["Total de falhas", sum(1 for x in d24 if x["familia"] == "religador"),
                sum(1 for x in d24 if x["familia"] == "regulador"), len(d24)])
    fecha(ws3, ws3.max_row, 4)
    ws3.append([])
    ws3.append(["Equipamentos distintos",
                len({x["ativo"] for x in d24 if x["familia"] == "religador"}),
                len({x["ativo"] for x in d24 if x["familia"] == "regulador"}),
                len({x["ativo"] for x in d24})])
    for c in range(1, 5):
        ws3.cell(row=ws3.max_row, column=c).font = Font(bold=True)
    ws3.append([])
    repetidos = [a for a, n in Counter(x["ativo"] for x in d24).items() if n > 1]
    ws3.append([f"Ativos com duas falhas no ano: {', '.join(repetidos)} — "
                f"por peça contam {len(d24)}, por equipamento {len({x['ativo'] for x in d24})}"])

    # 4 — a régua
    ws4 = wb.create_sheet("Como foi feito")
    ws4.column_dimensions["A"].width = 98
    for t in [
        "O rol de falhas de 2024",
        "",
        "A fonte:",
        "data/analise_ia/leitura_ss_os/shard*.json — a leitura que os agentes fizeram",
        "do texto das SS e das OS, com peça, data, citação e revisão de cada falha.",
        "113 falhas confirmadas ao todo, de 127 apontadas: 14 foram derrubadas na",
        "revisão por não se sustentarem no texto.",
        "",
        "Por que essa fonte e não outra:",
        "Ela reproduz o rol que o gestor já tem. Em 2026 bate no número exato (28",
        "religadores e 9 reguladores) e em 2025 erra por um em cada tipo (36 contra 35,",
        "19 contra 18). Mesma régua, mesma leitura — então 2024 sai comparável com os",
        "anos que ele considera certos.",
        "",
        "A régua da peça, que é a do gestor de 21/08:",
        "Conta como falha o que exigiu peça grande — controle, tanque/parte ativa ou",
        "equipamento completo no religador; célula, relé, banco completo ou furto no",
        "regulador. Fica fora trafo auxiliar, chave faca, rádio, antena, bateria,",
        "aterramento, cabo, conector, poste, poda, ajuste de proteção, comissionamento",
        "e obra de equipamento novo.",
        "",
        "O ano é o da ocorrência:",
        "Todas as 17 linhas de 2024 têm data de ocorrência, não de abertura da SS. Uma",
        "delas tem SS de 2026 (7908686014, ETO-COEP 00086/2026) porque a SS foi aberta",
        "depois — a falha é de 15/10/2024 e é aí que ela conta.",
        "",
        "Ativo repetido:",
        "Dois ativos falharam duas vezes em 2024 — 7908708116 (tanque em dezembro e",
        "controle em setembro) e 5820790038 (célula em agosto e controle em março).",
        "Por peça são 17 linhas; por equipamento, 15. A régua do gestor conta",
        "equipamento, então a quantidade do ano é 15: 10 religadores e 5 reguladores.",
        "",
        "O que NÃO foi mexido:",
        "2025 e 2026 ficam exatamente como estão na planilha do gestor. Nenhuma linha",
        "dele foi recalculada, reclassificada ou reordenada.",
    ]:
        ws4.append([t])
        ws4.cell(row=ws4.max_row, column=1).alignment = Alignment(wrap_text=True)
    ws4["A1"].font = Font(bold=True, size=12, color=SINAL)
    for n in range(1, ws4.max_row + 1):
        v = ws4.cell(row=n, column=1).value
        if v and str(v).endswith(":"):
            ws4.cell(row=n, column=1).font = Font(bold=True, size=11)

    for aba in wb.worksheets:
        for linha in aba.iter_rows():
            for c in linha:
                if c.value is not None:
                    c.font = Font(name="Arial", bold=c.font.bold,
                                  size=c.font.size or 10, color=c.font.color)

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    wb.save(SAIDA)
    return SAIDA


def main():
    det, linhas = montar()
    rl = len({x["ativo"] for x in det if x["ano"] == 2024
              and x["familia"] == "religador"})
    rt = len({x["ativo"] for x in det if x["ano"] == 2024
              and x["familia"] == "regulador"})
    print(f"2024: {len(linhas)} falhas em {rl + rt} equipamentos — {rl} RL + {rt} RT")
    for ano in (2025, 2026):
        a = len({x["ativo"] for x in det if x["ano"] == ano
                 and x["familia"] == "religador"})
        b = len({x["ativo"] for x in det if x["ano"] == ano
                 and x["familia"] == "regulador"})
        print(f"  confere {ano}: leitura {a} RL · {b} RT | "
              f"rol do gestor {ROL_GESTOR[ano]['RL']} RL · {ROL_GESTOR[ano]['RT']} RT")
    print("gravado:", planilha(det, linhas))


if __name__ == "__main__":
    main()

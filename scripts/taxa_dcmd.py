"""
A taxa de falha de equipamento especial no DCMD — tudo que passou pelos postos do DCMD.

Pedido do gestor (14/09): «recalcula a taxa de falha de equipamentos especiais do
DCMD, todas as SS que passaram pelo DCMD».

QUEM É O DCMD: o posto com «-RD-» no nome — ETO-RD-AR, ETO-RD-GU, ETO-RD-DP,
ETO-RD-PA, ETO-RD-GR, ETO-RD-PS, ETO-RD-PO, ETO-RD-AG, ENC-RD-PS, DOLP-RD-PA,
DG-RD-PO, ESO-RD-PO, ESO-RD-PA, ESO-RD-PS. É a régua que o `coep_para_dcmd.py` já
usava: «equipe do DCMD é a que tem RD no nome do posto — é a régua do gestor para
separar campo de escritório».

O PROJETO TEM DUAS DEFINIÇÕES, e as duas saem na planilha:
  DCMD campo     — só os postos RD, quem executa
  DCMD + COEP    — os RD mais o ETO-COEP, porque na visão ETO o que está parado no
                   COEP também conta como DCMD (em aquisição ou em logística)

PASSOU PELO DCMD = o ativo teve ao menos uma SS num posto do DCMD naquele ano. É a
mesma régua de «passou pelo posto» que o COEP já usa: esteve lá em algum momento,
não só a que chegou no ano. As duas bases entram juntas — o recorte de SS/OS
(6.362 SS) e a base de repasse (10.386), que é a única que enxerga antes de 2022.

AS DUAS TAXAS, que respondem perguntas diferentes:

  PASSAGEM — ativos que passaram pelo DCMD ÷ parque. Quantos por cento do parque
             encostaram no DCMD no ano, por qualquer motivo.
  FALHA    — ativos com peça grande QUE passaram pelo DCMD ÷ parque. A régua do
             gestor de 21/08 (controle, tanque/parte ativa, completo no religador;
             célula, relé, banco completo, furto no regulador), restrita a quem
             passou pelo DCMD.

O parque de cada ano é o real, da série mensal reconstruída do cadastro — 2024 RL
1.083 · RT 130; 2025 RL 1.234 · RT 162; 2026 RL 1.287 · RT 182. O parque fixo de
1.307 e 207 sai ao lado, para comparar.

A RESSALVA DA FALHA: a coluna de falha depende da leitura das SS, que cobriu os 129
ativos da carteira. Quem falhou e foi resolvido saiu da carteira, então a falha de
2024 é piso. A coluna de PASSAGEM não tem esse problema — ela sai da base inteira.

Grava dist/TAXA_DCMD.xlsx.
Rodar: python3 scripts/taxa_dcmd.py
"""

import glob
import json
import os
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.drawing.colors import ColorChoice
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, "dist", "TAXA_DCMD.xlsx")

TINTA, PAPEL, SINAL = "FF211D15", "FFF2EFE6", "FFBC4B0E"
COR_A, COR_B = "1F7C50", "B8480C"
ANOS = ("2024", "2025", "2026")
FIXO = {"RL": 1307, "RT": 207}


def posto(num):
    return (num or "").strip().split()[0] if num else ""


def tipo(cod):
    if cod.startswith(("79", "78")):
        return "RL"
    return "RT" if cod.startswith("58") else None


def ano_de(bruto):
    if not bruto or len(str(bruto)) < 10:
        return ""
    s = str(bruto)
    return s[:4] if s[4:5] == "-" else s[6:10]


def eventos():
    """(ativo, ano, posto) das duas bases — o recorte de SS/OS e a de repasse."""
    ev = []
    caminho = os.path.join(RAIZ, "data", "missao", "ssos_min.json")
    with open(caminho, encoding="utf-8") as fh:
        for s in json.load(fh):
            a, an = s["NUM_TRAFO"], ano_de(s["DATA_ABERTURA_SS"])
            if tipo(a) and an:
                ev.append((a, an, posto(s["NUMERO_SS"])))
    caminho = os.path.join(RAIZ, "data", "missao", "ss_ocorrencia.json")
    with open(caminho, encoding="utf-8") as fh:
        for o in json.load(fh):
            a = str(o.get("EQUIPAMENTO", "")).strip()
            an = ano_de(o.get("DTA_ABERTURA") or "")
            if tipo(a) and an:
                ev.append((a, an, (o.get("POSTO_SGM") or "").strip()))
    return ev


def falhas_lidas():
    """Ativo com peça grande, por ano e tipo — a mesma leitura que gera o rol."""
    det = []
    padrao = os.path.join(RAIZ, "data", "analise_ia", "leitura_ss_os", "shard*.json")
    for caminho in sorted(glob.glob(padrao)):
        with open(caminho, encoding="utf-8") as fh:
            det += json.load(fh).get("detalhe", [])
    fal = defaultdict(set)
    for x in det:
        t = "RL" if x["familia"] == "religador" else "RT"
        fal[(str(x["ano"]), t)].add(x["ativo"])
    return fal


def parque_real():
    caminho = os.path.join(RAIZ, "data", "missao", "parque_mensal.json")
    with open(caminho, encoding="utf-8") as fh:
        serie = json.load(fh)
    saida = {}
    for ano in ANOS:
        ms = [m for m in serie if m.startswith(ano)]
        for t in ("RL", "RT"):
            saida[(ano, t)] = sum(serie[m]["parque"][t] for m in ms) / len(ms)
    return saida


def montar():
    ev = eventos()
    fal = falhas_lidas()
    parque = parque_real()

    campo, com_coep, todos = defaultdict(set), defaultdict(set), defaultdict(set)
    ss_por_posto, ss_dcmd = Counter(), 0
    for a, an, p in ev:
        t = tipo(a)
        todos[(an, t)].add(a)
        if "-RD-" in p:
            campo[(an, t)].add(a)
            com_coep[(an, t)].add(a)
            ss_por_posto[p] += 1
            ss_dcmd += 1
        elif p == "ETO-COEP":
            com_coep[(an, t)].add(a)

    linhas = []
    for an in ANOS:
        for t in ("RL", "RT"):
            k = (an, t)
            pr, pf = parque[k], FIXO[t]
            f_total = fal.get(k, set())
            f_campo = f_total & campo[k]
            f_coep = f_total & com_coep[k]
            linhas.append({
                "ano": an, "tipo": t,
                "parque_real": round(pr), "parque_fixo": pf,
                "com_ss": len(todos[k]),
                "passou": len(campo[k]), "passou_coep": len(com_coep[k]),
                "passagem_real": len(campo[k]) / pr,
                "passagem_fixo": len(campo[k]) / pf,
                "falhas": len(f_total),
                "falha_dcmd": len(f_campo), "falha_dcmd_coep": len(f_coep),
                "taxa_falha_real": len(f_campo) / pr,
                "taxa_falha_fixo": len(f_campo) / pf,
                "pct_falha_no_dcmd": (len(f_campo) / len(f_total)) if f_total else 0,
            })
    return linhas, ss_por_posto, ss_dcmd, len(ev)


# ------------------------------------------------------------------ a planilha
def cabeca(ws, linha, titulo, cols, larguras=None):
    ws.cell(row=linha, column=1, value=titulo).font = Font(bold=True, size=11,
                                                           color=SINAL)
    ws.append(cols)
    for i in range(1, len(cols) + 1):
        c = ws.cell(row=linha + 1, column=i)
        c.font = Font(bold=True, color=PAPEL, size=10)
        c.fill = PatternFill("solid", fgColor=TINTA)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    if larguras:
        for i, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    return linha + 2


def fecha(ws, linha, n):
    for i in range(1, n + 1):
        ws.cell(row=linha, column=i).font = Font(bold=True)
        ws.cell(row=linha, column=i).border = Border(
            top=Side(style="medium", color=TINTA))


def planilha(linhas, ss_por_posto, ss_dcmd, ss_total):
    wb = Workbook()

    # 1 — as duas taxas
    ws = wb.active
    ws.title = "Taxa no DCMD"
    r0 = cabeca(ws, 1, "Equipamento especial que passou pelo DCMD — passagem e falha",
                ["Ano", "Tipo", "Parque real", "Passou pelo DCMD", "Taxa de passagem",
                 "Falhas no ano", "Falhou E passou", "Taxa de falha no DCMD",
                 "Da falha, quanto passou", "Passou (com COEP)", "Falhou (com COEP)"],
                [8, 7, 12, 15, 15, 12, 13, 17, 16, 15, 15])
    for x in linhas:
        ws.append([x["ano"], x["tipo"], x["parque_real"], x["passou"],
                   x["passagem_real"], x["falhas"], x["falha_dcmd"],
                   x["taxa_falha_real"], x["pct_falha_no_dcmd"],
                   x["passou_coep"], x["falha_dcmd_coep"]])
        for col in (5, 8, 9):
            ws.cell(row=ws.max_row, column=col).number_format = "0.0%"
        ws.cell(row=ws.max_row, column=8).number_format = "0.00%"
    fecha(ws, ws.max_row, 11)
    r1 = ws.max_row

    for n, x in enumerate(linhas):
        ws.cell(row=r0 + n, column=13, value=f"{x['ano']} {x['tipo']}")
    cats = AxDataSource(strRef=StrRef(f=f"'{ws.title}'!$M${r0}:$M${r1}"))

    g = BarChart()
    g.type, g.grouping = "col", "clustered"
    g.title = "Taxa de passagem pelo DCMD, sobre o parque de cada ano"
    g.height, g.width, g.gapWidth = 9, 22, 60
    s = Series(Reference(ws, min_col=5, min_row=r0 - 1, max_row=r1),
               title_from_data=True)
    s.graphicalProperties.solidFill = ColorChoice(srgbClr=COR_B)
    s.graphicalProperties.line.noFill = True
    s.cat = cats
    g.series.append(s)
    g.x_axis.axPos, g.y_axis.axPos = "b", "l"
    g.x_axis.delete = g.y_axis.delete = False
    g.y_axis.numFmt = "0.0%"
    ws.add_chart(g, "A" + str(r1 + 3))
    ws.column_dimensions["M"].hidden = True

    # 2 — quem é o DCMD, posto a posto
    ws2 = wb.create_sheet("Postos do DCMD")
    cabeca(ws2, 1, f"Os postos com «-RD-» no nome — {ss_dcmd} SS de "
                   f"{ss_total} ({ss_dcmd / ss_total * 100:.1f}%)",
           ["Posto", "SS", "% do DCMD"], [16, 10, 12])
    for p, n in ss_por_posto.most_common():
        ws2.append([p, n, n / ss_dcmd])
        ws2.cell(row=ws2.max_row, column=3).number_format = "0.0%"
    ws2.append(["TOTAL", ss_dcmd, 1])
    ws2.cell(row=ws2.max_row, column=3).number_format = "0.0%"
    fecha(ws2, ws2.max_row, 3)

    # 3 — parque fixo contra parque real, na taxa do DCMD
    ws3 = wb.create_sheet("Fixo × real")
    cabeca(ws3, 1, "O que muda usando o parque fixo de 1.307 e 207",
           ["Ano", "Tipo", "Parque fixo", "Parque real", "Passagem (fixo)",
            "Passagem (real)", "Falha DCMD (fixo)", "Falha DCMD (real)"],
           [8, 7, 12, 12, 14, 14, 16, 16])
    for x in linhas:
        ws3.append([x["ano"], x["tipo"], x["parque_fixo"], x["parque_real"],
                    x["passagem_fixo"], x["passagem_real"],
                    x["taxa_falha_fixo"], x["taxa_falha_real"]])
        for col in (5, 6):
            ws3.cell(row=ws3.max_row, column=col).number_format = "0.0%"
        for col in (7, 8):
            ws3.cell(row=ws3.max_row, column=col).number_format = "0.00%"
    fecha(ws3, ws3.max_row, 8)

    # 4 — a régua
    ws4 = wb.create_sheet("Como foi feito")
    ws4.column_dimensions["A"].width = 98
    for t in [
        "A taxa do DCMD — tudo que passou pelos postos do DCMD",
        "",
        "Quem é o DCMD:",
        "O posto com «-RD-» no nome. São catorze: ETO-RD-AR, ETO-RD-GU, ETO-RD-DP,",
        "ETO-RD-PA, ETO-RD-GR, ETO-RD-PS, ETO-RD-PO, ETO-RD-AG, ENC-RD-PS,",
        "DOLP-RD-PA, DG-RD-PO, ESO-RD-PO, ESO-RD-PA e ESO-RD-PS. É a régua que o",
        "coep_para_dcmd.py já usava para separar campo de escritório.",
        "",
        "A planilha traz também a leitura com o COEP junto, porque na visão ETO o que",
        "está parado no COEP também é DCMD — em aquisição ou em logística. Quem decide",
        "qual das duas vale é o gestor.",
        "",
        "Passou pelo DCMD:",
        "O ativo teve ao menos uma SS num posto do DCMD naquele ano. Mesma régua de",
        "«passou pelo posto» que o COEP usa: esteve lá em algum momento. As duas bases",
        "entram juntas — o recorte de SS/OS e a base de repasse, que é a única que",
        "enxerga antes de 2022.",
        "",
        "As duas taxas, que respondem coisas diferentes:",
        "",
        "PASSAGEM — ativos que passaram pelo DCMD ÷ parque. Quantos por cento do parque",
        "encostaram no DCMD no ano, por qualquer motivo. Sai da base inteira, é sólida.",
        "",
        "FALHA — ativos com peça grande QUE passaram pelo DCMD ÷ parque. A régua do",
        "gestor de 21/08, restrita a quem passou pelo DCMD.",
        "",
        "A ressalva da coluna de falha:",
        "Ela depende da leitura das SS, que cobriu os 129 ativos da carteira — a foto do",
        "que está pendente. Quem falhou e foi resolvido saiu da carteira, e quanto mais",
        "para trás no tempo, mais gente saiu. Por isso a falha de 2024 é piso.",
        "A coluna de PASSAGEM não tem esse problema.",
        "",
        "O parque:",
        "O real de cada ano, da série mensal reconstruída do cadastro (parque_mensal_2024).",
        "2024 — RL 1.083 · RT 130 | 2025 — RL 1.234 · RT 162 | 2026 — RL 1.287 · RT 182.",
        "A aba «Fixo × real» mostra o que muda usando os 1.307 e 207 da régua antiga.",
    ]:
        ws4.append([t])
        ws4.cell(row=ws4.max_row, column=1).alignment = Alignment(wrap_text=True)
    ws4["A1"].font = Font(bold=True, size=12, color=SINAL)
    for n in range(1, ws4.max_row + 1):
        v = ws4.cell(row=n, column=1).value
        if v and str(v).endswith(":"):
            ws4.cell(row=n, column=1).font = Font(bold=True, size=11)

    for aba in wb.worksheets:
        for linha in aba.iter_rows():
            for c in linha:
                if c.value is not None:
                    c.font = Font(name="Arial", bold=c.font.bold,
                                  size=c.font.size or 10, color=c.font.color)

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    wb.save(SAIDA)
    return SAIDA


def main():
    linhas, ss_por_posto, ss_dcmd, ss_total = montar()
    print(f"SS em posto do DCMD: {ss_dcmd} de {ss_total} "
          f"({ss_dcmd / ss_total * 100:.1f}%)")
    print(f"\n{'ano':5s} {'tp':3s} {'parque':>7s} {'passou':>7s} {'passagem':>9s} "
          f"{'falhou':>7s} {'no DCMD':>8s} {'taxa falha':>11s}")
    for x in linhas:
        print(f"{x['ano']:5s} {x['tipo']:3s} {x['parque_real']:7d} {x['passou']:7d} "
              f"{x['passagem_real'] * 100:8.1f}% {x['falhas']:7d} "
              f"{x['falha_dcmd']:8d} {x['taxa_falha_real'] * 100:10.2f}%")
    print("\ngravado:", planilha(linhas, ss_por_posto, ss_dcmd, ss_total))


if __name__ == "__main__":
    main()

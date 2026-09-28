"""
Conferência da COEP 5 — o que mudou desde a COEP 4 entregue em 25/09, o que está errado e o
que a base de SS de 23/09 pede para atualizar.

Pedido do gestor (28/09): «verifica as informações da planilha de equipamentos especiais, só
verifica possíveis atualizações». **Só conferir**: a COEP 5 não é alterada; sai uma lista com aba,
célula, o que está e o que deveria estar, para ele acertar no Excel.

Fontes:
- `data/raw/GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_5.xlsx` — a planilha dele (28/09);
- `dist/GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4.xlsx` — a que entreguei em 25/09 (refeita por
  `compra_0309_por_ativo.py` se faltar);
- `data/raw/GESTAO_DE_EQUIPAMENTOS.xlsx` — cadastro de ajustes (marca, tensão, potência, alimentador);
- `data/raw/RELIGA_REGULA_23092026.xlsx` — base de SS. O «Religa_regula.xlsx» que ele mandou junto
  em 28/09 é **o mesmo arquivo** (MD5 dd6bf179…): nada depois de 23/09.

Três achados que explicam quase tudo:
1. **A tabela foi reordenada e as colunas de orçamento e PMA ficaram na ordem antiga** em quatro
   linhas: o RT 5800961074 e o RT 5853360007 ficaram com preço e PMA de religador completo, e os RL
   7900535058 e 7920024127 com o de outra linha. O total da coluna não muda, porque é troca.
2. **Cinco linhas novas (55 a 59) entraram pela metade**: ativo como texto em três, sem tipo, marca,
   tensão, alimentador ou orçamento, e o PMA do 5823916001 em #REF!.
3. **Fórmulas apontando para a pasta Downloads dele** (a carteira «Relação dos Equipamentos
   Indisponíveis - ETO (32)», a «COEP_4 (3)» e o «v6 (1)» do SLA): noutro computador viram #REF!.

Rodar: python3 scripts/confere_coep5.py
"""

import os
import re
import subprocess
import sys
import zipfile
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from compra_0309_por_ativo import BOM, PECA  # noqa: E402
from compra_2907_por_ativo import JA_TEM, MA_ALTA  # noqa: E402
from marca_tensao_53 import _varre, faixa  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(RAIZ, "data", "raw")
COEP5 = os.path.join(RAW, "GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_5.xlsx")
COEP4 = os.path.join(RAIZ, "dist", "GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4.xlsx")
AJUSTES = os.path.join(RAW, "GESTAO_DE_EQUIPAMENTOS.xlsx")
BASE_SS = os.path.join(RAW, "RELIGA_REGULA_23092026.xlsx")
SAIDA = os.path.join(RAIZ, "dist", "CONFERENCIA_COEP_5.xlsx")
ENTREGA = os.path.join(RAIZ, "dist", "GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_5.xlsx")   # entrega_coep5.py

SLA = {"Muito Alta": 12, "Alta": 26, "Média": 86, "Baixa": 134}
COLS_ORC = ("Orçamento MO", "Orçamento MAT", "Orçamento Total", "PMA")

# O que a base de SS (até 23/09) conta e a Gestão ainda não mostra. A SS e a situação são
# conferidas na base ao rodar (assert); o parecer vai na planilha puxado da própria base.
FATOS = [
    ("7926089013", "ETO-RD-PS 00409/2026", "SS PENDENTE",
     "Tanque sai para obra de construção: passa a precisar de RL Completo 34,5. A Observação já diz "
     "isso, mas Defeito, MAT (38.094,72) e Total seguem de «Controle 34,5»."),
    ("7944559149", "ETO-RD-PS 00410/2026", "SS PENDENTE",
     "Mesmo caso do 7926089013, aberto no mesmo dia: o tanque vai para obra de construção. Se "
     "sair, vira RL Completo 34,5 — hoje está como «Controle 34,5»."),
    ("5800961074", "ETO-COEP 00221/2025", "SS PENDENTE",
     "Células doadas para obra de construção (21/09); o parecer de 24/08 já pedia as 3 células. "
     "A Gestão tem 1 célula («Célula 200 34,5»)."),
    ("7931219078", "ETO-COEP 00201/2026", "SS PENDENTE",
     "A SS da Gestão (ETO-RD-PO 00183/2026) foi CANCELADA em 11/09; a SS viva é esta. Em 17/09 a "
     "DMSL diz que o trafo auxiliar ainda estava com problema. A Observação diz que já foi "
     "substituído: a base, até 23/09, não mostra a troca."),
    ("5856070091", "ETO-PROT 00193/2026", "SS PENDENTE",
     "Equipamento instalado; a PROT recebeu pedido de ajuste para comissionar (21/09). Na esteira: "
     "Concluído COCM e Estudo Proteção."),
    ("5858783119", "ETO-RD-AG 00928/2026", "SS ATENDIDA",
     "A linha viva retirou as 3 células com defeito em 09/09."),
    ("7903569004", "ETO-RD-AR 01419/2026", "SS ATENDIDA",
     "Controle retirado em 17/09 para instalar no 7923673004. Se foi instalado, o 7923673004 "
     "precisa só do tanque 13,8 e o controle 13,8 (38296) reservado a ele sobra."),
    ("7900535058", "ETO-COEP 00173/2025", "SS PENDENTE",
     "Em 16/09 o COEP pôs «em logística — substituir controle, alinhado com o Leandro». A Gestão "
     "segue «Gerado PMA». É COOPER: a DMSL pede o conjunto inteiro por um NOJA RC10."),
    ("7955946007", "ETO-TELE 01279/2026", "SS PENDENTE",
     "A Gestão diz Realizado (troca da placa de alimentação CA). A SS da TELE segue pendente, e o "
     "último parecer do COEP (04/09) pede para conferir se a placa foi mesmo trocada."),
    ("5803327001", "DOLP-RD-PA 00971/2026", "SS ATENDIDA",
     "Controle trocado em 15/09 (EMD 520501): confirma o Realizado. Segue na PROT (ETO-PROT "
     "00189/2026) para ajuste."),
    ("7937102148", "ETO-TELE 01368/2026", "SS PENDENTE",
     "COI confirmou em operação (18/09): confirma o Realizado. A SS segue na TELE."),
    ("7955986084", "ETO-COEP 00133/2026", "SS CANCELADA",
     "SS do COEP cancelada em 04/09 (placa de alimentação CA): confirma o Realizado."),
    ("5800440256", "ETO-TELE 01375/2026", "SS ATENDIDA",
     "Comissionamento atendido em 19/09, depois da troca da célula (09/09): confirma o Realizado."),
    ("7908705049", "ETO-COEP 00197/2026", "SS PENDENTE",
     "A SS nova que você pôs na Gestão confere com a base (aberta em 13/09)."),
    ("5823916001", "ETO-COEP 00196/2026", "SS PENDENTE",
     "É furto (parecer DMSL de 06/08): célula da fase C + controle completo. O Defeito diz só "
     "«Células» e o orçamento está vazio."),
    ("7949808058", "ETO-COEP 00195/2026", "SS PENDENTE",
     "Tanque danificado (DMSL 31/08); COEP prevê troca até 13/11. O religador é 34,5 kV (cadastro e "
     "os 14 RL do alimentador LD02010153), mas o Defeito diz «Tanque 13,8»."),
    ("5844630060", "ETO-RD-AR 01440/2026", "SS PENDENTE",
     "Célula boa; falta só a linha viva reconectar o cabo na chave faca da fase C. Não usa material "
     "do DCMD e, pela régua, não é falha."),
    ("7905320122", "ETO-RD-PA 00583/2026", "SS PENDENTE",
     "Chave de entrada da fase B danificada. Chave faca: pela régua, não é falha."),
]

TINTA, PAPEL, SINAL = "FF211D15", "FFF2EFE6", "FFBC4B0E"
PAPEL2, FILETE = "FFE9E5D8", "FFC8C2AF"


def txt(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def reais(v):
    if not isinstance(v, (int, float)):
        return txt(v) or "(vazio)"
    s = f"{v:,.2f}"
    return "R$ " + s.replace(",", "X").replace(".", ",").replace("X", ".")


# ------------------------------------------------------------------ leitura
def gestao(caminho):
    """{ativo: linha} só dentro da Table1 — embaixo dela há rascunho do gestor."""
    wf = openpyxl.load_workbook(caminho)
    fim = int(re.sub(r"\D", "", wf["Gestão"].tables["Table1"].ref.split(":")[1]))
    formulas = wf["Gestão"]
    ws = openpyxl.load_workbook(caminho, data_only=True)["Gestão"]
    cab = [c.value for c in ws[1]]
    col = {n: get_column_letter(i + 1) for i, n in enumerate(cab) if n}
    linhas = {}
    for r in range(2, fim + 1):
        d = {n: ws.cell(r, i + 1).value for i, n in enumerate(cab) if n}
        d["_r"], d["_f"] = r, {n: formulas.cell(r, i + 1).value for i, n in enumerate(cab) if n}
        linhas[txt(d["Ativo"])] = d
    return linhas, col


def tabela7(caminho):
    """{ativo: Counter(PMA)} das unidades da compra (Estoque, linhas 36-100)."""
    ws = openpyxl.load_workbook(caminho, data_only=True)["Estoque"]
    aloc = defaultdict(Counter)
    for r in range(36, 101):
        if ws.cell(r, 1).value is not None:
            aloc[txt(ws.cell(r, 11).value)][txt(ws.cell(r, 1).value)] += 1
    return aloc


def cadastro(alvo):
    wb = openpyxl.load_workbook(AJUSTES, read_only=True, data_only=True)
    rl = _varre(wb["Ajustes RL Poste"], alvo, "EQUIPAMENTO",
                {"marca": ("RELE",), "tensao": ("TENSÃO",), "alim": ("ALIMENTADOR",)})
    rt = _varre(wb["Ajustes Reguladores de Tensão"], alvo, "CÓDIGO",
                {"marca": ("PARTE ATIVA",), "tensao": ("TENSÃO PRIMÁRIA [Kv]",),
                 "kva": ("POTÊNCIA [Kvar]",), "alim": ("ALIMENTADOR",)})
    wb.close()
    return {**rl, **rt}


def base_ss():
    """{SS: registro} — a planilha grava a dimensão como A1:A1; read_only=True leria uma célula."""
    ws = openpyxl.load_workbook(BASE_SS, data_only=True)["Exportar Planilha"]
    linhas = list(ws.iter_rows(values_only=True))
    cab = list(linhas[0])
    por = {}
    for r in linhas[1:]:
        if r[1]:
            por.setdefault(txt(r[1]), dict(zip(cab, r)))
    return por


def parecer(reg, n=260):
    return txt(reg.get("DESCRIÇÃO"))[:n]


# ------------------------------------------------------------------ conferências
def o_que_mudou(g4, g5):
    """Linhas novas, saídas e campos alterados, casando pelo ativo (a ordem mudou)."""
    out = []
    for a in g5:
        if a not in g4:
            d = g5[a]
            out.append([d["_r"], a, "linha nova", "",
                        f"{d['Criticidade']} · {d['Status']} · {d['Defeito']} · {txt(d['SS SGM'])}"])
    for a in g4:
        if a not in g5:
            out.append(["", a, "saiu da Gestão", "", ""])
    for a, d4 in g4.items():
        if a not in g5:
            continue
        d5 = g5[a]
        for k, v4 in d4.items():
            if k.startswith("_") or k in ("Dias Pendente", "Status Atendimento", "Status Prazo"):
                continue
            v5 = d5.get(k)
            igual = (num(v4) == num(v5)) if num(v4) is not None and num(v5) is not None \
                else txt(v4) == txt(v5)
            if not igual:
                fmt = reais if k.startswith("Orçamento") else txt
                out.append([d5["_r"], a, k, fmt(v4) if v4 not in (None, "") else "(vazio)",
                            fmt(v5) if v5 not in (None, "") else "(vazio)"])
    for a, d5 in g5.items():
        if d5.get("Observação"):
            out.append([d5["_r"], a, "Observação (coluna nova)", "", txt(d5["Observação"])])
    return sorted(out, key=lambda x: (x[0] or 999, x[2]))


def veredito(mudou, visto, g5):
    """Confere? — errado (está na aba Corrigir), confirmado pela base de SS ou pelo preço."""
    leitura = {a: l for a, _, _, l in FATOS}
    par = defaultdict(Counter)
    for x in g5.values():
        par[x["Defeito"]][(num(x["Orçamento MO"]), num(x["Orçamento MAT"]))] += 1
    for linha in mudou:
        a, campo = linha[1], linha[2]
        l = leitura.get(a, "")
        d = g5.get(a, {})
        if (a, campo) in visto:
            v = "não — ver aba Corrigir"
        elif campo in ("Status", "SS SGM", "Observação (coluna nova)") and (
                "confirma o" in l or "confere" in l or "Observação já diz" in l):
            v = "sim — a base de SS confirma" + (" (falta mudar o Defeito)"
                                                 if (a, "Defeito") in visto else "")
        elif campo in ("Status", "Observação (coluna nova)") and ("conferir" in l or "não mostra" in l):
            v = "a base de SS (até 23/09) ainda não confirma"
        elif campo in ("Defeito", "Orçamento MAT", "Orçamento Total") and d and \
                par[d["Defeito"]].most_common(1)[0] == ((num(d["Orçamento MO"]), num(d["Orçamento MAT"])),
                                                        par[d["Defeito"]].most_common(1)[0][1]) and \
                par[d["Defeito"]].most_common(1)[0][1] > 1:
            v = "sim — preço confere com o Defeito"
        elif campo == "linha nova":
            v = "incompleta — ver aba Corrigir" if any(x == a for x, _ in visto) else "sim"
        else:
            v = ""
        linha.append(v)
    return mudou


def corrigir(g4, g5, col, aloc, cad):
    """Erros de conteúdo da Gestão: [o quê, aba, célula, ativo, está, deveria, por quê]."""
    out = []
    visto = set()                                   # (ativo, coluna) já apontado

    def add(grupo, d, campos, esta, deveria, porque):
        campos = [campos] if isinstance(campos, str) else campos
        cel = ", ".join(f"{col[c]}{d['_r']}" for c in campos)
        out.append([grupo, "Gestão", cel, txt(d["Ativo"]), esta, deveria, porque])
        visto.update((txt(d["Ativo"]), c) for c in campos)

    # 1) orçamento e PMA que ficaram na posição antiga depois de reordenar a tabela
    g1 = "1 · Orçamento e PMA de outra linha"
    pos4 = {d["_r"]: a for a, d in g4.items()}
    for a, d in sorted(g5.items(), key=lambda x: x[1]["_r"]):
        if a not in g4 or d["_r"] == g4[a]["_r"]:
            continue
        v4 = [g4[a][k] for k in COLS_ORC]
        v5 = [d[k] for k in COLS_ORC]
        vizinho = pos4.get(d["_r"])
        if [txt(x) for x in v4] != [txt(x) for x in v5] and vizinho and \
                [txt(g4[vizinho][k]) for k in COLS_ORC[:3]] == [txt(x) for x in v5[:3]]:
            out.append([g1, "Gestão", f"{col['Orçamento MO']}{d['_r']}:{col['PMA']}{d['_r']}", a,
                        " · ".join(reais(x) for x in v5[:3]) + f" · PMA {txt(v5[3])}",
                        " · ".join(reais(x) for x in v4[:3]) + f" · PMA {txt(v4[3])}",
                        f"{d['Defeito']}: ficou com o orçamento do {vizinho}, que estava na linha "
                        f"{d['_r']} antes de reordenar a tabela"])
            visto.update((a, k) for k in COLS_ORC)

    # 2) defeito ou preço que não bate com o próprio texto ou com as outras linhas
    g2 = "2 · Defeito ou preço que não bate"
    d = g5.get("7955946007")
    if d and d["Defeito"] == "Controle 34,5" and num(d["Orçamento MAT"]) == 76246.2:
        add(g2, d, ["Orçamento MAT", "Orçamento Total"], f"MAT {reais(76246.2)} · Total {reais(89455.54)}",
            f"MAT {reais(38094.72)} · Total {reais(51304.06)}",
            "Defeito «Controle 34,5» com MAT de RL Completo: soma R$ 38.151,48 a mais no total")
    d = g5.get("7926089013")
    if d and d["Defeito"] == "Controle 34,5" and "Completo" in txt(d.get("Observação")):
        add(g2, d, ["Defeito", "Orçamento MAT", "Orçamento Total"],
            f"Controle 34,5 · MAT {reais(38094.72)} · Total {reais(51304.06)}",
            f"RL Completo 34,5 · MAT {reais(76246.2)} · Total {reais(89455.54)}",
            "a Observação diz Completo; a SS ETO-RD-PS 00409/2026 (17/09) leva o tanque para obra")
    d = g5.get("7949808058")
    if d and d["Defeito"] == "Tanque 13,8" and faixa(cad.get("7949808058", {}).get("tensao")) == "34,5 kV":
        add(g2, d, "Defeito", "Tanque 13,8",
            f"Tanque 34,5 · MO {reais(13209.34)} · MAT {reais(38151.48)} · Total {reais(51360.82)}",
            "o religador é 34,5 kV: cadastro de ajustes e os 14 RL do alimentador LD02010153")
    par = defaultdict(Counter)
    for x in g5.values():
        par[x["Defeito"]][(num(x["Orçamento MO"]), num(x["Orçamento MAT"]))] += 1
    for a, d in g5.items():
        (mo, mat), n = par[d["Defeito"]].most_common(1)[0]
        if num(d["Orçamento MO"]) is None or num(d["Orçamento MAT"]) is None:
            continue                                # orçamento vazio: vai no grupo 3
        if n > 1 and (num(d["Orçamento MO"]), num(d["Orçamento MAT"])) != (mo, mat) \
                and (a, "Orçamento MAT") not in visto:
            add(g2, d, "Orçamento MAT", f"MO {reais(d['Orçamento MO'])} · MAT {reais(d['Orçamento MAT'])}",
                f"MO {reais(mo)} · MAT {reais(mat)}", f"é o preço das outras {n} linhas de «{d['Defeito']}»")
    for a, d in g5.items():
        mo, mat, tot = (num(d[k]) for k in COLS_ORC[:3])
        if None not in (mo, mat, tot) and abs(mo + mat - tot) > 0.01:
            add(g2, d, "Orçamento Total", reais(tot), reais(mo + mat), "Total = MO + MAT")

    # 3) linha nova entrou pela metade: uma linha por ativo, com o que falta e de onde tirar
    g3 = "3 · Linha nova incompleta"
    for a, d in sorted(g5.items(), key=lambda x: x[1]["_r"]):
        if a in g4:
            continue
        c = cad.get(a, {})
        falta, valor = [], []
        if txt(d.get("Tipo")) != a[:2]:
            falta.append("Tipo"); valor.append(f"Tipo {a[:2]}")
        if not d.get("Marca") and c.get("marca"):
            falta.append("Marca"); valor.append(c["marca"])
        if not d.get("Faixa de Tensão") and c.get("tensao"):
            falta.append("Faixa de Tensão"); valor.append(faixa(c["tensao"]))
        if a[:2] == "58" and not d.get("Potencia") and c.get("kva"):
            falta.append("Potencia"); valor.append(f"{c['kva']} kVA")
        if None in (num(d[k]) for k in COLS_ORC[:3]):
            falta.append("Orçamento MO"); valor.append("orçamento pelo Defeito")
        if not txt(d.get("PMA")):
            falta.append("PMA"); valor.append("Sem PMA")
        if d["Criticidade"] in SLA and d.get("SLA_Total") != SLA[d["Criticidade"]]:
            falta.append("SLA_Total")
            valor.append(f"SLA {SLA[d['Criticidade']]} ({d['Criticidade']}), Dias Pendente e Status Prazo")
        if d.get("Índice") in (None, ""):
            falta.append("Índice"); valor.append("Índice")
        if not d.get("Alimentador") and c.get("alim"):
            falta.append("Alimentador"); valor.append(c["alim"])
        if not d.get("Município"):
            falta.append("Município"); valor.append("município, polo e regional")
        if falta:
            add(g3, d, falta, "vazio: " + ", ".join(f.replace("Orçamento MO", "orçamento")
                                                    .replace("Faixa de Tensão", "tensão")
                                                    .replace("SLA_Total", "SLA") for f in falta),
                " · ".join(valor), f"{d['Criticidade']} · {d['Status']} · {d['Defeito']} — marca, "
                "tensão, potência e alimentador saem do cadastro de ajustes")

    # 4) ativo gravado como texto
    txts = [d for d in g5.values() if isinstance(d["Ativo"], str)]
    if txts:
        out.append(["4 · Ativo gravado como texto", "Gestão", ", ".join(f"{col['Ativo']}{d['_r']}" for d in txts),
                    ", ".join(txt(d["Ativo"]) for d in txts), "texto", "número",
                    "os outros são número; procura por ativo em outra aba não acha o texto"])

    # 5) fórmula com erro ou presa num arquivo da pasta Downloads
    g5_ = "5 · Fórmula com erro ou ligada à pasta Downloads"
    for a, d in g5.items():
        if txt(d.get("PMA")).startswith("#"):
            add(g5_, d, "PMA", txt(d["PMA"]), "Sem PMA",
                "a fórmula procura na «GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4 (3).xlsx» da pasta Downloads")
        f = d["_f"].get("SS SGM")
        f = f.text if hasattr(f, "text") else f
        if isinstance(f, str) and "[1]" in f:
            add(g5_, d, "SS SGM", "fórmula na carteira da pasta Downloads", txt(d["SS SGM"]),
                "colar como valor: noutro computador vira #REF!")
    na = [d for d in g5.values() if isinstance(d.get("Descrição"), str) and d["Descrição"].startswith("#")]
    if na:
        out.append([g5_, "Gestão", ", ".join(f"{col['Descrição']}{d['_r']}" for d in na),
                    ", ".join(txt(d["Ativo"]) for d in na), "#N/A", "parecer da SS, ou vazio",
                    "a SS não está na aba BASE SS_OS, que só guarda SS pendente (os 5 Realizado e "
                    "o 7931219078); pôr SEERRO(…;\"\") na fórmula"])

    # 6) dado que estava preenchido na COEP 4 e sumiu
    sumiu = [d for a, d in g5.items() if a in g4 and not d.get("Alimentador") and g4[a].get("Alimentador")]
    if sumiu:
        out.append(["6 · Dado apagado", "Gestão", ", ".join(f"{col['Alimentador']}{d['_r']}" for d in sumiu),
                    ", ".join(txt(d["Ativo"]) for d in sumiu), "(vazio)",
                    " · ".join(txt(g4[txt(d['Ativo'])]["Alimentador"]) for d in sumiu),
                    "o alimentador estava preenchido na COEP 4 (é o mesmo nos três: Rio Sono e "
                    "Pedro Afonso)"])
        visto.update((txt(d["Ativo"]), "Alimentador") for d in sumiu)
    return out, visto


def outras_abas(caminho=COEP5):
    """Dinâmicas, vínculos externos, quadros e fórmulas das outras abas."""
    out = []
    z = zipfile.ZipFile(caminho)
    rels = z.read("xl/pivotCache/_rels/pivotCacheDefinition4.xml.rels").decode()
    if "Downloads" in rels:
        out.append(["7 · Outras abas", "SLA de Manutenção", "A3:K17", "", "tabela dinâmica lê «v6 (1).xlsx» da pasta Downloads",
                    "ler da Tabela1 da aba Base SLA de Manutenção (mesmas 48 linhas)",
                    "noutro computador não atualiza; Analisar › Alterar Fonte de Dados"])
    for n in z.namelist():
        if re.match(r"xl/pivotCache/pivotCacheDefinition\d+\.xml$", n):
            x = z.read(n).decode()
            if 'worksheetSource name="Table1"' in x and 'refreshOnLoad="1"' not in x:
                qtd = int(re.search(r'recordCount="(\d+)"', x).group(1))
                out.append(["7 · Outras abas", "Orçamento", "A3:D14", "",
                            f"tabela dinâmica por status de 21/09: {qtd} ativos, 1 Realizado, "
                            "R$ 4.969.638,46",
                            "58 ativos, 5 Realizado, R$ 5.154.017,55 na coluna Total",
                            "Dados › Atualizar Tudo"])
    ws = openpyxl.load_workbook(caminho)["SLA de Manutenção"]
    wv = openpyxl.load_workbook(caminho, data_only=True)["SLA de Manutenção"]
    if txt(ws["I44"].value) == "es":
        out.append(["7 · Outras abas", "SLA de Manutenção", "I44", "", "«es» (texto)", '=SEERRO(I42/(I42+I43);"")',
                    "% de ETO-RD-PA em julho; a fórmula foi apagada"])
    if txt(ws["A107"].value) == "ETO-RD-GU":
        out.append(["7 · Outras abas", "SLA de Manutenção", "A107", "", "ETO-RD-GU", "Total",
                    "a linha soma todas as equipes (C107 = C39+C42+…+C66)"])
    cidade = {"P92": ("K133", "Araguaína"), "P95": ("K136", "Guaraí"),
              "P98": ("K139", "Dianópolis"), "P101": ("K142", "Gurupi")}
    for c, (fim, nome) in cidade.items():
        f = txt(ws[c].value)
        if f.startswith("=SUM("):
            out.append(["7 · Outras abas", "SLA de Manutenção", c, "", f"{f} = {wv[c].value}",
                        f"={fim} = {wv[fim].value}",
                        f"{nome}: soma o acumulado de todos os meses; o total é o último mês"])
    e = openpyxl.load_workbook(caminho)["Estoque"]
    if e["B1"].value is None and e["F7"].value is None:
        out.append(["6 · Dado apagado", "Estoque", "A6:H16 e B1", "", "quadro «Status PMA» e saldo apagados",
                    "as 10 linhas dos PMA e o saldo (9 peças)",
                    "a lista peça a peça (linhas 35 a 100) está igual à de 25/09; se apagou de "
                    "propósito, ignore"])
    o = openpyxl.load_workbook(caminho)["Orçamento"]
    if o["C19"].value is None:
        out.append(["6 · Dado apagado", "Orçamento", "B16:E30", "", "bloco apagado",
                    "8481/8495, orçado 6,1 mi, realizado, execução, logística e BC (R$ 315.372,40)",
                    "nenhum gráfico usa essas células; se apagou de propósito, ignore"])
    return out


def base(g5, por_ss):
    out = []
    for a, ss, situacao, leitura in FATOS:
        reg = por_ss.get(ss)
        assert reg, f"{ss} não está na base"
        assert txt(reg["EQUIPAMENTO"]) == a, (ss, reg["EQUIPAMENTO"], a)
        assert reg["STATUS"] == situacao, (ss, reg["STATUS"], situacao)
        d = g5.get(a, {})
        out.append([d.get("_r", ""), a, txt(d.get("Status")), ss, reg["POSTO_SGM"], situacao,
                    reg["DTA_ABERTURA"].strftime("%d/%m/%Y"), leitura, parecer(reg)])
    return out


def na_entrega(erros, g4, cad):
    """Cada item da lista, conferido de novo na COEP 5 entregue: corrigido, em parte ou com ele."""
    g, col = gestao(ENTREGA)
    resto, _ = corrigir(g4, g, col, tabela7(ENTREGA), cad)
    resto += outras_abas(ENTREGA)
    for e in erros:
        grupo, aba, cel, ativo = e[:4]
        mesmo = [r for r in resto if r[0] == grupo and r[1] == aba and (r[3] == ativo if ativo else r[2] == cel)]
        if not mesmo:
            e.append("orçamento: atualiza sozinha ao abrir o arquivo" if (aba, cel) == ("Orçamento", "A3:D14")
                     else "corrigido na entrega")
        elif mesmo[0][2] == cel and mesmo[0][4] == e[4]:
            e.append("fica com você")
        else:
            e.append("em parte — ainda " + mesmo[0][4])
    return erros


def mudou_na_entrega():
    """Célula a célula, da COEP 5 que ele mandou para a entregue — o que a entrega mexeu."""
    from compra_2907_por_ativo import todas_as_celulas
    antes, depois = todas_as_celulas(COEP5), todas_as_celulas(ENTREGA)
    va = openpyxl.load_workbook(COEP5, data_only=True)
    vd = openpyxl.load_workbook(ENTREGA, data_only=True)
    out, desc, ssos = [], 0, 0
    from openpyxl.utils import column_index_from_string as ci
    for k in sorted(set(antes) | set(depois), key=lambda k: (k[0], int(re.sub(r"\D", "", k[1])),
                                                            ci(re.sub(r"\d", "", k[1])))):
        if antes.get(k) == depois.get(k) or k[0].startswith("Ajustes"):
            continue
        aba, ref = k
        if aba == "BASE SS_OS":
            ssos += 1
            continue
        if aba == "Gestão" and ref.startswith("G"):
            desc += 1
            continue
        a = txt(vd[aba][f"A{re.sub(r'[A-Z]', '', ref)}"].value) if aba == "Gestão" else ""
        v0, v1 = va[aba][ref].value, vd[aba][ref].value
        dinheiro = aba == "Gestão" and re.sub(r"\d", "", ref) in ("L", "M", "N")
        fmt = (lambda v: reais(v) if dinheiro and isinstance(v, (int, float)) else txt(v))
        if txt(v0) == txt(v1) and type(v0) is not type(v1):         # ativo: texto → número
            out.append([aba, ref, a, f"{txt(v0)} (texto)", f"{txt(v1)} (número)"])
            continue
        out.append([aba, ref, a, fmt(v0) or "(vazio)", fmt(v1) or "(vazio)"])
    out.append(["Gestão", "coluna Descrição", "", "PROCX; #N/A quando a SS não está na BASE SS_OS",
                f"SEERRO(PROCX(…);\"\") nas {desc} fórmulas"])
    out.append(["BASE SS_OS", "BN:BP", "", "(não existia)",
                f"Faixa de Tensão · Potência · Marca — colunas fora da consulta, {ssos} células"])
    out.append(["Ajustes RL Poste · Ajustes Reguladores de Tensão", "abas ocultas", "", "(não existiam)",
                "cadastro de ajustes da proteção (GESTAO_DE_EQUIPAMENTOS.xlsx, 26/08), só valores"])
    return out


def compras(g5, aloc):
    """Peça reservada a quem virou Realizado, e para quem ela iria pela régua das compras."""
    ordem = sorted(g5.values(), key=lambda d: (d["Criticidade"] not in MA_ALTA,
                                               d["Índice"] if isinstance(d["Índice"], (int, float)) else 999))
    # defeito que a base e a Observação mandam trocar (a confirmar pelo gestor)
    novo_defeito = {"7926089013": "RL Completo 34,5", "7944559149": "RL Completo 34,5",
                    "7949808058": "Tanque 34,5"}
    out = []
    for a, d in g5.items():
        if d["Status"] in JA_TEM and aloc.get(a):
            for pma, q in aloc[a].items():
                peca = PECA[pma]
                fila = []
                for x in ordem:
                    b = txt(x["Ativo"])
                    if x["Status"] in JA_TEM or b == a:
                        continue
                    defeito = novo_defeito.get(b, x["Defeito"])
                    falta = Counter(BOM.get(defeito, [])) - Counter(PECA[p] for p in aloc.get(b, {})
                                                                    for _ in range(aloc[b][p]))
                    if falta.get(peca) and sum(falta.values()) == 1:
                        fila.append(f"{b} ({x['Criticidade']}, Índice {x['Índice']}"
                                    + (", se o Defeito virar " + defeito if b in novo_defeito else "")
                                    + ")")
                if PECA[pma] in ("controle RT", "célula 400"):
                    continue
                out.append([pma, peca, q, a, d["Status"], txt(d.get("Observação")),
                            fila[0] if fila else "Reserva", "; ".join(fila[1:4]),
                            "sugestão — você decide (religador)"])
    # peças de regulador: decisão do gestor em 28/09, aplicada na entrega (entrega_coep5.py)
    from entrega_coep5 import RT_TABELA7
    motivo = {"5856070091": "recebeu em 27/07 as peças do 5862236091 e já foi instalado; a PROT "
                            "recebeu pedido de ajuste para comissionar em 21/09 (ETO-PROT 00193/2026)",
              "5862236091": "COEP, 11/09: as duas células foram para obra de construção, não será "
                            "mantido este ano (ETO-COEP 00158/2026)"}
    for (pma, antes, depois), q in Counter(RT_TABELA7.values()).items():
        out.append([pma, PECA[pma], q, antes, "", motivo[antes], depois, "",
                    "feito na entrega (regulador)"])
    out.append(["38291", "controle RT", 1, "Reserva", "", "segue de reserva: nenhum outro RT fecha "
                "o conserto com ele (o 5823916001 e o 5855411017 estão em logística)", "Reserva", "",
                "feito na entrega (regulador)"])
    return out


# ------------------------------------------------------------------ planilha
def cabeca(ws, titulo, texto, cols, larg):
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13, color=SINAL)
    ws["A2"] = texto
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    ws.row_dimensions[2].height = 48
    for i, (c, w) in enumerate(zip(cols, larg), 1):
        cel = ws.cell(row=4, column=i, value=c)
        cel.font = Font(bold=True, color=PAPEL)
        cel.fill = PatternFill("solid", fgColor=TINTA)
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"


def corpo(ws, linhas):
    lado = Side(style="thin", color=FILETE)
    for j, l in enumerate(linhas):
        ws.append(l)
        r = ws.max_row
        for i in range(1, len(l) + 1):
            cel = ws.cell(row=r, column=i)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.border = Border(bottom=lado)
            if j % 2:
                cel.fill = PatternFill("solid", fgColor=PAPEL2)


def main():
    if not os.path.exists(COEP4):
        subprocess.run([sys.executable, os.path.join(RAIZ, "scripts", "compra_0309_por_ativo.py")],
                       check=True)
    g4, _ = gestao(COEP4)
    g5, col = gestao(COEP5)
    aloc = tabela7(COEP5)
    assert tabela7(COEP4) == aloc, "a lista peça a peça mudou — rever a aba Compras"
    cad = cadastro(set(g5))
    por_ss = base_ss()

    corr, visto = corrigir(g4, g5, col, aloc, cad)
    erros = sorted(corr + outras_abas(), key=lambda x: x[0])
    tem_entrega = os.path.exists(ENTREGA)
    if tem_entrega:
        erros = na_entrega(erros, g4, cad)
    mudou = veredito(o_que_mudou(g4, g5), visto, g5)
    fatos = base(g5, por_ss)
    livres = compras(g5, aloc)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Corrigir"
    cabeca(ws, "O que está errado na COEP 5 que você mandou",
           "Aba, célula, o que está e o que deveria estar. A última coluna diz o que a COEP 5 entregue "
           "em 28/09 já corrigiu e o que fica com você. As quatro primeiras linhas são o maior "
           "problema: depois de reordenar a Gestão, o orçamento e o PMA ficaram na posição antiga.",
           ["O quê", "Aba", "Célula", "Ativo", "Está", "Deveria", "Por quê", "Na entrega de 28/09"],
           [24, 16, 16, 13, 38, 42, 60, 30])
    corpo(ws, erros)

    ws = wb.create_sheet("Base de SS até 23-09")
    cabeca(ws, "O que a base de SS conta e a Gestão ainda não mostra",
           "Base RELIGA_REGULA de 23/09 — o «Religa_regula.xlsx» de 28/09 é o mesmo arquivo, então "
           "nada depois de 23/09 aparece aqui. «Leitura» é o que a SS muda na Gestão; «Parecer» é o "
           "texto mais recente da SS.",
           ["Linha", "Ativo", "Status na Gestão", "SS", "Posto", "Situação", "Aberta em",
            "Leitura", "Parecer (início)"], [7, 13, 16, 22, 13, 14, 11, 60, 70])
    corpo(ws, fatos)

    ws = wb.create_sheet("Compras")
    cabeca(ws, "Peças da compra que mudam de dono",
           "Régua das compras: Muito Alta e Alta primeiro, depois o Índice, só quem a peça conserta "
           "inteiro; Realizado e Em logística não recebem. Regulador: decisão sua de 28/09, já feita "
           "na aba Estoque e na coluna PMA. Religador: só sugestão — a aba Estoque não mudou.",
           ["PMA", "Peça", "Qtd", "Estava com", "Status", "Por quê", "Vai para",
            "Próximos da fila", "Situação"], [9, 14, 6, 13, 12, 50, 44, 50, 30])
    corpo(ws, livres)

    ws = wb.create_sheet("O que você mudou")
    cabeca(ws, "Da COEP 4 (25/09) para a COEP 5 (28/09), na aba Gestão",
           "Casado pelo ativo, não pela linha: a ordem das linhas mudou. Dias Pendente, Status "
           "Atendimento e Status Prazo não mudaram em nenhuma linha (são valor digitado, não fórmula).",
           ["Linha", "Ativo", "Campo", "COEP 4", "COEP 5", "Confere?"], [7, 13, 24, 40, 60, 30])
    corpo(ws, mudou)

    if tem_entrega:
        ws = wb.create_sheet("O que a entrega mudou")
        cabeca(ws, "Da COEP 5 que você mandou para a entregue em 28/09",
               "Célula a célula. Todo o resto — gráficos, dinâmicas, vínculos, a consulta da BASE "
               "SS_OS — ficou igual. As duas abas de ajuste estão ocultas: botão direito numa aba › "
               "Reexibir.", ["Aba", "Célula", "Ativo", "Antes", "Depois"], [22, 16, 13, 40, 60])
        corpo(ws, mudou_na_entrega())

    ws = wb.create_sheet("Como foi feito")
    for l in (__doc__ or "").strip().splitlines():
        ws.append([l])
    ws.column_dimensions["A"].width = 110

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    wb.save(SAIDA)
    print(f"{len(erros)} correções · {len(fatos)} fatos da base · {len(livres)} peças livres · "
          f"{len(mudou)} mudanças → {SAIDA}")
    return erros, fatos, livres, mudou


if __name__ == "__main__":
    main()

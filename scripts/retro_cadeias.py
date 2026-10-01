"""
A leitura das cadeias pela skill analise-equipamento (analista + verificador adversarial), mês a mês.

Lê `.analise/retro2026/retro.json` (montado por `juntar_retro.py`): uma linha por demanda — cadeia de
repasses que passou pelo COEP e teve movimento de 01/04 a 23/09/2026 — com a linha do tempo das
tratativas (data e como foi datada), se era backlog de 2025, se teve tratativa real em 2025, se o
serviço foi executado (com prova literal) e o desfecho em 23/09. O verificador manda sobre o analista.

Usado por `retrospectiva_coep.py`, que monta a pauta e a planilha.
"""

import datetime as dt
import json
import re
import os
from collections import Counter, defaultdict

from openpyxl.styles import Alignment

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RETRO = os.path.join(RAIZ, ".analise", "retro2026", "retro.json")
INICIO = dt.date(2026, 4, 1)
MESES = [(2026, m) for m in range(4, 10)]
NOME = ("janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro",
        "dezembro")
EXECUTOU = {"atendido", "executado, falta comissionar ou ajustar", "executado, sem prova de operação",
            "executado, mas não voltou a operar"}


def dia(s):
    """A data da tratativa. Quando o agente escreveu uma janela («08/07/2026 a 19/08/2026»), vale o fim dela —
    o «no mais tardar», como no parecer datado pelo repasse."""
    datas = re.findall(r"\d{2}/\d{2}/\d{4}", s or "")
    if not datas:
        return None
    try:
        return dt.datetime.strptime(datas[-1], "%d/%m/%Y").date()
    except ValueError:
        return None


def carrega(caminho=None):
    with open(caminho or RETRO, encoding="utf-8") as fh:
        d = json.load(fh)
    for r in d["demandas"]:
        r["_trat"] = [dict(t, _dia=dia(t.get("data")),
                           _janela=len(re.findall(r"\d{2}/\d{2}/\d{4}", t.get("data") or "")) > 1
                           or "janela" in (t.get("como_datou") or "") or "exports" in (t.get("como_datou") or ""))
                      for t in r.get("tratativas", []) if dia(t.get("data"))]
        r["_exec"] = dia(r.get("data_execucao")) if r.get("executada") else None
    return d


def no_periodo(t):
    return t["_dia"] >= INICIO and not t.get("antes")


def primeira_2026(r):
    xs = [t["_dia"] for t in r["_trat"] if t["_dia"].year == 2026 and not t.get("antes")]
    return min(xs) if xs else None


def por_mes(dem):
    """Por mês de abril a setembro: tratativas, demandas tocadas, ações, quem, execuções e o backlog que andou."""
    out = {}
    for a, m in MESES:
        tr = [(r, t) for r in dem for t in r["_trat"] if no_periodo(t) and (t["_dia"].year, t["_dia"].month) == (a, m)]
        ex = [r for r in dem if r["_exec"] and (r["_exec"].year, r["_exec"].month) == (a, m)]
        prim = [r for r in dem if r.get("backlog_2025") and primeira_2026(r) and (primeira_2026(r).year, primeira_2026(r).month) == (a, m)]
        out[(a, m)] = {"tratativas": len(tr), "tratativas_por_janela": sum(1 for _, t in tr if t["_janela"]),
                       "demandas": len({id(r) for r, _ in tr}),
                       "acoes": Counter(t.get("acao") for _, t in tr), "quem": Counter(t.get("quem") for _, t in tr),
                       "execucoes": len(ex), "execucoes_backlog": sum(1 for r in ex if r.get("backlog_2025")),
                       "execucoes_com_operacao": sum(1 for r in ex if r.get("voltou_a_operar") is True),
                       "backlog_primeira_tratativa": len(prim)}
    return out


def resumo(dem):
    """Conta por DEMANDA viva: a cadeia marcada como duplicada de outra (encerrada sem serviço e reaberta) é a mesma
    demanda e não conta duas vezes — fica só o número delas à parte."""
    duplicadas = [r for r in dem if (r.get("desfecho") or "").startswith("duplicada")]
    dem = [r for r in dem if r not in duplicadas]
    back = [r for r in dem if r.get("backlog_2025")]
    so_2026 = [r for r in back if not r.get("tratada_em_2025")]
    ex = [r for r in dem if r.get("executada")]
    return {"demandas": len(dem), "duplicadas": len(duplicadas), "ativos": len({r["ativo"] for r in dem}),
            "backlog": len(back), "backlog_sem_tratativa_em_2025": len(so_2026),
            "backlog_com_tratativa_em_2026": sum(1 for r in back if primeira_2026(r)),
            "desfecho": Counter(r.get("desfecho") for r in dem),
            "desfecho_backlog": Counter(r.get("desfecho") for r in back),
            "executadas": len(ex), "executadas_backlog": sum(1 for r in ex if r.get("backlog_2025")),
            "executadas_no_periodo": sum(1 for r in ex if r["_exec"] and r["_exec"] >= INICIO),
            "operando_provado": sum(1 for r in ex if r.get("voltou_a_operar") is True),
            "nao_voltou": sum(1 for r in ex if r.get("voltou_a_operar") is False),
            "sem_tratativa_no_periodo": sum(1 for r in dem if not any(no_periodo(t) for t in r["_trat"])),
            "dias_ate_executar": sorted((r["_exec"] - primeira_2026(r)).days for r in ex
                                        if r["_exec"] and primeira_2026(r) and r["_exec"] >= primeira_2026(r))}


# ------------------------------------------------------------------ abas
def aba_por_demanda(wb, dem, cabecalho, celulas):
    ws = wb.create_sheet("Por demanda")
    cols = [("Ativo", 12), ("RL/RT", 6), ("Cadeia (primeira SS)", 20), ("Aberta em", 11), ("O que era", 14),
            ("Backlog de 2025?", 9), ("Tratativa real em 2025?", 10), ("O que houve em 2025", 40),
            ("Tratativas de abr a set/2026", 9), ("Primeira tratativa de 2026", 11), ("Executada?", 9),
            ("Execução em", 11), ("Prova da execução", 50), ("Voltou a operar?", 9), ("Desfecho em 23/09", 26),
            ("Pendência que sobrou", 30), ("A verificação mudou?", 30), ("Confiança", 9)]
    cabecalho(ws, 1, [c for c, _ in cols], [w for _, w in cols])
    ws.row_dimensions[1].height = 45
    sim = lambda v: "sim" if v is True else ("não" if v is False else "não se sabe")
    linhas = []
    for r in sorted(dem, key=lambda r: (r.get("desfecho") or "", r["ativo"])):
        mud = [c for c in ("desfecho", "executada", "data_execucao", "voltou_a_operar", "backlog_2025", "tratada_em_2025")
               if c + "_analista" in r]
        linhas.append([r["ativo"], r.get("familia"), r.get("cadeia"), r.get("aberta_em"), r.get("o_que_era"),
                       sim(r.get("backlog_2025")), sim(r.get("tratada_em_2025")), r.get("o_que_houve_em_2025"),
                       sum(1 for t in r["_trat"] if no_periodo(t)),
                       primeira_2026(r).strftime("%d/%m/%Y") if primeira_2026(r) else "",
                       sim(r.get("executada")), r.get("data_execucao") or "", (r.get("prova_execucao") or "")[:400],
                       sim(r.get("voltou_a_operar")), r.get("desfecho"), r.get("pendencia_restante") or "",
                       ("; ".join(mud) + " — " + (r.get("porque_verificacao") or "")) if mud else "",
                       r.get("confianca_verificacao") or r.get("confianca")])
    celulas(ws, 2, linhas)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:R{ws.max_row}"


def aba_linha_do_tempo(wb, dem, cabecalho, celulas):
    ws = wb.create_sheet("Linha do tempo")
    cols = [("Data", 11), ("Como foi datada", 18), ("Ativo", 12), ("Cadeia", 20), ("Quem", 11), ("O que fez", 26),
            ("Resumo", 50), ("Evidência (literal)", 70), ("Backlog de 2025?", 9)]
    cabecalho(ws, 1, [c for c, _ in cols], [w for _, w in cols])
    linhas = []
    for r in dem:
        for t in r["_trat"]:
            if no_periodo(t):
                linhas.append([t["_dia"], t.get("como_datou"), r["ativo"], r.get("cadeia"), t.get("quem"), t.get("acao"),
                               t.get("resumo"), (t.get("evidencia") or "")[:300], "sim" if r.get("backlog_2025") else "não"])
    linhas.sort(key=lambda x: (x[0], x[2]))
    celulas(ws, 2, linhas)
    for c in ws["A"][1:]:
        c.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"


def aba_tratativas_mes(wb, dem, pm, cabecalho, celulas, titulo):
    ws = wb.create_sheet("Tratativas por mês")
    r0 = titulo(ws, 1, "O que foi feito nas cadeias, mês a mês — leitura de todas as SS, com verificação adversarial",
                "Uma tratativa é um passo da cadeia com data: despacho, compra, entrega, pergunta, cancelamento, execução em campo, "
                "comissionamento. A data é a escrita no texto ou a do SGM (repasse, cancelamento, conclusão). «Execuções» conta a "
                "demanda cujo serviço foi FEITO, com frase literal que prove, no mês da execução.", 8)
    acoes = sorted({a for v in pm.values() for a in v["acoes"]}, key=lambda a: -sum(v["acoes"][a] for v in pm.values()))
    quem = sorted({a for v in pm.values() for a in v["quem"]}, key=lambda a: -sum(v["quem"][a] for v in pm.values()))
    cab = ["", *[f"{NOME[m - 1][:3]}/{str(a)[2:]}" for a, m in MESES], "Total"]
    cabecalho(ws, r0, cab, [34] + [9] * (len(cab) - 1))
    linhas = [["Tratativas", *[pm[k]["tratativas"] for k in MESES], sum(pm[k]["tratativas"] for k in MESES)],
              ["   … datadas por janela (o fim dela)", *[pm[k]["tratativas_por_janela"] for k in MESES],
               sum(pm[k]["tratativas_por_janela"] for k in MESES)],
              ["Demandas com tratativa no mês", *[pm[k]["demandas"] for k in MESES], ""],
              ["Execuções com prova", *[pm[k]["execucoes"] for k in MESES], sum(pm[k]["execucoes"] for k in MESES)],
              ["   … de backlog de 2025", *[pm[k]["execucoes_backlog"] for k in MESES], sum(pm[k]["execucoes_backlog"] for k in MESES)],
              ["   … com volta à operação provada", *[pm[k]["execucoes_com_operacao"] for k in MESES],
               sum(pm[k]["execucoes_com_operacao"] for k in MESES)],
              ["Backlog de 2025 com a 1ª tratativa de 2026", *[pm[k]["backlog_primeira_tratativa"] for k in MESES],
               sum(pm[k]["backlog_primeira_tratativa"] for k in MESES)],
              ["O QUE FOI FEITO", *[""] * (len(MESES) + 1)]]
    linhas += [[f"   {a}", *[pm[k]["acoes"][a] for k in MESES], sum(pm[k]["acoes"][a] for k in MESES)] for a in acoes]
    linhas += [["QUEM FEZ", *[""] * (len(MESES) + 1)]]
    linhas += [[f"   {a}", *[pm[k]["quem"][a] for k in MESES], sum(pm[k]["quem"][a] for k in MESES)] for a in quem]
    celulas(ws, r0 + 1, linhas, negrito=(0, 3, 7, 8 + len(acoes)))
    for row in ws.iter_rows(min_row=r0 + 1, max_row=ws.max_row, min_col=1, max_col=1):
        row[0].alignment = Alignment(horizontal="left", vertical="top")

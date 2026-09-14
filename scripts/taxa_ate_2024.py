"""
A taxa de falha até 2024, com o parque de cada ano em vez do parque fixo.

Pedido do gestor (14/09), junto com a COEP 3: «aqui é taxa de falha faça até 2024».

O QUE ESTAVA ERRADO: a régua divide os três anos pelo MESMO parque — 1.307
religadores e 207 reguladores. Só que esse é o parque de 2026. Pela série mensal
reconstruída do cadastro (`parque_mensal_2024.py`), o parque médio de 2024 era
1.083 RL e 130 RT. Dividir a falha de 2024 por 1.307 e 207 é dividir por um parque
que ainda não existia — e no regulador o erro é de 37%.

O QUE MUDA:

                 parque fixo   parque real   taxa muda
    2024 RL         1.307         1.083        +20,7%
    2024 RT           207           130        +58,9%
    2025 RT           207           162        +28,0%

O CORTE HONESTO: a taxa de CHAMADA (toda ida a campo) sai daqui recalculada, porque
os eventos de cada ano já estão apurados em `taxa_falha.json`. A taxa de PEÇA GRANDE
de 2024 NÃO sai: ela precisa do texto do parecer para saber qual peça foi trocada, e
a base crua de SS/OS não está neste clone. O próprio `taxa_falha.json` registra a
lacuna — cobertura de evidência de componente de 8,1%, e em 2024 são 10 eventos
confirmados de 584 no religador.

Por isso a planilha traz TRÊS medidas lado a lado, e diz o método de cada uma:

  chamada        — todo evento no ano ÷ parque. Sai dos três anos.
  troca (AIC)    — obra de substituição concluída no ano ÷ parque. Sai dos três anos,
                   não depende de texto.
  rol do gestor  — o rol de peça grande lido SS a SS. Só existe em 2025 e 2026.

Grava dist/TAXA_ATE_2024.xlsx.
Rodar: python3 scripts/taxa_ate_2024.py
"""

import json
import os

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.drawing.colors import ColorChoice
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TAXA = os.path.join(RAIZ, "data", "missao", "taxa_falha.json")
SAIDA = os.path.join(RAIZ, "dist", "TAXA_ATE_2024.xlsx")
ROL_COEP3 = os.path.join(RAIZ, "data", "raw", "GESTAO_EQ_ESPECIAIS_COEP_3.xlsx")

TINTA, PAPEL, SINAL = "FF211D15", "FFF2EFE6", "FFBC4B0E"
COR_A, COR_B = "1F7C50", "B8480C"

PARQUE_FIXO = {"RL": 1307, "RT": 207}
ANOS = ("2024", "2025", "2026")
NOME = {"RL": "religador", "RT": "regulador"}
# o rol de peça grande da planilha do gestor — só 2025 e 2026 existem
ROL_GESTOR = {"2025": {"RL": 35, "RT": 18}, "2026": {"RL": 28, "RT": 9}}


def parque_mensal():
    """A série mensal do parque, do parque_mensal_2024. Refaz se não houver cache."""
    cache = os.path.join(RAIZ, "data", "missao", "parque_mensal.json")
    if os.path.exists(cache):
        with open(cache, encoding="utf-8") as fh:
            return json.load(fh)
    import parque_mensal_2024 as pm
    _, _, entrada, parque, meses, inicio = pm.montar(pm.PADRAO_ENTRADA)
    serie = {m: {"entrada": dict(entrada[m]), "parque": parque[m]} for m in meses}
    os.makedirs(os.path.dirname(cache), exist_ok=True)
    with open(cache, "w", encoding="utf-8") as fh:
        json.dump(serie, fh, ensure_ascii=False, indent=1)
    return serie


def medio_do_ano(serie, ano):
    """Parque médio do ano — média dos fechamentos mensais, não a ponta."""
    ms = [m for m in serie if m.startswith(ano)]
    return {t: sum(serie[m]["parque"][t] for m in ms) / len(ms) for t in ("RL", "RT")}, \
        len(ms)


def montar():
    serie = parque_mensal()
    with open(TAXA, encoding="utf-8") as fh:
        tf = json.load(fh)["serie_por_ano"]

    linhas = []
    for ano in ANOS:
        real, n_meses = medio_do_ano(serie, ano)
        for t in ("RL", "RT"):
            x = tf[ano][NOME[t]]
            fixo = PARQUE_FIXO[t]
            linhas.append({
                "ano": ano, "tipo": t, "meses": n_meses,
                "parque_fixo": fixo, "parque_real": round(real[t]),
                "eventos": x["eventos"], "ativos": x["ativos_distintos"],
                "trocas": x["trocas_confirmadas"],
                "rol": ROL_GESTOR.get(ano, {}).get(t),
                "chamada_fixo": x["eventos"] / fixo * 100,
                "chamada_real": x["eventos"] / real[t] * 100,
                "troca_fixo": x["trocas_confirmadas"] / fixo * 100,
                "troca_real": x["trocas_confirmadas"] / real[t] * 100,
                "rol_real": (ROL_GESTOR[ano][t] / real[t] * 100)
                if ano in ROL_GESTOR else None,
                "incidencia_real": x["ativos_distintos"] / real[t] * 100,
            })
    return serie, linhas


# ------------------------------------------------------------------ a planilha
def cabeca(ws, linha, titulo, cols, larguras=None):
    ws.cell(row=linha, column=1, value=titulo).font = Font(bold=True, size=11,
                                                           color=SINAL)
    ws.append(cols)
    for i in range(1, len(cols) + 1):
        c = ws.cell(row=linha + 1, column=i)
        c.font = Font(bold=True, color=PAPEL, size=10)
        c.fill = PatternFill("solid", fgColor=TINTA)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    if larguras:
        for i, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    return linha + 2


def fecha(ws, linha, n):
    for i in range(1, n + 1):
        ws.cell(row=linha, column=i).font = Font(bold=True)
        ws.cell(row=linha, column=i).border = Border(
            top=Side(style="medium", color=TINTA))


def planilha(serie, linhas):
    wb = Workbook()

    # 1 — a taxa dos três anos, fixo contra real
    ws = wb.active
    ws.title = "Taxa por ano"
    r0 = cabeca(ws, 1, "Taxa de falha por 100 equipamentos — parque fixo × parque real",
                ["Ano", "Tipo", "Parque fixo", "Parque real", "Eventos",
                 "Chamada (fixo)", "Chamada (real)", "Quanto muda",
                 "Trocas", "Troca (fixo)", "Troca (real)",
                 "Rol do gestor", "Rol ÷ parque real"],
                [8, 7, 11, 11, 9, 13, 13, 12, 8, 12, 12, 12, 15])
    for x in linhas:
        ws.append([x["ano"], x["tipo"], x["parque_fixo"], x["parque_real"],
                   x["eventos"], x["chamada_fixo"] / 100, x["chamada_real"] / 100,
                   x["chamada_real"] / x["chamada_fixo"] - 1,
                   x["trocas"], x["troca_fixo"] / 100, x["troca_real"] / 100,
                   x["rol"] if x["rol"] is not None else "não existe",
                   (x["rol_real"] / 100) if x["rol_real"] is not None else "—"])
        for col in (6, 7, 10, 11):
            ws.cell(row=ws.max_row, column=col).number_format = "0.0%"
        ws.cell(row=ws.max_row, column=8).number_format = "+0.0%;-0.0%"
        if x["rol_real"] is not None:
            ws.cell(row=ws.max_row, column=13).number_format = "0.0%"
    fecha(ws, ws.max_row, 13)
    r1 = ws.max_row

    g = BarChart()
    g.type, g.grouping = "col", "clustered"
    g.title = "Taxa de chamada: o que o parque fixo esconde"
    g.height, g.width, g.gapWidth = 9, 22, 60
    for col, cor in ((6, COR_A), (7, COR_B)):
        s = Series(Reference(ws, min_col=col, min_row=r0 - 1, max_row=r1 - 1),
                   title_from_data=True)
        s.graphicalProperties.solidFill = ColorChoice(srgbClr=cor)
        s.graphicalProperties.line.noFill = True
        g.series.append(s)
    rotulos = [f"{x['ano']} {x['tipo']}" for x in linhas]
    for n, r in enumerate(rotulos):
        ws.cell(row=r0 + n, column=15, value=r)
    cats = AxDataSource(strRef=StrRef(f=f"'{ws.title}'!$O${r0}:$O${r1 - 1}"))
    for s in g.series:
        s.cat = cats
    g.x_axis.axPos, g.y_axis.axPos = "b", "l"
    g.x_axis.delete = g.y_axis.delete = False
    g.y_axis.numFmt = "0.0%"
    ws.add_chart(g, "A" + str(r1 + 3))
    ws.column_dimensions["O"].hidden = True

    # 2 — o parque de cada ano, mês a mês, que é de onde sai o denominador
    ws2 = wb.create_sheet("Parque de cada ano")
    cabeca(ws2, 1, "O parque mês a mês — a origem do denominador real",
           ["Mês", "Parque RL", "Parque RT"], [12, 12, 12])
    for mes in sorted(serie):
        ws2.append([mes, serie[mes]["parque"]["RL"], serie[mes]["parque"]["RT"]])
    fecha(ws2, ws2.max_row, 3)
    ws2.append([])
    ws2.append(["Ano", "Meses", "Parque médio RL", "Parque médio RT"])
    for i in range(1, 5):
        c = ws2.cell(row=ws2.max_row, column=i)
        c.font, c.fill = Font(bold=True, color=PAPEL, size=10), \
            PatternFill("solid", fgColor=TINTA)
    for ano in ANOS:
        real, n = medio_do_ano(serie, ano)
        ws2.append([ano, n, round(real["RL"]), round(real["RT"])])
    ws2.append([])
    ws2.append(["A série sai de parque_mensal_2024.py: cadastro do ativo (DTA_ORIG) "
                "filtrado de atualização, ancorado no parque de agosto de 2026."])

    # 3 — o que falta para fechar 2024 de verdade
    ws3 = wb.create_sheet("O que falta em 2024")
    ws3.column_dimensions["A"].width = 98
    for t in [
        "O que NÃO dá para fazer de 2024, e por quê",
        "",
        "A taxa de PEÇA GRANDE de 2024 não sai daqui:",
        "A régua do gestor conta falha pela peça — controle, tanque/parte ativa ou "
        "equipamento completo no religador; célula, relé, banco completo ou furto no "
        "regulador. Saber qual peça foi trocada exige LER o parecer da SS.",
        "",
        "O texto do parecer não está neste clone. A base crua de SS/OS "
        "(BASE_SS_OS_ddmmaaaa.txt, 36 MB) está no .gitignore, e o recorte local "
        "(ssos_min.json, 6.362 SS) guarda só os campos estruturados — número, datas, "
        "situação, TIPOSS — sem a descrição.",
        "",
        "O tamanho da lacuna está medido no próprio taxa_falha.json:",
        "cobertura de evidência de componente = 8,1%. Em 2024, no religador, são 10 "
        "eventos com peça grande confirmada em 584 — os outros 574 não têm evidência "
        "de componente. Chamar isso de «taxa de falha de 2024» seria inventar.",
        "",
        "As três medidas da primeira aba, e o que cada uma aguenta:",
        "",
        "CHAMADA — todo evento do ano ÷ parque. Sai dos três anos e é a mais sólida, "
        "porque só depende de data de ocorrência e código do ativo. É a que responde "
        "«quantas vezes por ano esse parque chama equipe».",
        "",
        "TROCA (AIC) — obra de substituição concluída no ano ÷ parque. Sai dos três "
        "anos e também não depende de texto, porque vem do AIC. É a mais perto de "
        "«peça grande», mas conta só o que virou obra paga.",
        "",
        "ROL DO GESTOR — as 90 linhas lidas SS a SS, com peça, causa raiz e citação. "
        "É a melhor das três, e só existe em 2025 (35 RL + 18 RT) e 2026 (28 RL + 9 "
        "RT). Para 2024 não foi feita.",
        "",
        "O que fecharia 2024:",
        "Um export da base de SS/OS com a coluna de descrição, cobrindo 2024. Com ele "
        "o rol de 2024 se monta pela mesma régua das 90 linhas, e aí as três colunas "
        "ficam comparáveis de verdade.",
    ]:
        ws3.append([t])
        ws3.cell(row=ws3.max_row, column=1).alignment = Alignment(wrap_text=True)
    ws3["A1"].font = Font(bold=True, size=12, color=SINAL)
    for n in range(1, ws3.max_row + 1):
        v = ws3.cell(row=n, column=1).value
        if v and str(v).endswith(":"):
            ws3.cell(row=n, column=1).font = Font(bold=True, size=11)

    # 4 — a régua
    ws4 = wb.create_sheet("Como foi feito")
    ws4.column_dimensions["A"].width = 98
    for t in [
        "A taxa até 2024, com o parque de cada ano",
        "",
        "O denominador mudou, e é só isso:",
        "A régua oficial divide os três anos pelo mesmo parque — 1.307 religadores e "
        "207 reguladores —, que é o parque de 2026. Aqui cada ano é dividido pelo "
        "parque médio DAQUELE ano, tirado da série mensal reconstruída do cadastro.",
        "",
        "Parque médio por ano (média dos doze fechamentos mensais):",
        "  2024 — RL 1.083 · RT 130",
        "  2025 — RL 1.234 · RT 162",
        "  2026 — RL 1.287 · RT 182   (oito meses)",
        "",
        "O numerador não mudou:",
        "Eventos, ativos distintos e trocas confirmadas vêm de taxa_falha.json, como "
        "já estavam. Nenhuma falha foi reclassificada aqui.",
        "",
        "Onde o erro é grande:",
        "No regulador de 2024 — 207 contra 130 reais. A taxa de chamada sobe de 33,8 "
        "para 53,7 por 100, quase 59% a mais. O regulador é onde o parque mais cresceu "
        "(66,7% em dois anos e meio), então é onde o parque fixo mais engana.",
        "",
        "No religador de 2026 quase não muda (1,5%), porque 1.307 é justamente o "
        "parque de agora. Quanto mais para trás no tempo, maior o erro.",
        "",
        "A ressalva do parque:",
        "A série mensal vem do cadastro e não enxerga BAIXA de equipamento — retirada, "
        "desativação, troca de código. Se houve baixa no período, o parque de 2024 era "
        "maior que 1.083 e a correção aqui é um teto, não um ponto exato.",
    ]:
        ws4.append([t])
        ws4.cell(row=ws4.max_row, column=1).alignment = Alignment(wrap_text=True)
    ws4["A1"].font = Font(bold=True, size=12, color=SINAL)
    for n in range(1, ws4.max_row + 1):
        v = ws4.cell(row=n, column=1).value
        if v and str(v).endswith(":"):
            ws4.cell(row=n, column=1).font = Font(bold=True, size=11)

    for aba in wb.worksheets:
        for linha in aba.iter_rows():
            for c in linha:
                if c.value is not None:
                    c.font = Font(name="Arial", bold=c.font.bold,
                                  size=c.font.size or 10, color=c.font.color)

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    wb.save(SAIDA)
    return SAIDA


def main():
    serie, linhas = montar()
    print(f"{'ano':6s} {'tipo':5s} {'fixo':>6s} {'real':>6s} "
          f"{'chamada fixo':>13s} {'chamada real':>13s} {'muda':>8s}")
    for x in linhas:
        print(f"{x['ano']:6s} {x['tipo']:5s} {x['parque_fixo']:6d} "
              f"{x['parque_real']:6d} {x['chamada_fixo']:12.1f} "
              f"{x['chamada_real']:12.1f} "
              f"{(x['chamada_real'] / x['chamada_fixo'] - 1) * 100:+7.1f}%")
    print("gravado:", planilha(serie, linhas))


if __name__ == "__main__":
    main()

"""
COEP 5 — entrega de 28/09: bases de ajuste da proteção ocultas, três colunas na BASE SS_OS, peças
de RT da primeira compra redistribuídas e as correções certas da conferência.

Pedidos do gestor em 28/09, na ordem:
1. «verifica as informações da planilha … só verifica possíveis atualizações» — virou
   `confere_coep5.py` (a lista do que está errado e do que a base de SS pede).
2. «tem também que atualizar as peças que compramos pra RT também … só as de regulador».
   Decisão dele (28/09, entre três opções): **liberar as peças do 5856070091 e do 5862236091 e
   redistribuir pela régua**. A base de SS mostra que o 5856070091 recebeu em 27/07 as peças que
   eram do 5862236091 («peças redirecionadas 5856070091 prioridade maior») e já foi instalado — a
   PROT recebeu pedido de ajuste para comissionar em 21/09 —, e que o 5862236091 não vai ser mantido
   este ano (COEP, 11/09: «utilizamos as duas células pra outra obra de construção»). Pela régua da
   primeira compra (Muito Alta e Alta primeiro, depois o Índice, só quem a peça conserta inteiro,
   não recebe «Em logística» nem «Realizado», o resto é Reserva): 1 célula 400 → 5836786094;
   3 células 400 + 1 controle de RT → 5856156091 (RT Completo); o controle de reserva segue reserva.
3. «coloca a base de ajuste de proteção e oculta elas … na aba de SS OS tem uma query que puxa as
   SS que têm 79 e 58 e estão no tipo indisponibilidade para operação do DCMD; seria muito
   interessante se colocasse uma coluna, sem atrapalhar a query, pra que venha a faixa de tensão, a
   potência e a marca». As abas «Ajustes RL Poste» e «Ajustes Reguladores de Tensão» entram
   **ocultas**, com os valores do cadastro (`GESTAO_DE_EQUIPAMENTOS.xlsx`, gravado em 26/08). A
   BASE SS_OS é uma consulta do Power Query (lê `L:\\COEP\\BASE SS_OS.txt`); as três colunas novas
   entram **à direita, fora da consulta** (colunas não ligadas da tabela — `dataBound="0"`), como a
   «Coluna1» que já existia. Ao atualizar a consulta, elas ficam e a fórmula desce sozinha.
4. «Sim, corrigir o que é certo»: as 4 linhas com orçamento e PMA de outra linha, os campos vazios das
   5 linhas novas (do cadastro de ajustes e da carteira), ativo gravado como texto, fórmulas presas na
   pasta Downloads, #N/A da Descrição, alimentador apagado e as fórmulas quebradas da aba SLA. O que
   depende de decisão dele (defeito do 7926089013, 7949808058 e 7955946007; orçamento do 5823916001;
   Dias Pendente e Índice das linhas novas) fica só na lista.

Edição direta no XML, como nas entregas anteriores: a COEP 5 tem 8 gráficos, 5 dinâmicas, 4
vínculos externos e uma consulta do Power Query — regravar com openpyxl perderia parte disso.

Rodar: python3 scripts/entrega_coep5.py
"""

import datetime as dt
import os
import re
import sys
import uuid
import zipfile
from collections import Counter
from xml.etree import ElementTree
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from compra_2907_por_ativo import Textos, todas_as_celulas  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.path.join(RAIZ, "data", "raw", "GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_5.xlsx")
AJUSTES = os.path.join(RAIZ, "data", "raw", "GESTAO_DE_EQUIPAMENTOS.xlsx")
SAIDA = os.path.join(RAIZ, "dist", "GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_5.xlsx")

ABAS_AJUSTE = [("Ajustes RL Poste", 14), ("Ajustes Reguladores de Tensão", 15)]   # colunas A:N e A:O
RL, RT = "'Ajustes RL Poste'", "'Ajustes Reguladores de Tensão'"
COD = "VALUE(TRIM(BASE_SS_OS[[#This Row],[NUM_TRAFO]]))"
# As três colunas da BASE SS_OS. Código RL na coluna A das duas abas; RL: RELE=K, TENSÃO=M;
# RT: PARTE ATIVA=D, POTÊNCIA=E, TENSÃO PRIMÁRIA=I. O cadastro grava a tensão como «34.500»,
# «13.800» ou 34500 — as duas primeiras letras bastam para a faixa.
NOVAS = [
    ("Faixa de Tensão",
     f'IFERROR(_xlfn.SWITCH(LEFT(_xlfn.XLOOKUP({COD},{RL}!$A:$A,{RL}!$M:$M,'
     f'_xlfn.XLOOKUP({COD},{RT}!$A:$A,{RT}!$I:$I,""))&"",2),"34","34,5 kV","13","13,8 kV",""),"")'),
    ("Potência",
     f'IFERROR(IF(_xlfn.XLOOKUP({COD},{RT}!$A:$A,{RT}!$E:$E)&""="","",'
     f'_xlfn.XLOOKUP({COD},{RT}!$A:$A,{RT}!$E:$E)&" kVA"),"")'),
    ("Marca",
     f'IFERROR(_xlfn.XLOOKUP({COD},{RL}!$A:$A,{RL}!$K:$K,'
     f'_xlfn.XLOOKUP({COD},{RT}!$A:$A,{RT}!$D:$D,""))&"","")'),
]

# ---- Gestão: o que é certo corrigir (conferência de 28/09) ----------------------------------------
# 1) as 4 linhas que ficaram com orçamento e PMA de outra linha depois de reordenar a tabela: volta o
#    da COEP 4 entregue em 25/09, que é o que a lista peça a peça da aba Estoque dá a cada ativo
ORCAMENTO = {
    "5800961074": (80318.5, 51705.75, 132024.25, "Sem PMA"),
    "5853360007": (0, 15000, 15000, "Sem PMA"),
    "7900535058": (13209.34, 76246.2, 89455.54, "38295+38297"),
    "7920024127": (13209.34, 76246.2, 89455.54, "38295+38297"),
}
# 2) linhas novas: marca, tensão e potência do cadastro de ajustes; alimentador e município da base
#    de SS; polo e regional da carteira (ATUALIZADA 16) ou, fora dela, do mesmo município na Gestão
NOVAS_LINHAS = {
    "5844630060": {"Potencia": "239 kVA", "PMA": "Sem PMA", "Alimentador": "AL01054060",
                   "Município": "NOVA OLINDA", "Polo": "COLINAS DO TOCANTINS", "Regional": "NORTE"},
    "5800859011": {"Ativo": 5800859011, "PMA": "Sem PMA", "Alimentador": "AL01065015",
                   "Município": "TUPIRAMA", "Polo": "GUARAI", "Regional": "NORTE"},
    "7905320122": {"SLA_Total": 26, "Alimentador": "AL06060122",
                   "Município": "PALMAS", "Polo": "PALMAS", "Regional": "CENTRO"},
    "5823916001": {"Ativo": 5823916001, "Marca": "1 e 2 TOSHIBA, 3 ITB", "Faixa de Tensão": "13,8 kV",
                   "Tipo": "=58", "Potencia": "167 kVA", "SS SGM": "ETO-COEP 00196/2026",
                   "PMA": "Sem PMA", "Alimentador": "AL06072001",
                   "Município": "PORTO NACIONAL", "Polo": "PORTO NACIONAL", "Regional": "CENTRO"},
    "7949808058": {"Ativo": 7949808058, "Marca": "NOJA RC10", "Faixa de Tensão": "34,5 kV",
                   "Tipo": "=79", "SS SGM": "ETO-COEP 00195/2026", "PMA": "Sem PMA",
                   "Alimentador": "LD02010153", "Município": "SANTA FE DO ARAGUAIA",
                   "Polo": "ARAGUAINA", "Regional": "NORTE"},
}
ALIMENTADOR_APAGADO = {"7947203070": "LD02065015", "7900525015": "LD02065015",
                       "7925733015": "LD02065015"}
# 3) peças de RT (decisão do gestor, 28/09): quem recebe cada unidade na Tabela7 e o PMA na Gestão
RT_TABELA7 = {  # linha da Tabela7: (PMA, estava com, passa a)
    36: ("38291", "5856070091", "5856156091"),
    39: ("38292", "5856070091", "5856156091"),
    40: ("38292", "5856070091", "5856156091"),
    41: ("38292", "5856070091", "5856156091"),
    42: ("38292", "5862236091", "5836786094"),
}
RT_PMA = {"5856070091": "Sem PMA", "5862236091": "Sem PMA", "5836786094": "38292",
          "5856156091": "38291+38292"}
# 4) aba SLA de Manutenção
SLA_CEL = {
    "I44": ("f", 'IFERROR(I42/(I42+I43),"")', 1),     # «es» no lugar da fórmula de julho
    "A107": ("s", "Total", None),                    # a linha soma todas as equipes
    "P92": ("f", "K133", 13), "P95": ("f", "K136", 1),
    "P98": ("f", "K139", 2), "P101": ("f", "K142", 5),
}


# ------------------------------------------------------------------ XML
class _Celula:
    def __init__(self, xml, a, b):
        self.a, self.b, self.s = a, b, xml[a:b]

    def start(self):
        return self.a

    def end(self):
        return self.b

    def group(self, i=0):
        return self.s


def celula(xml, ref):
    """A célula inteira. Não dá para achar por regex só: numa célula vazia (`<c r="Z57" s="1"/>`)
    um `>.*?</c>` atravessa a linha e engole o começo da próxima — pega a tag de abertura primeiro."""
    m = re.search(rf'<c r="{ref}"[\s/>]', xml)
    if not m:
        return None
    fim_tag = xml.index(">", m.start())
    if xml[fim_tag - 1] == "/":
        return _Celula(xml, m.start(), fim_tag + 1)
    return _Celula(xml, m.start(), xml.index("</c>", fim_tag) + 4)


def estilo(xml, ref):
    m = celula(xml, ref)
    if not m:
        return None
    s = re.search(r' s="(\d+)"', m.group(0)[:m.group(0).find(">") + 1])
    return s.group(1) if s else None


def monta(ref, valor, s, textos, formula=None, extra=""):
    est = f' s="{s}"' if s is not None else ""
    if formula is not None:
        if isinstance(valor, (int, float)):
            return f'<c r="{ref}"{est}{extra}><f>{escape(formula)}</f><v>{valor}</v></c>'
        return f'<c r="{ref}"{est} t="str"{extra}><f>{escape(formula)}</f><v>{escape(str(valor))}</v></c>'
    if isinstance(valor, str):
        return f'<c r="{ref}"{est} t="s"><v>{textos.de(valor)}</v></c>'
    return f'<c r="{ref}"{est}><v>{valor}</v></c>'


def poe(xml, ref, nova):
    """Troca a célula `ref` por `nova`, ou insere na ordem das colunas da linha."""
    m = celula(xml, ref)
    if m:
        return xml[:m.start()] + nova + xml[m.end():]
    col, lin = re.match(r"([A-Z]+)(\d+)", ref).groups()
    ci = column_index_from_string(col)
    lm = re.search(rf'<row r="{lin}"[^>]*?(?:/>|>(.*?)</row>)', xml, re.S)
    assert lm, f"linha {lin} não existe"
    if lm.group(0).endswith("/>"):
        nova_linha = lm.group(0)[:-2] + ">" + nova + "</row>"
        return xml[:lm.start()] + nova_linha + xml[lm.end():]
    corpo, ini = lm.group(1), lm.start(1)
    for c in re.finditer(r'<c r="([A-Z]+)\d+"', corpo):
        if column_index_from_string(c.group(1)) > ci:
            pos = ini + c.start()
            return xml[:pos] + nova + xml[pos:]
    pos = lm.end(1)
    return xml[:pos] + nova + xml[pos:]


def linha_da_tabela(xml_tabela):
    ref = re.search(r'<table [^>]*ref="([^"]+)"', xml_tabela).group(1)
    return ref


def mapa_partes(z):
    wb = z.read("xl/workbook.xml").decode()
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    alvo = {m.group(1): m.group(2) for m in re.finditer(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
    alvo.update({m.group(2): m.group(1) for m in re.finditer(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    abas = {m.group(1): "xl/" + alvo[m.group(2)].lstrip("/").replace("xl/", "")
            for m in re.finditer(r'<sheet name="([^"]+)"[^>]*r:id="([^"]+)"', wb)}
    tabelas = {}
    for n in z.namelist():
        if re.match(r"xl/tables/table\d+\.xml$", n):
            tabelas[re.search(r'<table [^>]*name="([^"]+)"', z.read(n).decode()).group(1)] = n
    return abas, tabelas


def aba_oculta(linhas, ncol):
    """Planilha só com valores (texto inline, sem mexer na lista de textos compartilhados)."""
    rows = []
    for i, r in enumerate(linhas, 1):
        cs = []
        for j in range(ncol):
            v = r[j] if j < len(r) else None
            if v is None or v == "":
                continue
            ref = f"{get_column_letter(j + 1)}{i}"
            if isinstance(v, bool):
                cs.append(f'<c r="{ref}" t="b"><v>{int(v)}</v></c>')
            elif isinstance(v, (int, float)):
                cs.append(f'<c r="{ref}"><v>{v}</v></c>')
            elif isinstance(v, (dt.datetime, dt.date)):
                serial = (dt.datetime(v.year, v.month, v.day) - dt.datetime(1899, 12, 30)).days
                cs.append(f'<c r="{ref}" s="12"><v>{serial}</v></c>')   # xf 12 = data (numFmt 14)
            else:
                cs.append(f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">'
                          f'{escape(str(v))}</t></is></c>')
        rows.append(f'<row r="{i}">{"".join(cs)}</row>')
    ult = f"{get_column_letter(ncol)}{len(linhas)}"
    cols = "".join(f'<col min="{j}" max="{j}" width="{16 if j > 1 else 13}" customWidth="1"/>'
                   for j in range(1, ncol + 1))
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<dimension ref="A1:{ult}"/><sheetViews><sheetView workbookViewId="0">'
            '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
            '</sheetView></sheetViews><sheetFormatPr defaultRowHeight="15"/>'
            f'<cols>{cols}</cols><sheetData>{"".join(rows)}</sheetData>'
            '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
            '</worksheet>')


# ------------------------------------------------------------------ dados
def ajustes():
    """As duas abas do cadastro, só com os valores e só até a última linha com código."""
    wb = openpyxl.load_workbook(AJUSTES, data_only=True)
    out = {}
    for nome, ncol in ABAS_AJUSTE:
        linhas = [list(r[:ncol]) for r in wb[nome].iter_rows(values_only=True)]
        while linhas and linhas[-1][0] in (None, ""):
            linhas.pop()
        out[nome] = linhas
    return out


def calcula_novas(base_ssos, aj):
    """O que as três fórmulas devolvem — vai como valor em cache (o Excel recalcula ao abrir)."""
    rl = {r[0]: r for r in aj["Ajustes RL Poste"][1:] if isinstance(r[0], int)}
    rt = {r[0]: r for r in aj["Ajustes Reguladores de Tensão"][1:] if isinstance(r[0], int)}
    out = []
    for cod in base_ssos:
        try:
            k = int(str(cod).strip())
        except ValueError:
            out.append(("", "", ""))
            continue
        a, b = rl.get(k), rt.get(k)
        t = str(a[12] if a and a[12] is not None else (b[8] if b and b[8] is not None else ""))
        faixa = {"34": "34,5 kV", "13": "13,8 kV"}.get(t[:2], "")
        pot = "" if not b or b[4] in (None, "") else f"{b[4]} kVA"
        marca = str(a[10] if a and a[10] is not None else (b[3] if b and b[3] is not None else ""))
        out.append((faixa, pot, marca))
    return out


# ------------------------------------------------------------------ edição
def edita(parts, abas, tabelas, textos, aj):
    log = []
    ge = parts[abas["Gestão"]].decode()
    fim = int(re.sub(r"\D", "", linha_da_tabela(parts[tabelas["Table1"]].decode()).split(":")[1]))
    cab = {}
    for m in re.finditer(r'<c r="([A-Z]+)1"[^>]*t="s"[^>]*><v>(\d+)</v></c>', linha_xml_(ge, 1)):
        cab[textos.texto(int(m.group(2)))] = m.group(1)
    col = cab.__getitem__
    linha_de = {}
    for n in range(2, fim + 1):
        m = celula(ge, f"A{n}")
        v = re.search(r"<v>(.*?)</v>", m.group(0)).group(1)
        a = textos.texto(int(v)) if 't="s"' in m.group(0) else v
        linha_de[a.strip()] = n

    def gp(a, campo, valor, formula=None, s=None):
        nonlocal ge
        ref = f"{col(campo)}{linha_de[a]}"
        if s is None:
            # célula que já existe guarda o estilo dela; célula nova herda o da mesma coluna na linha 54
            s = estilo(ge, ref) if celula(ge, ref) else estilo(ge, f"{col(campo)}54")
        ge = poe(ge, ref, monta(ref, valor, s, textos, formula))
        log.append(("Gestão", ref, a, campo, valor))

    for a, (mo, mat, tot, pma) in ORCAMENTO.items():
        for campo, v in zip(("Orçamento MO", "Orçamento MAT", "Orçamento Total", "PMA"), (mo, mat, tot, pma)):
            gp(a, campo, v)
    for a, campos in NOVAS_LINHAS.items():
        for campo, v in campos.items():
            if campo == "Tipo":
                gp(a, campo, v[1:], formula="LEFT(Table1[[#This Row],[Ativo]],2)", s="133")
            else:
                gp(a, campo, v)
    for a, v in ALIMENTADOR_APAGADO.items():
        gp(a, "Alimentador", v)
    for a, v in RT_PMA.items():
        gp(a, "PMA", v)

    # Descrição: SEERRO em volta do PROCX, para a SS que não está mais pendente não virar #N/A
    cd = col("Descrição")
    n_desc = 0
    for n in range(2, fim + 1):
        m = celula(ge, f"{cd}{n}")
        if not m or "<f" not in m.group(0):
            continue
        c = m.group(0)
        f = re.search(r"(<f[^>]*>)(.*?)(</f>)", c, re.S)
        if f.group(2).startswith("IFERROR("):
            continue
        novo = c[:f.start(2)] + f"IFERROR({f.group(2)},&quot;&quot;)" + c[f.end(2):]
        if 't="e"' in novo:
            novo = novo.replace('t="e"', 't="str"')
            novo = re.sub(r"<v>#N/A</v>", "<v></v>", novo)
            log.append(("Gestão", f"{cd}{n}", "", "Descrição", "#N/A → vazio"))
        ge = ge[:m.start()] + novo + ge[m.end():]
        n_desc += 1
    parts[abas["Gestão"]] = ge.encode()

    # Estoque: a Tabela7 diz quem recebe cada unidade de RT
    es = parts[abas["Estoque"]].decode()
    for lin, (pma, antes, depois) in RT_TABELA7.items():
        m = celula(es, f"A{lin}")
        assert re.search(rf"<v>{pma}</v>", m.group(0)), (lin, m.group(0))
        k = celula(es, f"K{lin}")
        assert textos.texto(int(re.search(r"<v>(\d+)</v>", k.group(0)).group(1))) == antes, (lin, antes)
        es = poe(es, f"K{lin}", monta(f"K{lin}", depois, estilo(es, f"K{lin}"), textos))
        log.append(("Estoque", f"K{lin}", antes, pma, depois))
    parts[abas["Estoque"]] = es.encode()

    # SLA de Manutenção
    sl = parts[abas["SLA de Manutenção"]].decode()
    for ref, (tipo, v, cache) in SLA_CEL.items():
        s = estilo(sl, ref)
        nova = monta(ref, cache, s, textos, formula=v) if tipo == "f" else monta(ref, v, s, textos)
        sl = poe(sl, ref, nova)
        log.append(("SLA de Manutenção", ref, "", "", v))
    parts[abas["SLA de Manutenção"]] = sl.encode()
    return log, n_desc


def linha_xml_(xml, n):
    m = re.search(rf'<row r="{n}"[^>]*?(?:/>|>.*?</row>)', xml, re.S)
    return m.group(0)


def colunas_ssos(parts, abas, tabelas, textos, aj):
    """Três colunas não ligadas à direita da consulta, com fórmula de coluna calculada."""
    tab_path = tabelas["BASE_SS_OS"]
    tb = parts[tab_path].decode()
    ref = linha_da_tabela(tb)                          # A1:BM54
    c0, c1 = re.match(r"A1:([A-Z]+)(\d+)", ref).groups()
    ult_lin = int(c1)
    ci = column_index_from_string(c0)
    novas_cols = [get_column_letter(ci + i + 1) for i in range(len(NOVAS))]
    novo_ref = f"A1:{novas_cols[-1]}{ult_lin}"

    # tabela
    ids = [int(x) for x in re.findall(r'<tableColumn id="(\d+)"', tb)]
    prox = max(ids) + 1
    extra = ""
    for i, (nome, formula) in enumerate(NOVAS):
        uid = "{" + str(uuid.uuid4()).upper() + "}"
        extra += (f'<tableColumn id="{prox + i}" xr3:uid="{uid}" uniqueName="{prox + i}" '
                  f'name="{escape(nome)}" queryTableFieldId="{prox + i}">'
                  f'<calculatedColumnFormula>{escape(formula)}</calculatedColumnFormula></tableColumn>')
    n = len(ids) + len(NOVAS)
    tb = re.sub(r'<tableColumns count="\d+">', f'<tableColumns count="{n}">', tb)
    tb = tb.replace("</tableColumns>", extra + "</tableColumns>")
    tb = tb.replace(f'ref="{ref}"', f'ref="{novo_ref}"')
    tb = tb.replace(f'ref="A2:{c0}{ult_lin}"', f'ref="A2:{novas_cols[-1]}{ult_lin}"')
    parts[tab_path] = tb.encode()

    # consulta: campos não ligados à direita
    rel = parts[tab_path.replace("tables/", "tables/_rels/") + ".rels"].decode()
    q_path = "xl/" + re.search(r'Target="\.\./([^"]+)"', rel).group(1)
    q = parts[q_path].decode()
    nq = int(re.search(r'<queryTableFields count="(\d+)"', q).group(1))
    livre = int(re.search(r'unboundColumnsRight="(\d+)"', q).group(1))
    campos = "".join(f'<queryTableField id="{prox + i}" dataBound="0" tableColumnId="{prox + i}"/>'
                     for i in range(len(NOVAS)))
    q = q.replace(f'<queryTableFields count="{nq}">', f'<queryTableFields count="{nq + len(NOVAS)}">')
    q = q.replace("</queryTableFields>", campos + "</queryTableFields>")
    q = re.sub(r'nextId="\d+"', f'nextId="{prox + len(NOVAS)}"', q)
    q = q.replace(f'unboundColumnsRight="{livre}"', f'unboundColumnsRight="{livre + len(NOVAS)}"')
    parts[q_path] = q.encode()

    # aba: cabeçalho, fórmula e valor em cache em cada linha
    sh = parts[abas["BASE SS_OS"]].decode()
    ntrafo_col = None
    for m in re.finditer(r'<c r="([A-Z]+)1"[^>]*t="s"[^>]*><v>(\d+)</v></c>', linha_xml_(sh, 1)):
        if textos.texto(int(m.group(2))) == "NUM_TRAFO":
            ntrafo_col = m.group(1)
    codigos = []
    for lin in range(2, ult_lin + 1):
        m = celula(sh, f"{ntrafo_col}{lin}")
        v = re.search(r"<v>(.*?)</v>", m.group(0)).group(1) if m else ""
        codigos.append(textos.texto(int(v)) if m and 't="s"' in m.group(0) else v)
    valores = calcula_novas(codigos, aj)
    for i, (nome, formula) in enumerate(NOVAS):
        c = novas_cols[i]
        sh = poe(sh, f"{c}1", monta(f"{c}1", nome, None, textos))
        for lin in range(2, ult_lin + 1):
            sh = poe(sh, f"{c}{lin}", monta(f"{c}{lin}", valores[lin - 2][i], None, textos, formula=formula))
    sh = sh.replace(f'<dimension ref="{ref}"/>', f'<dimension ref="{novo_ref}"/>')
    sh = re.sub(r'spans="1:65"', f'spans="1:{ci + len(NOVAS)}"', sh)
    larg = "".join(f'<col min="{ci + i + 1}" max="{ci + i + 1}" width="{w}" customWidth="1"/>'
                   for i, w in enumerate((15, 12, 22)))
    sh = sh.replace("</cols>", larg + "</cols>")
    parts[abas["BASE SS_OS"]] = sh.encode()
    return novas_cols, codigos, valores


def abas_ocultas(parts, aj):
    wb = parts["xl/workbook.xml"].decode()
    rels = parts["xl/_rels/workbook.xml.rels"].decode()
    ct = parts["[Content_Types].xml"].decode()
    app = parts["docProps/app.xml"].decode()
    n_aba = max(int(x) for x in re.findall(r"xl/worksheets/sheet(\d+)\.xml", " ".join(parts)))
    sid = max(int(x) for x in re.findall(r'sheetId="(\d+)"', wb))
    novas = []
    for i, (nome, ncol) in enumerate(ABAS_AJUSTE, 1):
        caminho = f"xl/worksheets/sheet{n_aba + i}.xml"
        rid = f"rId{900 + i}"
        assert f'Id="{rid}"' not in rels
        parts[caminho] = aba_oculta(aj[nome], ncol).encode()
        wb = wb.replace("</sheets>", f'<sheet name="{escape(nome)}" sheetId="{sid + i}" state="hidden" '
                                     f'r:id="{rid}"/></sheets>')
        rels = rels.replace("</Relationships>",
                            f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/'
                            f'officeDocument/2006/relationships/worksheet" '
                            f'Target="worksheets/sheet{n_aba + i}.xml"/></Relationships>')
        ct = ct.replace("</Types>", f'<Override PartName="/{caminho}" ContentType="application/'
                                    f'vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
        novas.append(nome)
    # docProps/app.xml: a lista de abas
    m = re.search(r"<vt:lpstr>Planilhas</vt:lpstr></vt:variant><vt:variant><vt:i4>(\d+)</vt:i4>", app)
    if m:
        k = int(m.group(1))
        app = app.replace(m.group(0), m.group(0).replace(f"<vt:i4>{k}</vt:i4>", f"<vt:i4>{k + 2}</vt:i4>"))
        t = re.search(r'<TitlesOfParts><vt:vector size="(\d+)" baseType="lpstr">(.*?)</vt:vector>', app, re.S)
        novo = (f'<TitlesOfParts><vt:vector size="{int(t.group(1)) + 2}" baseType="lpstr">{t.group(2)}'
                + "".join(f"<vt:lpstr>{escape(x)}</vt:lpstr>" for x in novas) + "</vt:vector>")
        app = app.replace(t.group(0), novo)
    parts["xl/workbook.xml"], parts["xl/_rels/workbook.xml.rels"] = wb.encode(), rels.encode()
    parts["[Content_Types].xml"], parts["docProps/app.xml"] = ct.encode(), app.encode()


def dinamicas(parts):
    """Orçamento: atualiza ao abrir. SLA: lê a Tabela1 da própria planilha, não o «v6 (1).xlsx»."""
    feito = []
    for n in list(parts):
        if not re.match(r"xl/pivotCache/pivotCacheDefinition\d+\.xml$", n):
            continue
        x = parts[n].decode()
        if 'worksheetSource name="Table1"' in x or 'worksheetSource name="Tabela1"' in x:
            rel = n.replace("pivotCache/", "pivotCache/_rels/") + ".rels"
            if 'worksheetSource name="Tabela1" r:id=' in x:
                rid = re.search(r'worksheetSource name="Tabela1" r:id="([^"]+)"', x).group(1)
                x = x.replace(f' r:id="{rid}"/>', "/>", 1)
                r = parts[rel].decode()
                r2 = re.sub(rf'<Relationship Id="{rid}"[^>]*/>', "", r)
                assert r2 != r and "Downloads" not in r2
                parts[rel] = r2.encode()
                feito.append((n, "fonte: Tabela1 da aba Base SLA de Manutenção"))
            if "refreshOnLoad=" not in x:
                x = re.sub(r'(<pivotCacheDefinition [^>]*?)( refreshedBy=)', r'\1 refreshOnLoad="1"\2', x, count=1)
                assert 'refreshOnLoad="1"' in x, n
                feito.append((n, "atualiza ao abrir"))
            parts[n] = x.encode()
    return feito


def sem_calcchain(parts):
    """Fórmula que vira valor deixa entrada velha no calcChain, e o Excel pede reparo (lição de 24/09).
    Sai a parte; o Excel refaz ao abrir. E recálculo completo ao abrir, pelas fórmulas novas."""
    parts.pop("xl/calcChain.xml", None)
    ct = parts["[Content_Types].xml"].decode()
    ct = re.sub(r'<Override PartName="/xl/calcChain.xml"[^>]*/>', "", ct)
    parts["[Content_Types].xml"] = ct.encode()
    r = parts["xl/_rels/workbook.xml.rels"].decode()
    r = re.sub(r'<Relationship [^>]*Target="calcChain.xml"/>', "", r)
    parts["xl/_rels/workbook.xml.rels"] = r.encode()
    wb = parts["xl/workbook.xml"].decode()
    if "fullCalcOnLoad" not in wb:
        wb = re.sub(r"<calcPr ", '<calcPr fullCalcOnLoad="1" ', wb, count=1)
    parts["xl/workbook.xml"] = wb.encode()


def conta_textos(parts):
    return sum(len(re.findall(rb'<c [^>]*t="s"', v)) for k, v in parts.items()
               if re.match(r"xl/worksheets/sheet\d+\.xml$", k))


# ------------------------------------------------------------------ conferência
def confere(parts_antes, parts, permitidas):
    # 1) só as partes esperadas mudaram; gráficos, desenhos, vínculos e a consulta ficam byte a byte
    mudou = {k for k in parts if parts_antes.get(k) != parts[k]}
    sumiu = set(parts_antes) - set(parts)
    assert mudou <= permitidas, mudou - permitidas
    assert sumiu == {"xl/calcChain.xml"}, sumiu
    for k in mudou:
        if k.endswith(".xml") or k.endswith(".rels"):
            ElementTree.fromstring(parts[k])                    # XML bem formado
    # 2) textos compartilhados: count exato
    sst = parts["xl/sharedStrings.xml"].decode()
    m = re.search(r'<sst [^>]*count="(\d+)" uniqueCount="(\d+)"', sst)
    assert int(m.group(1)) == conta_textos(parts), (m.group(1), conta_textos(parts))
    assert int(m.group(2)) == len(re.findall(r"<si>", sst))
    return sorted(mudou)


def main():
    z = zipfile.ZipFile(BASE)
    parts_antes = {n: z.read(n) for n in z.namelist()}
    parts = dict(parts_antes)
    abas, tabelas = mapa_partes(z)
    textos = Textos(parts["xl/sharedStrings.xml"].decode())
    aj = ajustes()

    log, n_desc = edita(parts, abas, tabelas, textos, aj)
    novas_cols, codigos, valores = colunas_ssos(parts, abas, tabelas, textos, aj)
    abas_ocultas(parts, aj)
    din = dinamicas(parts)
    sem_calcchain(parts)
    parts["xl/sharedStrings.xml"] = textos.grava(conta_textos(parts)).encode()
    rel_q = parts[tabelas["BASE_SS_OS"].replace("tables/", "tables/_rels/") + ".rels"].decode()
    permitidas = {abas[a] for a in ("Gestão", "Estoque", "SLA de Manutenção", "BASE SS_OS")} | {
        "xl/sharedStrings.xml", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml",
        "docProps/app.xml", tabelas["BASE_SS_OS"], "xl/" + re.search(r'Target="\.\./([^"]+)"', rel_q).group(1)}
    permitidas |= {n for n, _ in din} | {n.replace("pivotCache/", "pivotCache/_rels/") + ".rels" for n, _ in din}
    permitidas |= {k for k in parts if k not in parts_antes}          # as duas abas novas
    mudou = confere(parts_antes, parts, permitidas)

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    with zipfile.ZipFile(SAIDA, "w", zipfile.ZIP_DEFLATED) as out:
        for n in z.namelist():                            # mesma ordem de partes do original
            if n in parts:
                out.writestr(z.getinfo(n), parts[n])
        for n in parts:
            if n not in z.namelist():
                out.writestr(n, parts[n])

    confere_saida(log, novas_cols, codigos, valores, aj)
    print(f"{len(log)} células na Gestão/Estoque/SLA · {n_desc} fórmulas da Descrição com SEERRO · "
          f"{len(codigos)} linhas na BASE SS_OS × {len(NOVAS)} colunas ({novas_cols[0]}:{novas_cols[-1]}) · "
          f"abas ocultas {[n for n, _ in ABAS_AJUSTE]} · dinâmicas {din}")
    print("partes alteradas:", mudou)
    return log


def confere_saida(log, novas_cols, codigos, valores, aj):
    """Abre a saída com openpyxl e confere célula a célula contra a COEP 5 que ele mandou."""
    antes, depois = todas_as_celulas(BASE), todas_as_celulas(SAIDA)
    tocadas = {(aba, ref) for aba, ref, *_ in log}
    tocadas |= {("BASE SS_OS", f"{c}{n}") for c in novas_cols for n in range(1, len(codigos) + 2)}
    desc = [k for k in depois if k[0] == "Gestão" and k[1].startswith("G")]
    tocadas |= set(desc)
    novas_abas = {n for n, _ in ABAS_AJUSTE}
    for k in set(antes) | set(depois):
        if k in tocadas or k[0] in novas_abas:
            continue
        assert antes.get(k) == depois.get(k), (k, antes.get(k), depois.get(k))
    wb = openpyxl.load_workbook(SAIDA)
    assert [wb[n].sheet_state for n, _ in ABAS_AJUSTE] == ["hidden", "hidden"]
    ws = wb["BASE SS_OS"]
    t = ws.tables["BASE_SS_OS"]
    assert t.ref.endswith(f"{novas_cols[-1]}{len(codigos) + 1}"), t.ref
    assert [ws[f"{c}1"].value for c in novas_cols] == [n for n, _ in NOVAS]
    v = openpyxl.load_workbook(SAIDA, data_only=True)
    vs = v["BASE SS_OS"]
    for i, lin in enumerate(range(2, len(codigos) + 2)):
        assert tuple(vs[f"{c}{lin}"].value or "" for c in novas_cols) == valores[i], lin
    g = v["Gestão"]
    for aba, ref, a, campo, val in log:
        if aba == "Gestão" and campo != "Descrição":
            assert g[ref].value == val, (ref, g[ref].value, val)
        if aba == "Estoque":
            assert v["Estoque"][ref].value == val, (ref, v["Estoque"][ref].value, val)
    for n, ncol in ABAS_AJUSTE:
        ws = v[n]
        assert ws.max_row == len(aj[n]), (n, ws.max_row, len(aj[n]))
        assert [c.value for c in ws[1]][:ncol] == aj[n][0][:ncol]


if __name__ == "__main__":
    main()

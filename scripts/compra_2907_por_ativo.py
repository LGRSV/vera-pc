"""
Preenche a compra de 29/07/2026 na planilha COEP 4: quem recebe cada peça (aba Estoque,
Tabela7) e o PMA de cada ativo (aba Gestão, coluna M).

Pedido do gestor (24/09): «relacionar os equipamentos fora de Alta e Muito Alta com os PMAs
que a gente fez. A compra foi feita para Muito Alta e Alta, mas resolvemos alguns casos antes,
então temos que aproveitar o material que pedimos». Depois: «desconsidere a base de SS, o
intuito é só preencher a planilha — a aba Estoque e a Gestão com os PMAs certinho».

**Então tudo sai da própria planilha**: a compra por unidade (Tabela7), a Criticidade, o
Status, o Defeito e o Índice da aba Gestão, e o que ele já tinha digitado na coluna Ativo
(foto de 24/09 — as 11 linhas de 38295 e as 6 de 38297).

**A régua:**

1. O que ele já digitou fica.
2. Muito Alta e Alta primeiro; depois a ordem da coluna Índice (a prioridade dele).
3. Só recebe quem a compra conserta por inteiro: RL Completo = tanque + controle; RT Completo
   = três células + controle. Faltou uma peça, a vez passa para o próximo.
4. Não recebe quem a própria Gestão diz que já tem peça andando («Em logistica (N1>N3)») ou
   que já foi feito («Realizado»).
5. O que sobra fica como «Reserva».

As duas escolhas de Média que ele fez (7927713200 e 7937102148, Índice 61 e 62) seguem essa
ordem. As sobras de 34,5 vão para o 7900535058 (Índice 60, RL Completo: leva o último
controle) e o 7925733015 (67). O 7944559149 (63) e o 7908249152 (65) ficam sem: precisam de
controle de 34,5, e os sete já estão com dono.

**Edição cirúrgica**, como em 21/09: a COEP 4 tem 37 gráficos, 8 tabelas dinâmicas e vínculos
externos com valor só em cache — regravar com openpyxl apaga o cache. O script mexe só no XML
das células: aba Estoque (sheet5), coluna M da Gestão (sheet3), textos novos em sharedStrings e
um estilo de data em styles. E **tira o calcChain.xml**: o arquivo de 21/09 ainda listava a
Gestão!K44 como fórmula depois que ela virou valor, e isso faz o Excel pedir para reparar o
arquivo. Sem o calcChain, o Excel monta outro ao abrir.

Base: `data/raw/GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4_ENTREGUE_2109.xlsx` (o que foi entregue em
21/09). Saída: `dist/GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4.xlsx`.

Rodar: python3 scripts/compra_2907_por_ativo.py
"""

import datetime as dt
import os
import re
import zipfile
from collections import Counter, defaultdict
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.path.join(RAIZ, "data", "raw", "GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4_ENTREGUE_2109.xlsx")
SAIDA = os.path.join(RAIZ, "dist", "GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4.xlsx")

ABA_GESTAO, ABA_ESTOQUE = "xl/worksheets/sheet3.xml", "xl/worksheets/sheet5.xml"
ORDEM_PMA = ["38291", "38292", "38294", "38295", "38296", "38297"]
CHEGADA = {"38295": dt.date(2026, 10, 5), "38297": dt.date(2026, 10, 5)}   # informado pelo gestor
CHEGADA_RESTO = dt.date(2026, 11, 15)
STATUS_UNIDADE = "Em logística de entrega"                  # o texto que ele já usa na coluna
RESERVA = "Reserva"

# Peças de cada defeito. Provada no centavo pela coluna Orçamento MAT da própria aba Gestão.
BOM = {"RL Completo 13,8": ["38294", "38296"], "RL Completo 34,5": ["38295", "38297"],
       "Tanque 13,8": ["38294"], "Tanque 34,5": ["38295"], "Controle 34,5": ["38297"],
       "RT Completo 400 34,5": ["38292", "38292", "38292", "38291"],
       "Célula 400 34,5": ["38292"], "Controle 200 34,5": ["38291"],
       "Controle 200 13,8": ["38291"], "Controle 167 13,8": ["38291"]}

# O que ele já tinha digitado na coluna Ativo (foto de 24/09), na ordem das linhas.
DIGITADO = {
    "38295": ["7947203070", "7927413001", "7900453110", "7903112004", "7920024127",
              "7925744087", "7925871077", "7927646026", "7930428001", "7927713200",
              "7937102148"],
    "38297": ["7926089013", "7947203070", "7927413001", "7900453110", "7903112004",
              "7920024127"],
}
# Status da Gestão que querem dizer «não precisa desta compra».
JA_TEM = {"Em logistica (N1>N3)": "a Gestão diz Em logística (N1>N3): a peça já está andando",
          "Realizado": "a Gestão diz Realizado"}
MA_ALTA = {"Muito Alta", "Alta"}


def _t(v):
    return "" if v is None else str(v).strip()


def excel_data(d):
    return (d - dt.date(1899, 12, 30)).days


# ------------------------------------------------------------------ leitura da planilha
def ler(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb["Gestão"]
    col = {_t(c.value): i for i, c in enumerate(ws[1], 1)}
    g = {}
    # só a tabela (Table1, A1:AG54): embaixo dela, nas linhas 59-99, está a conta de kits
    # do gestor e uma cópia do plano de 16/07, com os mesmos códigos de ativo
    fim = int(re.sub(r"\D", "", ws.tables["Table1"].ref.split(":")[1]))
    for r in range(2, fim + 1):
        a = _t(ws.cell(row=r, column=col["Ativo"]).value)
        if a:
            v = lambda n: ws.cell(row=r, column=col[n]).value
            g[a] = dict(linha=r, ativo=a, crit=_t(v("Criticidade")), status=_t(v("Status")),
                        defeito=_t(v("Defeito")), indice=v("Índice"), pma=_t(v("PMA")))
    we = wb["Estoque"]
    ini, fim = we.tables["Tabela7"].ref.split(":")
    r0, r1 = int(re.sub(r"\D", "", ini)), int(re.sub(r"\D", "", fim))
    unidades = [dict(linha=r, pma=_t(we.cell(row=r, column=1).value),
                     data=we.cell(row=r, column=9).value, req=we.cell(row=r, column=10).value)
                for r in range(r0 + 1, r1 + 1)]
    return g, unidades


# ------------------------------------------------------------------ a divisão
def divide(g, qtd):
    lista = {p: [None] * qtd[p] for p in ORDEM_PMA}
    recebe = defaultdict(Counter)
    for p, ativos in DIGITADO.items():
        for i, a in enumerate(ativos):
            assert a in g, f"{a}, digitado no {p}, não está na Gestão"
            assert p in BOM.get(g[a]["defeito"], []), f"{a} ({g[a]['defeito']}) não usa o {p}"
            assert g[a]["status"] not in JA_TEM, f"{a} digitado, mas {JA_TEM[g[a]['status']]}"
            lista[p][i] = a
            recebe[a][p] += 1
    fila = sorted(g.values(), key=lambda x: (0 if x["crit"] in MA_ALTA else 1, x["indice"]))
    sem = {}
    for x in fila:
        a = x["ativo"]
        if x["status"] in JA_TEM or not BOM.get(x["defeito"]):
            continue
        falta = Counter(BOM[x["defeito"]]) - recebe[a]
        if not falta:
            continue
        if all(lista[p].count(None) >= q for p, q in falta.items()):
            for p, q in falta.items():
                for _ in range(q):
                    lista[p][lista[p].index(None)] = a
                    recebe[a][p] += 1
        else:
            sem[a] = sorted(p for p, q in falta.items() if lista[p].count(None) < q)
    for p in ORDEM_PMA:
        lista[p] = [a if a else RESERVA for a in lista[p]]
    recebe = {a: c for a, c in recebe.items() if c}     # o defaultdict cria vazios ao consultar
    return lista, recebe, sem


def pma_da_gestao(g, recebe):
    """O que vai na coluna M: os PMAs de quem recebe; a anotação dele onde já havia uma
    (ex.: «Controle Rua que voltou de dianopólis»); «Sem PMA» no resto."""
    novo = {}
    for a, x in g.items():
        if a in recebe:
            novo[a] = "+".join(sorted(recebe[a]))
        elif x["pma"] and x["pma"] != "Sem PMA" and not re.fullmatch(r"[\d+]+", x["pma"]):
            novo[a] = x["pma"]
        else:
            novo[a] = "Sem PMA"
    return novo


# ------------------------------------------------------------------ XML
class Textos:
    """sharedStrings: reaproveita o índice do texto que já existe, acrescenta o que falta."""

    def __init__(self, xml):
        self.xml = xml
        self.si = re.findall(r"<si>.*?</si>", xml, re.S)
        self.idx = {}
        for i, s in enumerate(self.si):
            t = "".join(re.findall(r"<t[^>]*>(.*?)</t>", s, re.S))
            self.idx.setdefault(t, i)
        self.novos = []

    def de(self, texto):
        chave = escape(texto)
        if chave not in self.idx:
            self.idx[chave] = len(self.si) + len(self.novos)
            self.novos.append(f'<si><t xml:space="preserve">{chave}</t></si>')
        return self.idx[chave]

    def texto(self, i):
        s = self.si[i] if i < len(self.si) else self.novos[i - len(self.si)]
        return "".join(re.findall(r"<t[^>]*>(.*?)</t>", s, re.S))

    def grava(self, referencias):
        """`count` é o total de células que apontam para a lista. O Excel grava exato; a edição
        de 21/09 deixou 32 a menos. Aqui vai o número contado nas abas depois da edição."""
        m = re.search(r'<sst [^>]*count="(\d+)" uniqueCount="(\d+)"', self.xml)
        count, uniq = referencias, len(self.si) + len(self.novos)
        xml = self.xml.replace(m.group(0), m.group(0).replace(f'count="{m.group(1)}"', f'count="{count}"')
                               .replace(f'uniqueCount="{m.group(2)}"', f'uniqueCount="{uniq}"'), 1)
        return xml.replace("</sst>", "".join(self.novos) + "</sst>")


def cel(ref, valor=None, s=None, textos=None, formula=None):
    est = f' s="{s}"' if s is not None else ""
    if formula is not None:
        v = "" if valor is None else (f"<v>{valor}</v>" if isinstance(valor, (int, float))
                                      else f"<v>{escape(str(valor))}</v>")
        t = "" if isinstance(valor, (int, float)) else ' t="str"'
        return f'<c r="{ref}"{est}{t}><f>{escape(formula)}</f>{v}</c>'
    if isinstance(valor, str):
        return f'<c r="{ref}"{est} t="s"><v>{textos.de(valor)}</v></c>'
    return f'<c r="{ref}"{est}><v>{valor}</v></c>'


def linha_xml(xml, n):
    m = re.search(rf'<row r="{n}"[^>]*?(?:/>|>.*?</row>)', xml, re.S)
    return m.group(0) if m else None


def edita_estoque(xml, textos, unidades, lista, estilo_data):
    # 1) Tabela7: Ativo, Status e Previsão em cada unidade
    pos = defaultdict(int)
    for u in unidades:
        p, n = u["pma"], u["linha"]
        ativo = lista[p][pos[p]]
        pos[p] += 1
        velha = linha_xml(xml, n)
        nova = re.sub(rf'<c r="[KLM]{n}"[^>]*?(?:/>|>.*?</c>)', "", velha, flags=re.S)
        nova = re.sub(r'spans="\d+:\d+"', 'spans="1:13"', nova)
        nova = nova.replace("</row>", cel(f"K{n}", ativo, textos=textos)
                            + cel(f"L{n}", STATUS_UNIDADE, textos=textos)
                            + cel(f"M{n}", excel_data(CHEGADA.get(p, CHEGADA_RESTO)), s=12) + "</row>")
        xml = xml.replace(velha, nova, 1)

    # 2) o quadro «Status PMA» que ele começou em A6:D7, completo e com o saldo
    cab = linha_xml(xml, 6)
    extra = "".join(cel(f"{c}6", t, s=142, textos=textos) for c, t in
                    (("E", "Previsão de Chegada no N1"), ("F", "Comprado"),
                     ("G", "Com equipamento"), ("H", "Saldo")))
    xml = xml.replace(cab, re.sub(r'spans="\d+:\d+"', 'spans="1:8"', cab).replace("</row>", extra + "</row>"), 1)
    por_pma = {p: [u for u in unidades if u["pma"] == p] for p in ORDEM_PMA}
    novas = []
    for i, p in enumerate(ORDEM_PMA):
        n = 7 + i
        u0 = por_pma[p][0]
        compr = len(por_pma[p])
        com = sum(1 for a in lista[p] if a != RESERVA)
        rng_p, rng_k = "$A$36:$A$75", "$K$36:$K$75"
        celulas = (cel(f"A{n}", int(p), s=141)
                   + cel(f"B{n}", excel_data(u0["data"].date() if hasattr(u0["data"], "date") else u0["data"]), s=estilo_data)
                   + cel(f"C{n}", int(u0["req"]), s=141)
                   + cel(f"D{n}", STATUS_UNIDADE, s=141, textos=textos)
                   + cel(f"E{n}", excel_data(CHEGADA.get(p, CHEGADA_RESTO)), s=estilo_data)
                   + cel(f"F{n}", compr, s=141, formula=f"COUNTIF({rng_p},A{n})")
                   + cel(f"G{n}", com, s=141,
                         formula=f'COUNTIFS({rng_p},A{n},{rng_k},"<>",{rng_k},"<>{RESERVA}")')
                   + cel(f"H{n}", compr - com, s=141, formula=f"F{n}-G{n}"))
        novas.append(f'<row r="{n}" spans="1:8" x14ac:dyDescent="0.25">{celulas}</row>')
    velha7 = linha_xml(xml, 7)
    xml = xml.replace(velha7, "".join(novas), 1)
    for n in range(8, 13):
        assert linha_xml(xml.replace("".join(novas), ""), n) is None, f"linha {n} da aba Estoque não estava vazia"

    # 3) ao lado do «Saldo» que ele deixou em A1
    saldo = sum(1 for p in ORDEM_PMA for a in lista[p] if a == RESERVA)
    l1 = linha_xml(xml, 1)
    xml = xml.replace(l1, re.sub(r'spans="\d+:\d+"', 'spans="1:3"', l1).replace(
        "</row>", cel("B1", saldo, formula="SUM(H7:H12)")
        + cel("C1", "peça(s) da compra ainda sem equipamento (Reserva na coluna Ativo)", textos=textos)
        + "</row>"), 1)
    return xml, saldo


def edita_gestao(xml, textos, g, novo):
    mudou = []
    for a, x in g.items():
        ref = f"M{x['linha']}"
        m = re.search(rf'<c r="{ref}"([^>]*?)(?:/>|>(.*?)</c>)', xml, re.S)
        attrs = re.sub(r'\s*t="[^"]*"', "", m.group(1))
        antes = x["pma"]
        if antes == novo[a]:
            continue
        xml = xml.replace(m.group(0), f'<c r="{ref}"{attrs} t="s"><v>{textos.de(novo[a])}</v></c>', 1)
        mudou.append((x["linha"], a, antes, novo[a]))
    return xml, mudou


def estilo_de_data(styles):
    """O quadro dele usa o estilo 141 (preenchimento 12, borda 19). Para data, o mesmo com
    formato de data: reaproveita se existir, senão acrescenta ao fim de cellXfs."""
    bloco = re.search(r"<cellXfs count=\"(\d+)\">(.*?)</cellXfs>", styles, re.S)
    xfs = re.findall(r"<xf [^>]*?(?:/>|>.*?</xf>)", bloco.group(2), re.S)
    for i, xf in enumerate(xfs):
        if all(k in xf for k in ('numFmtId="14"', 'fillId="12"', 'borderId="19"')):
            return styles, i
    novo = ('<xf numFmtId="14" fontId="0" fillId="12" borderId="19" xfId="0" '
            'applyNumberFormat="1" applyFill="1" applyBorder="1"/>')
    n = int(bloco.group(1))
    styles = styles.replace(bloco.group(0), f'<cellXfs count="{n + 1}">' + bloco.group(2) + novo + "</cellXfs>", 1)
    return styles, n


# ------------------------------------------------------------------ conferência
def todas_as_celulas(caminho):
    wb = openpyxl.load_workbook(caminho)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula):
                    v = ("ArrayFormula", v.ref, v.text)
                if v is not None:
                    out[(ws.title, c.coordinate)] = v
    return out


def main():
    g, unidades = ler(BASE)
    qtd = Counter(u["pma"] for u in unidades)
    assert dict(qtd) == {"38291": 3, "38292": 4, "38294": 8, "38295": 13, "38296": 5, "38297": 7}
    lista, recebe, sem = divide(g, qtd)
    for a, c in recebe.items():           # cada um recebe exatamente o que o defeito pede
        assert c == Counter(BOM[g[a]["defeito"]]), (a, c)
    for p, ativos in DIGITADO.items():    # nada do que ele digitou saiu do lugar
        assert lista[p][:len(ativos)] == ativos, p
    for a, x in g.items():                # Muito Alta e Alta com material nesta compra: todos
        if x["crit"] in MA_ALTA and BOM.get(x["defeito"]) and x["status"] not in JA_TEM:
            assert a in recebe, a
    novo = pma_da_gestao(g, recebe)

    with zipfile.ZipFile(BASE) as z:
        infos = z.infolist()
        partes = {i.filename: z.read(i.filename) for i in infos}
    textos = Textos(partes["xl/sharedStrings.xml"].decode("utf-8"))
    styles, estilo_data = estilo_de_data(partes["xl/styles.xml"].decode("utf-8"))
    estoque, saldo = edita_estoque(partes[ABA_ESTOQUE].decode("utf-8"), textos, unidades, lista, estilo_data)
    gestao, mudou = edita_gestao(partes[ABA_GESTAO].decode("utf-8"), textos, g, novo)
    partes[ABA_ESTOQUE] = estoque.encode("utf-8")
    partes[ABA_GESTAO] = gestao.encode("utf-8")
    refs = sum(len(re.findall(r'<c [^>]*t="s"', b.decode("utf-8"))) for n, b in partes.items()
               if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n))
    partes["xl/sharedStrings.xml"] = textos.grava(refs).encode("utf-8")
    partes["xl/styles.xml"] = styles.encode("utf-8")
    # sem calcChain: o de 21/09 ainda apontava a Gestão!K44 como fórmula
    ct = partes["[Content_Types].xml"].decode("utf-8")
    partes["[Content_Types].xml"] = re.sub(r'<Override PartName="/xl/calcChain.xml"[^>]*/>', "", ct).encode("utf-8")
    rels = partes["xl/_rels/workbook.xml.rels"].decode("utf-8")
    partes["xl/_rels/workbook.xml.rels"] = re.sub(r'<Relationship [^>]*Target="calcChain.xml"/>', "", rels).encode("utf-8")
    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    with zipfile.ZipFile(SAIDA, "w", zipfile.ZIP_DEFLATED) as z:
        for i in infos:
            if i.filename == "xl/calcChain.xml":
                continue
            z.writestr(i, partes[i.filename], compress_type=zipfile.ZIP_DEFLATED)

    # ---- conferência do arquivo gravado ----
    with zipfile.ZipFile(BASE) as a, zipfile.ZipFile(SAIDA) as b:
        mexidas = {n for n in a.namelist() if n in b.namelist() and a.read(n) != b.read(n)}
        sumiu = set(a.namelist()) - set(b.namelist())
        assert set(b.namelist()) <= set(a.namelist())
    esperadas = {ABA_ESTOQUE, ABA_GESTAO, "xl/sharedStrings.xml", "xl/styles.xml",
                 "[Content_Types].xml", "xl/_rels/workbook.xml.rels"}
    assert mexidas <= esperadas and sumiu == {"xl/calcChain.xml"}, (mexidas, sumiu)
    antes, depois = todas_as_celulas(BASE), todas_as_celulas(SAIDA)
    dif = {k for k in set(antes) | set(depois) if antes.get(k) != depois.get(k)}
    permitido = ({("Estoque", f"{c}{n}") for c in "KLM" for n in range(36, 76)}
                 | {("Estoque", f"{c}{n}") for c in "ABCDEFGH" for n in range(6, 13)}
                 | {("Estoque", "B1"), ("Estoque", "C1")}
                 | {("Gestão", f"M{linha}") for linha, *_ in mudou})
    assert dif <= permitido, sorted(dif - permitido)
    wv = openpyxl.load_workbook(SAIDA, data_only=True)
    e, gg = wv["Estoque"], wv["Gestão"]
    pos = defaultdict(int)
    for u in unidades:
        p = u["pma"]
        assert _t(e.cell(row=u["linha"], column=11).value) == lista[p][pos[p]]
        assert e.cell(row=u["linha"], column=13).value.date() == CHEGADA.get(p, CHEGADA_RESTO)
        pos[p] += 1
    for i, p in enumerate(ORDEM_PMA):
        n = 7 + i
        assert e.cell(row=n, column=6).value == qtd[p]
        assert e.cell(row=n, column=6).value == e.cell(row=n, column=7).value + e.cell(row=n, column=8).value
    assert e["B1"].value == saldo
    for a, x in g.items():
        assert _t(gg.cell(row=x["linha"], column=13).value) == novo[a], a

    # ---- relatório ----
    print(f"gravado {SAIDA}")
    print(f"partes alteradas: {sorted(mexidas)} · removida: xl/calcChain.xml")
    print(f"células alteradas: {len(dif)} (todas dentro do previsto)")
    nome = {"38291": "Controle RT", "38292": "Célula RT 400", "38294": "Tanque 13,8",
            "38295": "Tanque 34,5", "38296": "Controle 13,8", "38297": "Controle 34,5"}
    for p in ORDEM_PMA:
        print(f"\n{p} {nome[p]} ({len(lista[p])} un., chega {CHEGADA.get(p, CHEGADA_RESTO):%d/%m})")
        for k, a in enumerate(lista[p], 1):
            x = g.get(a, {})
            ja = "você" if a in DIGITADO.get(p, []) and k <= len(DIGITADO[p]) else "novo"
            if a == RESERVA:
                ja = ""
            print(f"   {k:2d} {a:<11} {x.get('crit', ''):<13} {x.get('indice', '') or '':>4} "
                  f"{x.get('defeito', ''):<22} {ja}")
    print("\nGestão, coluna M — o que mudou em relação a 21/09:")
    for linha, a, antes, depois in mudou:
        print(f"   M{linha:<3} {a} {g[a]['crit']:<13} {antes:<12} -> {depois}")
    print("\nFicam sem material desta compra (precisavam de peça que acabou):")
    for a, ps in sem.items():
        print(f"   {a} {g[a]['crit']:<13} Índice {g[a]['indice']:>3} {g[a]['defeito']:<22} "
              f"faltou {'/'.join(ps)} · status {g[a]['status']}")
    print(f"\nReserva (sem equipamento na Gestão que use): {saldo}")


if __name__ == "__main__":
    main()

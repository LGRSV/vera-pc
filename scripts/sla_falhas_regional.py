"""
SLA das falhas da taxa por regional — 2025 e 2026, da falha até a peça trocada.

Pedido do gestor (28/09): «calcular com base na taxa de falha o SLA do ano passado e o SLA
desse ano por regional». Entre três leituras, ele escolheu **«da falha até a troca da peça»**:
para cada falha da taxa, o tempo total, da falha até a peça trocada em campo, contra o prazo
da proposta DCMD.

UNIVERSO  as 90 linhas da aba «Falha Equipamentos» da COEP 5 — o rol que faz a taxa de falha
          (2025: 35 RL + 18 RT; 2026: 28 RL + 9 RT). O ano é o da fatia («RL 2025»), não o da
          coluna Ano (numa das 90 as duas discordam — lição da planilha base).
INÍCIO    a DATA DA FALHA do rol (a ocorrência — a mesma data que põe a falha no ano da taxa).
          A primeira ideia era a abertura da primeira SS da demanda (a cadeia de repasses que
          contém a SS da falha, para trás e para a frente, em todos os ramos), mas ela não data a
          falha: em 12 das 90 a primeira SS da base vem mais de 90 dias depois — até 431 —, porque o
          COEP recriou SS em lote (as de 29/06/2026 com número de 2025). Contar dali daria «10 dias,
          no prazo» a um religador parado desde dezembro. A conta pela abertura da SS vai ao lado,
          como segunda leitura.
TROCA     o texto manda: a primeira frase de execução («foi substituído», «substituído conforme
          solicitado», «instalado com sucesso», «feita a substituição»…) nas SS do ativo dentro da
          janela da falha (da data da falha até a próxima falha do mesmo ativo), com a data
          escrita junto da frase ou, sem data, a saída da SS onde a frase aparece pela primeira vez
          — frase que já estava em SS encerrada antes da falha é troca antiga repetida e não vale;
          frase em pergunta («conferir se a placa foi substituída») ou só de acessório («foi
          instalado o rádio») também não, nem a que o formulário da DMSL fecha com «ficou em
          operação? não» (troca parcial, a falha seguiu); data sem ano pega o ano da SS. SS aberta
          depois do serviço (conclusão antes da abertura — regularização) conta 0 dia. Quando o
          texto não diz, vale a primeira execução em campo pela cadeia, nesta ordem:
            1. o COCM sai para PROT, TELE, SE ou DMSL, ou fecha a SS atendida, depois de passar
               pelo COEP — é a esteira do gestor: saiu do campo para ajuste ou comissionamento,
               a peça foi trocada;
            2. o COEP repassa direto para a PROT ou a SE («instalado com sucesso, segue para
               Proteção») — troca feita fora da cadeia;
            3. a TELE, a SE ou a DMSL fecha a SS atendida — a própria DMSL trocou (comum em controle);
            4. o COEP fecha a SS atendida.
          Texto e regras 1–2 juntos: vale a data mais cedo (o COCM às vezes só repassa semanas
          depois de trocar). O rol diz se a troca foi feita («Troca feita», leitura dos textos até
          19/08). Quando ele diz «sim», a data sai do texto ou das quatro regras. Quando diz «não»,
          só valem texto posterior a 19/08 e as regras 1 e 2 — troca pela esteira, com a divergência
          anotada —; as regras 3 e 4 não bastam para desmentir quem leu o texto. Ativo que segue
          pendente na Gestão não tem troca.
PRAZO     proposta DCMD, pela criticidade do equipamento: Muito Alta 11 · Alta 20 · Média 40 ·
          Baixa 60 · sem classificação 32,75 (a média das quatro, régua do quadro dele). A
          criticidade vem da Gestão da COEP 5 e, fora dela, da aba de mapeamento da carteira.
SITUAÇÃO  trocada no prazo · trocada fora do prazo · em aberto (prazo estourado ou dentro) ·
          encerrada sem troca · trocada sem data na base.
SLA       trocadas no prazo ÷ trocadas — a mesma conta do «% de substituição no prazo» do quadro
          dele. Ao lado, a leitura que conta também quem está em aberto com o prazo estourado.
REGIONAL  da carteira (ATUALIZADA 16), ativo a ativo; quem não está lá vai pela localidade do
          cadastro de ajustes.
TEMPO     quanto levou da falha à troca, por ano da falha (aba «Tempo até consertar», pergunta de
          28/09: «quanto tempo demorávamos pra consertar em 2025 e quanto tempo demoramos em
          2026»). A mediana das trocadas sozinha engana: das falhas de 2026 só as rápidas já
          terminaram, e as abertas ainda correm. Vão junto a mediana mínima (a que sairia se todas
          as abertas fossem trocadas hoje) e a conta na mesma janela — trocadas em até 30, 60, 90,
          180 e 365 dias, entrando só a falha que já teve esse tempo todo para ser trocada.

Base de SS: RELIGA_REGULA de 23/09 (posição do «hoje» para quem está em aberto).
Grava dist/SLA_FALHAS_REGIONAL.xlsx e data/missao/sla_falhas_regional.json.
Rodar: python3 scripts/sla_falhas_regional.py
"""

import datetime as dt
import json
import os
import re
import sys
from collections import Counter, defaultdict
from statistics import median

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sla_manutencao as sm  # noqa: E402 — norm() e o mapeamento de criticidade da carteira

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(RAIZ, "data", "raw")
COEP5 = os.path.join(RAW, "GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_5.xlsx")
BASE_SS = os.path.join(RAW, "RELIGA_REGULA_23092026.xlsx")
CARTEIRA = os.path.join(RAW, "EQUIPAMENTOS_INDISPONIVEIS_ATUALIZADA16.xlsx")
AJUSTES = os.path.join(RAW, "GESTAO_DE_EQUIPAMENTOS.xlsx")
SAIDA = os.path.join(RAIZ, "dist", "SLA_FALHAS_REGIONAL.xlsx")
JSON = os.path.join(RAIZ, "data", "missao", "sla_falhas_regional.json")

HOJE = dt.datetime(2026, 9, 23, 23, 59)          # posição da base de SS
CORTE_DO_ROL = dt.datetime(2026, 8, 19)          # o rol leu os textos até aqui
PRAZO = {"Muito Alta": 11, "Alta": 20, "Média": 40, "Baixa": 60}
PRAZO_SEM = 32.75                                # a média das quatro, régua do quadro dele
REGIONAIS = ("NORTE", "CENTRO", "SUL")
ANOS = (2025, 2026)
RD = re.compile(r"-RD-")
FORA_DCMD = re.compile(r"PROT|TELE|SE-|DMSL|DEOP")


def txt(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


# ------------------------------------------------------------------ leitura
def rol():
    ws = openpyxl.load_workbook(COEP5, data_only=True)["Falha Equipamentos"]
    cab = [c.value for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(cab, r))
        if not d.get("Ativo"):
            continue
        tipo, ano = txt(d["Ano/TIPO"]).split()
        out.append({"ativo": txt(d["Ativo"]), "tipo": tipo, "ano": int(ano), "ss": sm.norm(txt(d["SS"])),
                    "ocorrencia": txt(d["Data"]), "peca": txt(d["Peça (modo)"]),
                    "troca_rol": txt(d["Troca feita"]).lower()})
    assert Counter((x["tipo"], x["ano"]) for x in out) == {("RL", 2025): 35, ("RT", 2025): 18,
                                                           ("RL", 2026): 28, ("RT", 2026): 9}
    return out


def base():
    """{SS: [linhas]} (a mesma SS vem repetida quando o repasse bifurca) e as ligações para trás.
    A planilha grava a dimensão como A1:A1: read_only=True leria uma célula só."""
    ws = openpyxl.load_workbook(BASE_SS, data_only=True)["Exportar Planilha"]
    linhas = list(ws.iter_rows(values_only=True))
    cab = linhas[0]
    por, antes, por_ativo = defaultdict(list), defaultdict(set), defaultdict(list)
    for r in linhas[1:]:
        d = dict(zip(cab, r))
        k = sm.norm(txt(d["SS_ORIGINAL"]))
        if not k:
            continue
        d["_seg"] = sm.norm(txt(d["SS_APOS_REPASSE"])) if d["SS_APOS_REPASSE"] else ""
        por[k].append(d)
        por_ativo[txt(d["EQUIPAMENTO"])].append(d)
        if d["_seg"]:
            antes[d["_seg"]].add(k)
    return por, antes, por_ativo


def gestao():
    """Criticidade e status de cada ativo na Gestão da COEP 5 (só dentro da Table1)."""
    wf = openpyxl.load_workbook(COEP5)
    fim = int(re.sub(r"\D", "", wf["Gestão"].tables["Table1"].ref.split(":")[1]))
    ws = openpyxl.load_workbook(COEP5, data_only=True)["Gestão"]
    cab = [c.value for c in ws[1]]
    out = {}
    for r in ws.iter_rows(min_row=2, max_row=fim, values_only=True):
        d = dict(zip(cab, r))
        out[txt(d["Ativo"])] = {"criticidade": txt(d["Criticidade"]), "status": txt(d["Status"])}
    return out


def regionais():
    """{ativo: (município, polo, regional)} da carteira; {localidade: (polo, regional)} para o resto."""
    wb = openpyxl.load_workbook(CARTEIRA, read_only=True, data_only=True)
    por_ativo, por_local = {}, {}
    for ws in wb.worksheets:
        linhas = ws.iter_rows(values_only=True)
        try:
            cab = [txt(x).upper() for x in next(linhas)]
        except StopIteration:
            continue
        if "REGIONAL" not in cab or "LOCALIDADE" not in cab:
            continue
        il, ir = cab.index("LOCALIDADE"), cab.index("REGIONAL")
        ip = cab.index("POLO") if "POLO" in cab else None
        for r in linhas:
            if not r or len(r) <= ir or not r[ir]:
                continue
            reg = txt(r[ir]).upper()
            if reg not in REGIONAIS:
                continue
            loc, polo = txt(r[il]).upper(), (txt(r[ip]).upper() if ip is not None else "")
            a = txt(r[0])
            if a[:2] in ("79", "78", "58"):
                por_ativo.setdefault(a, (loc, polo, reg))
            if loc:
                por_local.setdefault(loc, (polo, reg))
    wa = openpyxl.load_workbook(AJUSTES, read_only=True, data_only=True)
    local_cad = {}
    for aba, il in (("Ajustes RL Poste", 13), ("Ajustes Reguladores de Tensão", 14)):
        for r in list(wa[aba].iter_rows(values_only=True))[1:]:
            if r[0] is not None and len(r) > il and r[il]:
                local_cad[txt(r[0])] = txt(r[il]).upper()
    return por_ativo, por_local, local_cad


# ------------------------------------------------------------------ a demanda de cada falha
def demanda(ss, por, antes):
    """Todas as SS ligadas por repasse à SS da falha — para trás e para a frente, todos os ramos."""
    vis, fila = set(), [ss]
    while fila:
        x = fila.pop()
        if x in vis or x not in por:
            continue
        vis.add(x)
        fila.extend(d["_seg"] for d in por[x] if d["_seg"])
        fila.extend(antes.get(x, ()))
    # a SS repassada pode vir com a mesma abertura da anterior, no segundo: no empate, quem
    # repassou vem antes de quem recebeu (sem isso a cabeça trocava de uma rodada para outra)
    fundo = {}

    def prof(x, caminho=()):
        if x not in fundo:
            ants = [a for a in antes.get(x, ()) if a in vis and a not in caminho]
            fundo[x] = 1 + max((prof(a, caminho + (x,)) for a in ants), default=-1)
        return fundo[x]
    return sorted(vis, key=lambda x: (por[x][0]["DTA_ABERTURA"], prof(x), x))


def saida(ss, por):
    """(data, como) — quando a SS saiu do posto e para onde."""
    d0 = por[ss][0]
    st = txt(d0["STATUS"]).upper()
    if st in ("SS ATENDIDA", "SS CANCELADA"):
        return d0["DTA_CONCLUSAO"], st.split()[-1], ""
    if st == "SS PENDENTE":
        return None, "PENDENTE", ""
    seg = sorted((por[d["_seg"]][0]["DTA_ABERTURA"], txt(por[d["_seg"]][0]["POSTO_SGM"]))
                 for d in por[ss] if d["_seg"] and d["_seg"] in por)
    if not seg:
        return d0["DTA_ABERTURA"], "REPASSADA sem seguinte", ""
    return seg[0][0], "REPASSADA", seg[0][1]


def troca_na_base(nos, por, inicio, limite):
    """A primeira execução em campo da demanda dentro da janela, pelas quatro regras."""
    def na_janela(d):
        return d is not None and inicio <= d <= limite

    def regularizada(x):
        """SS aberta depois do serviço (conclusão antes da abertura), aberta dentro da janela."""
        d0 = por[x][0]
        return (txt(d0["STATUS"]).upper() == "SS ATENDIDA" and d0["DTA_CONCLUSAO"] is not None
                and d0["DTA_CONCLUSAO"] < d0["DTA_ABERTURA"] and inicio <= d0["DTA_ABERTURA"] + dt.timedelta(days=1)
                and d0["DTA_ABERTURA"] <= limite)

    coep = [x for x in nos if "COEP" in txt(por[x][0]["POSTO_SGM"])]
    depois_do_coep = por[coep[0]][0]["DTA_ABERTURA"] if coep else None
    # 1) COCM depois do COEP sai para ajuste/comissionamento ou fecha atendida
    for x in nos:
        posto = txt(por[x][0]["POSTO_SGM"])
        if not RD.search(posto) or (depois_do_coep and por[x][0]["DTA_ABERTURA"] < depois_do_coep):
            continue
        data, como, dest = saida(x, por)
        if regularizada(x):
            return max(data, inicio), 1, f"{x} ({posto}) atendida — SS aberta depois do serviço (regularização)"
        if na_janela(data) and (como == "ATENDIDA" or (como == "REPASSADA" and FORA_DCMD.search(dest))):
            return data, 1, f"{x} ({posto}) {'atendida' if como == 'ATENDIDA' else 'repassou para ' + dest}"
    # 2) COEP repassa direto para PROT/SE
    for x in coep:
        data, como, dest = saida(x, por)
        if na_janela(data) and como == "REPASSADA" and re.search(r"PROT|SE-", dest):
            return data, 2, f"{x} repassou para {dest}"
    # 3) TELE/SE/DMSL fecha atendida
    for x in nos:
        posto = txt(por[x][0]["POSTO_SGM"])
        if FORA_DCMD.search(posto) and "PROT" not in posto:
            data, como, _ = saida(x, por)
            if regularizada(x):
                return max(data, inicio), 3, f"{x} ({posto}) atendida — SS aberta depois do serviço (regularização)"
            if na_janela(data) and como == "ATENDIDA":
                return data, 3, f"{x} ({posto}) atendida"
    # 4) COEP fecha atendida
    for x in coep:
        data, como, _ = saida(x, por)
        if na_janela(data) and como == "ATENDIDA":
            return data, 4, f"{x} atendida no COEP"
    return None, None, ""


EXEC = re.compile(
    r"(?<!N[ÃA]O )\b(?:(?:FOI|FORAM|J[ÁA] FOI)\s+(?:REALIZAD[OA]\s+A\s+)?(?:SUBSTITU[ÍI]D[OA]S?|"
    r"TROCAD[OA]S?|INSTALAD[OA]S?)|(?:SUBSTITU[ÍI]D[OA]S?|TROCAD[OA]S?)\s+(?:CONFORME|COM\s+SUCESSO|"
    r"NA\s+OBRA|EM\s+\d{1,2}[/.]\d{1,2})|INSTALAD[OA]\s+COM\s+SUCESSO|(?:REALIZAD[OA]|EFETUAD[OA]|"
    r"FEIT[OA])\s+A\s+(?:SUBSTITUI[CÇ][AÃ]O|TROCA)\b)")
# pergunta ou condição logo antes («conferir se realmente a placa foi substituída»)
DUVIDA = re.compile(r"\b(?:SE|CASO|QUANDO|AP[ÓO]S|SER|SEJA)\b[^.]{0,45}$")
# o que vem depois fala só de acessório («foi instalado o rádio»)
ACESSORIO = re.compile(r"R[ÁA]DIO|ANTENA|BATERIA|PARA[- ]?RAIO|CHAVE|CABO|ATERRAMENTO|FUS[ÍI]VEL|ELO|MODEM")
PECA_GRANDE = re.compile(r"TANQUE|PARTE ATIVA|CONTROLE|REL[ÉE]|C[ÉE]LULA|EQUIPAMENTO|RELIGADOR|REGULADOR|"
                         r"FASE|PLACA|CONJUNTO|RL|RT\b|BANCO")
DATA = re.compile(r"(?<!\d)(\d{1,2})[/.-](\d{1,2})(?:[/.-](\d{4}|\d{2}))?(?![\d/])")
# formulário da DMSL: troca que não pôs o equipamento para operar não resolveu a falha
NAO_OPERA = re.compile(r"EM OPERA[ÇC][ÃA]O\s*\??\s*:?\s*N[ÃA]O")


def _datas_perto(texto, ini, fim, perto_de, raio=80):
    """Datas escritas perto da frase. Sem ano, vale o ano que põe a data junto da SS (±60 dias)."""
    out = []
    for m in DATA.finditer(texto[max(0, ini - raio):fim + raio]):
        d, mes, a = int(m.group(1)), int(m.group(2)), m.group(3)
        anos = [int(a) + (2000 if len(a) == 2 else 0)] if a else [perto_de.year, perto_de.year - 1]
        for ano in anos:
            try:
                x = dt.datetime(ano, mes, d)
            except ValueError:
                continue
            if a or abs((x - perto_de).days) <= 60:
                out.append(x)
                break
    return out


def troca_no_texto(a, inicio, limite, por, por_ativo):
    """A primeira frase de execução nova dentro da janela [início, limite] — (data, evidência)."""
    ss_do_ativo = []
    for d in por_ativo.get(a, []):
        ss = sm.norm(txt(d["SS_ORIGINAL"]))
        ss_do_ativo.append((d, ss, saida(ss, por)[0]))
    velhas = set()                    # frases que já estavam em SS encerradas antes da falha
    for d, ss, sai in ss_do_ativo:
        if sai is not None and sai < inicio:
            t = txt(d["DESCRIÇÃO"]).upper()
            velhas |= {t[max(0, m.start() - 40):m.end() + 40] for m in EXEC.finditer(t)}
    achados = []
    for d, ss, sai in sorted(ss_do_ativo, key=lambda x: x[0]["DTA_ABERTURA"]):
        if d["DTA_ABERTURA"] > limite or (sai is not None and sai < inicio):
            continue
        t = txt(d["DESCRIÇÃO"]).upper()
        for m in EXEC.finditer(t):
            trecho = t[max(0, m.start() - 40):m.end() + 40]
            depois = t[m.end():m.end() + 45]
            if trecho in velhas or DUVIDA.search(t[max(0, m.start() - 50):m.start()]):
                continue
            if ACESSORIO.search(depois) and not PECA_GRANDE.search(depois):
                continue
            if NAO_OPERA.search(t[m.end():m.end() + 300]):
                continue              # «ficou em operação? não»: troca parcial, a falha seguiu
            datas = _datas_perto(t, m.start(), m.end(), d["DTA_ABERTURA"])
            if datas:
                datas = [x for x in datas if inicio <= x <= min(limite, HOJE)]
                if not datas:
                    continue          # a frase fala de outra época
                quando = min(datas, key=lambda x: abs((x - d["DTA_ABERTURA"]).days))
                como = "data escrita junto"
            else:
                quando = sai or d["DTA_ABERTURA"]
                if not (inicio <= quando <= limite):
                    continue
                como = "saída da SS"
            achados.append((quando, f"texto em {ss} ({como}): «…{t[max(0, m.start() - 50):m.end() + 35].strip()}…»"))
    return min(achados, key=lambda o: o[0]) if achados else (None, "")


def criticidade(a, ges, mapa):
    c = ges.get(a, {}).get("criticidade", "")
    if c in PRAZO:
        return c, "Gestão COEP 5"
    if a in mapa:
        return mapa[a], "carteira (mapeamento)"
    return "Sem classificação", "sem criticidade"


def data_da_falha(f):
    return dt.datetime.strptime(f["ocorrencia"], "%d/%m/%Y")


def avalia(f, limite, por, antes, por_ativo, ges, mapa, reg_ativo, reg_local, local_cad):
    a = f["ativo"]
    nos = demanda(f["ss"], por, antes)
    if not nos:
        raise AssertionError(f"SS da falha fora da base: {f['ss']}")
    cabeca = nos[0]
    abertura = por[cabeca][0]["DTA_ABERTURA"]
    inicio = data_da_falha(f)
    d_cad, regra, ev_cad = troca_na_base(nos, por, inicio, limite)
    d_txt, ev_txt = troca_no_texto(a, inicio, limite, por, por_ativo)
    forte = (d_cad, regra, ev_cad) if regra in (1, 2) else None     # esteira: saiu do campo/COEP para ajuste
    nota = ""
    pendente_na_gestao = a in ges and ges[a]["status"] not in ("Realizado",) and limite >= HOJE
    if f["troca_rol"] == "sim":
        opcoes = [o for o in ((d_txt, "texto", ev_txt) if d_txt else None, forte) if o]
        if opcoes:
            data, regra, evid = min(opcoes, key=lambda o: o[0])
        elif d_cad:
            data, evid = d_cad, ev_cad
        else:
            data, regra, evid = None, None, ""
        trocada = data is not None
        if not trocada:
            nota = "o rol confirma a troca, mas a base não dá a data"
    else:
        novo_txt = (d_txt, "texto", ev_txt) if d_txt and d_txt > CORTE_DO_ROL else None
        opcoes = [o for o in (novo_txt, forte) if o]
        data, regra, evid = min(opcoes, key=lambda o: o[0]) if opcoes else (None, None, "")
        trocada = data is not None
        if trocada and data <= CORTE_DO_ROL:
            nota = "o rol diz sem troca; a base mostra o campo saindo para ajuste/comissionamento"
        elif trocada:
            nota = "troca depois da leitura do rol (19/08)"
        if trocada and pendente_na_gestao and data <= CORTE_DO_ROL:
            trocada, nota = False, "segue pendente na Gestão — a saída do campo não foi a troca"
    crit, fonte_crit = criticidade(a, ges, mapa)
    prazo = PRAZO.get(crit, PRAZO_SEM)
    dias_ss = None
    if trocada:
        dias = (data - inicio).days
        dias_ss = max(0, (data - abertura).days)
        sit = "trocada no prazo" if dias <= prazo else "trocada fora do prazo"
    elif f["troca_rol"] == "sim":
        dias, sit = None, "trocada sem data na base"
    else:
        ainda = limite < HOJE or pendente_na_gestao or any(
            txt(d["STATUS"]).upper() == "SS PENDENTE" for d in por_ativo.get(a, []))
        if ainda:
            fim = min(limite, HOJE)
            dias = (fim - inicio).days
            dias_ss = max(0, (fim - abertura).days)
            sit = "em aberto, prazo estourado" if dias > prazo else "em aberto, dentro do prazo"
            if limite < HOJE:
                nota = "falhou de novo antes de trocar — conta até a falha seguinte"
        else:
            dias, sit = None, "encerrada sem troca"
    sit_ss = ""
    if dias_ss is not None:
        sit_ss = ("no prazo" if dias_ss <= prazo else "fora do prazo") if trocada else \
            ("aberto estourado" if dias_ss > prazo else "aberto no prazo")
    if a in reg_ativo:
        mun, polo, reg = reg_ativo[a]
        fonte_reg = "carteira"
    else:
        mun = local_cad.get(a, "")
        polo, reg = reg_local.get(mun, ("", ""))
        fonte_reg = "localidade do cadastro de ajustes"
    return {**f, "regional": reg, "polo": polo, "municipio": mun, "fonte_regional": fonte_reg,
            "cabeca": cabeca, "abertura_ss": abertura.strftime("%d/%m/%Y"), "inicio": inicio.strftime("%d/%m/%Y"),
            "postos": " → ".join(txt(por[x][0]["POSTO_SGM"]).replace("ETO-", "") for x in nos),
            "criticidade": crit, "fonte_criticidade": fonte_crit, "prazo": prazo,
            "troca": data.strftime("%d/%m/%Y") if trocada else "", "regra": regra if trocada else "",
            "evidencia": evid if trocada else "", "dias": dias, "situacao": sit,
            "dias_pela_ss": dias_ss, "situacao_pela_ss": sit_ss, "nota": nota}


# ------------------------------------------------------------------ resumo
def resumo(itens, chave):
    grupos = defaultdict(list)
    for x in itens:
        grupos[chave(x)].append(x)
    out = {}
    for k, xs in grupos.items():
        c = Counter(x["situacao"] for x in xs)
        troc = c["trocada no prazo"] + c["trocada fora do prazo"]
        vencidas = troc + c["em aberto, prazo estourado"]
        dias_troca = [x["dias"] for x in xs if x["situacao"].startswith("trocada ") and x["dias"] is not None]
        out[k] = {"falhas": len(xs), "trocadas": troc, "no_prazo": c["trocada no prazo"],
                  "fora_do_prazo": c["trocada fora do prazo"],
                  "aberto_estourado": c["em aberto, prazo estourado"],
                  "aberto_no_prazo": c["em aberto, dentro do prazo"],
                  "sem_troca": c["encerrada sem troca"], "sem_data": c["trocada sem data na base"],
                  "sla": round(c["trocada no prazo"] / troc, 4) if troc else None,
                  "sla_com_abertos": round(c["trocada no prazo"] / vencidas, 4) if vencidas else None,
                  "sla_pela_ss": round(sum(1 for x in xs if x["situacao"].startswith("trocada ")
                                           and x["situacao_pela_ss"] == "no prazo") / troc, 4) if troc else None,
                  "mediana_dias": median(dias_troca) if dias_troca else None}
    return out


# ------------------------------------------------------------------ tempo até consertar
JANELAS = (30, 60, 90, 180, 365)
MESES = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")


def _dia(s):
    return dt.datetime.strptime(s, "%d/%m/%Y")


def _trocada(x):
    return x["situacao"].startswith("trocada ") and x["dias"] is not None


def _bloco(xs):
    dt_ = [x["dias"] for x in xs if _trocada(x)]
    da = [x["dias"] for x in xs if x["situacao"].startswith("em aberto")]
    return {"falhas": len(xs), "trocadas": len(dt_),
            "mediana_trocadas": median(dt_) if dt_ else None,
            "media_trocadas": round(sum(dt_) / len(dt_), 1) if dt_ else None,
            "mais_rapida": min(dt_) if dt_ else None, "mais_lenta": max(dt_) if dt_ else None,
            "abertas": len(da), "mediana_abertas": median(da) if da else None,
            "mais_antiga_aberta": max(da) if da else None,
            "mediana_minima": median(dt_ + da) if dt_ + da else None,
            "fora_da_conta": len(xs) - len(dt_) - len(da)}


def tempo(itens):
    """Da falha à troca, por ano da falha. A trocada dá o tempo que terminou; a aberta, o que já
    passou e ainda corre — por isso a mediana mínima (abertas trocadas hoje) e a mesma janela."""
    out = {"por_ano": {}, "por_tipo": {}, "janela": {}, "ano_da_troca": {}, "mes_da_troca": {}}
    mes = defaultdict(Counter)
    for ano in ANOS:
        xs = [x for x in itens if x["ano"] == ano]
        out["por_ano"][ano] = _bloco(xs)
        for tipo in ("RL", "RT"):
            out["por_tipo"][(ano, tipo)] = _bloco([x for x in xs if x["tipo"] == tipo])
        # mesma janela: só entra a falha que já teve N dias para ser trocada
        for n in JANELAS:
            base = [x for x in xs if (_trocada(x) and (HOJE - _dia(x["inicio"])).days >= n)
                    or (x["situacao"].startswith("em aberto") and x["dias"] >= n)]
            feitas = sum(1 for x in base if _trocada(x) and x["dias"] <= n)
            out["janela"][(ano, n)] = {"trocadas": feitas, "base": len(base),
                                       "pct": round(feitas / len(base), 4) if base else None}
        mesmo, depois = [], []
        for x in filter(_trocada, xs):
            t = _dia(x["troca"])
            mes[(t.year, t.month)][ano] += 1
            (mesmo if t.year == ano else depois).append(x["dias"])
        out["ano_da_troca"][ano] = {"mesmo_ano": len(mesmo), "depois": len(depois),
                                    "faixa_mesmo": (min(mesmo), max(mesmo)) if mesmo else None,
                                    "faixa_depois": (min(depois), max(depois)) if depois else None}
    out["mes_da_troca"] = {k: dict(v) for k, v in sorted(mes.items())}
    for ano in ANOS:                              # trocadas + abertas + fora da conta = falhas
        s = out["por_ano"][ano]
        assert s["trocadas"] + s["abertas"] + s["fora_da_conta"] == s["falhas"]
        assert sum(v.get(ano, 0) for v in out["mes_da_troca"].values()) == s["trocadas"]
    return out


# ------------------------------------------------------------------ planilha
TINTA, PAPEL, SINAL = "FF211D15", "FFF2EFE6", "FFBC4B0E"
PAPEL2, FILETE = "FFE9E5D8", "FFC8C2AF"


def cabeca_aba(ws, titulo, texto, cols, larg, linha=1):
    ws.cell(row=linha, column=1, value=titulo).font = Font(bold=True, size=13, color=SINAL)
    if texto:
        ws.cell(row=linha + 1, column=1, value=texto).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=linha + 1, start_column=1, end_row=linha + 1, end_column=len(cols))
        ws.row_dimensions[linha + 1].height = 44
    r = linha + 3
    for i, (c, w) in enumerate(zip(cols, larg), 1):
        cel = ws.cell(row=r, column=i, value=c)
        cel.font = Font(bold=True, color=PAPEL)
        cel.fill = PatternFill("solid", fgColor=TINTA)
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        if w:
            ws.column_dimensions[get_column_letter(i)].width = w
    return r + 1


def linhas(ws, r0, dados, formatos=None):
    lado = Side(style="thin", color=FILETE)
    for j, l in enumerate(dados):
        for i, v in enumerate(l, 1):
            cel = ws.cell(row=r0 + j, column=i, value=v)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.border = Border(bottom=lado)
            if formatos and i in formatos:
                cel.number_format = formatos[i]
            if j % 2:
                cel.fill = PatternFill("solid", fgColor=PAPEL2)
    return r0 + len(dados)


def _faixa(f):
    return "—" if not f else (f"{f[0]} dias" if f[0] == f[1] else f"{f[0]} a {f[1]} dias")


def aba_tempo(ws, tmp):
    negrito = Font(bold=True)
    cols = ["Ano da falha", "Falhas", "Trocadas (com data)", "Mediana das trocadas (dias)",
            "Média das trocadas (dias)", "Mais rápida (dias)", "Mais lenta (dias)", "Ainda sem troca",
            "Dias parados nas abertas (mediana)", "A mais antiga aberta (dias)",
            "Mediana mínima: se as abertas fossem trocadas hoje", "Fora da conta (sem data ou sem troca)"]
    r = cabeca_aba(ws, "Quanto tempo levamos para consertar — falhas de 2025 × falhas de 2026",
                   "Da data da falha até a peça trocada em campo, a mesma régua do SLA. A mediana das "
                   "trocadas engana na comparação: das falhas de 2026 só as rápidas já terminaram, e as "
                   "que seguem abertas vão puxar o número para cima. Por isso vão junto a mediana mínima "
                   "— a que sairia se todas as abertas fossem trocadas hoje, 23/09 — e, no quadro de "
                   "baixo, a comparação na mesma janela.",
                   cols, [11, 8, 11, 12, 12, 10, 10, 10, 13, 12, 17, 13])
    fim = linhas(ws, r, [[ano, s["falhas"], s["trocadas"], s["mediana_trocadas"], s["media_trocadas"],
                          s["mais_rapida"], s["mais_lenta"], s["abertas"], s["mediana_abertas"],
                          s["mais_antiga_aberta"], s["mediana_minima"], s["fora_da_conta"]]
                         for ano, s in tmp["por_ano"].items()])

    r = cabeca_aba(ws, "Na mesma janela — quantas foram trocadas em até N dias",
                   "Só entra a falha que já teve os N dias para ser trocada: a de 2026 que aconteceu há "
                   "60 dias não entra na janela de 90. Assim os dois anos são medidos pela mesma régua.",
                   ["Em até", "2025 — trocadas", "2025 — base", "2025 — %",
                    "2026 — trocadas", "2026 — base", "2026 — %"], [None] * 7, linha=fim + 2)
    dados = []
    for n in JANELAS:
        l = [f"{n} dias"]
        for ano in ANOS:
            j = tmp["janela"][(ano, n)]
            l += [j["trocadas"], j["base"], j["pct"]] if j["base"] else ["—", 0, "ainda não completou"]
        dados.append(l)
    fim = linhas(ws, r, dados, {4: "0%", 7: "0%"})

    r = cabeca_aba(ws, "Religador e regulador", None,
                   ["Ano da falha", "Tipo", "Falhas", "Trocadas (com data)", "Mediana das trocadas (dias)",
                    "Ainda sem troca", "Dias parados nas abertas (mediana)",
                    "Mediana mínima: se as abertas fossem trocadas hoje"], [None] * 8, linha=fim + 2)
    fim = linhas(ws, r, [[ano, tipo, s["falhas"], s["trocadas"], s["mediana_trocadas"], s["abertas"],
                          s["mediana_abertas"], s["mediana_minima"]]
                         for (ano, tipo), s in tmp["por_tipo"].items()])

    r = cabeca_aba(ws, "Trocada no próprio ano da falha ou só no ano seguinte", None,
                   ["Ano da falha", "Trocadas no mesmo ano", "Levaram", "Trocadas só no ano seguinte",
                    "Levaram"], [None] * 5, linha=fim + 2)
    fim = linhas(ws, r, [[ano, a["mesmo_ano"], _faixa(a["faixa_mesmo"]), a["depois"], _faixa(a["faixa_depois"])]
                         for ano, a in tmp["ano_da_troca"].items()])

    r = cabeca_aba(ws, "Mês em que a peça foi trocada",
                   "Cada troca com data, pelo mês em que aconteceu e pelo ano da falha que ela resolveu.",
                   ["Mês da troca", "Falhas de 2025", "Falhas de 2026", "Total"], [None] * 4, linha=fim + 2)
    dados = [[f"{MESES[m - 1]}/{a}", v.get(2025, 0), v.get(2026, 0), sum(v.values())]
             for (a, m), v in tmp["mes_da_troca"].items()]
    dados.append(["Total", sum(l[1] for l in dados), sum(l[2] for l in dados), sum(l[3] for l in dados)])
    fim = linhas(ws, r, dados)
    for c in range(1, 5):
        ws.cell(row=fim - 1, column=c).font = negrito


def grava(itens, por_reg, por_reg_tipo, por_ano, tmp):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo"
    cols = ["Ano", "Regional", "Falhas", "Trocadas", "No prazo", "Fora do prazo", "SLA (no prazo ÷ trocadas)",
            "Em aberto, prazo estourado", "SLA contando os abertos vencidos", "Em aberto, dentro do prazo",
            "Encerradas sem troca", "Trocadas sem data", "Mediana de dias até a troca",
            "SLA contando da abertura da SS"]
    r = cabeca_aba(ws, "SLA das falhas da taxa, por regional",
                   "Da data da falha (a mesma da taxa) até a peça trocada em campo, contra o prazo da "
                   "proposta DCMD (Muito Alta 11 · Alta 20 · Média 40 · Baixa 60 · sem classificação "
                   "32,75). SLA = trocadas no prazo ÷ trocadas, a mesma conta do «% de substituição no "
                   "prazo». Em aberto conta até 23/09/2026, a posição da base de SS. A última coluna conta "
                   "da abertura da primeira SS — segunda leitura, que não data a falha (ver «Como foi feito»).",
                   cols, [7, 11, 8, 10, 9, 9, 14, 14, 14, 13, 12, 11, 13, 14])
    dados = []
    for ano in ANOS:
        for reg in REGIONAIS:
            s = por_reg.get((ano, reg))
            if s:
                dados.append([ano, reg.title(), s["falhas"], s["trocadas"], s["no_prazo"], s["fora_do_prazo"],
                              s["sla"], s["aberto_estourado"], s["sla_com_abertos"], s["aberto_no_prazo"],
                              s["sem_troca"], s["sem_data"], s["mediana_dias"], s["sla_pela_ss"]])
        s = por_ano[ano]
        dados.append([ano, "Total", s["falhas"], s["trocadas"], s["no_prazo"], s["fora_do_prazo"], s["sla"],
                      s["aberto_estourado"], s["sla_com_abertos"], s["aberto_no_prazo"], s["sem_troca"],
                      s["sem_data"], s["mediana_dias"], s["sla_pela_ss"]])
    fim = linhas(ws, r, dados, {7: "0%", 9: "0%", 14: "0%"})
    for rr in range(r, fim):
        if ws.cell(row=rr, column=2).value == "Total":
            for c in range(1, len(cols) + 1):
                ws.cell(row=rr, column=c).font = Font(bold=True)

    r = cabeca_aba(ws, "Por regional e tipo", None, ["Ano", "Regional", "Tipo", "Falhas", "Trocadas",
                                                     "No prazo", "SLA", "Em aberto, prazo estourado"],
                   [None] * 8, linha=fim + 2)
    dados = []
    for ano in ANOS:
        for reg in REGIONAIS:
            for tipo in ("RL", "RT"):
                s = por_reg_tipo.get((ano, reg, tipo))
                if s:
                    dados.append([ano, reg.title(), tipo, s["falhas"], s["trocadas"], s["no_prazo"], s["sla"],
                                  s["aberto_estourado"]])
    linhas(ws, r, dados, {7: "0%"})
    ws.freeze_panes = "A5"

    aba_tempo(wb.create_sheet("Tempo até consertar"), tmp)

    ws = wb.create_sheet("Por falha")
    cols = ["Ano", "Tipo", "Ativo", "Regional", "Polo", "Município", "Peça", "SS da falha", "Data da falha",
            "Primeira SS da demanda", "Abertura dela", "Postos da demanda", "Criticidade", "Prazo (dias)",
            "Troca em", "Como a troca foi achada", "Dias", "Situação", "Dias pela abertura da SS",
            "Situação pela abertura da SS", "Rol diz troca feita?", "Nota"]
    r = cabeca_aba(ws, "As 90 falhas, uma por linha",
                   "O texto da SS manda («foi substituído», «substituído conforme solicitado»… com a data "
                   "escrita junto). Sem texto, a cadeia: 1) COCM depois do COEP sai para PROT/TELE/SE/DMSL "
                   "ou fecha atendida; 2) COEP repassa direto para PROT/SE; 3) TELE/SE/DMSL fecha atendida; "
                   "4) COEP fecha atendida. Quando o rol diz que não houve troca, só valem as regras 1 e 2 "
                   "e texto posterior a 19/08.",
                   cols, [6, 5, 12, 9, 16, 18, 9, 18, 11, 18, 11, 40, 11, 8, 10, 50, 6, 22, 9, 16, 9, 40])
    ordem = sorted(itens, key=lambda x: (x["ano"], REGIONAIS.index(x["regional"]), x["tipo"], x["ativo"]))
    linhas(ws, r, [[x["ano"], x["tipo"], x["ativo"], x["regional"].title(), x["polo"].title(),
                    x["municipio"].title(), x["peca"], x["ss"], x["inicio"], x["cabeca"], x["abertura_ss"],
                    x["postos"], x["criticidade"], x["prazo"], x["troca"], x["evidencia"], x["dias"],
                    x["situacao"], x["dias_pela_ss"], x["situacao_pela_ss"], x["troca_rol"], x["nota"]]
                   for x in ordem])
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(ordem)}"

    ws = wb.create_sheet("Como foi feito")
    for l in (__doc__ or "").strip().splitlines():
        ws.append([l])
    ws.column_dimensions["A"].width = 110
    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    wb.save(SAIDA)


def main():
    falhas = rol()
    por, antes, por_ativo = base()
    ges = gestao()
    mapa = sm.criticidades()
    reg_ativo, reg_local, local_cad = regionais()
    # janela de cada falha: até o início da próxima falha do mesmo ativo (a troca dela não é desta)
    inicios = {id(f): data_da_falha(f) for f in falhas}
    itens = []
    for f in falhas:
        seguintes = [inicios[id(g)] for g in falhas if g["ativo"] == f["ativo"] and inicios[id(g)] > inicios[id(f)]]
        limite = min(seguintes) if seguintes else HOJE
        itens.append(avalia(f, limite, por, antes, por_ativo, ges, mapa, reg_ativo, reg_local, local_cad))
    assert all(x["regional"] in REGIONAIS for x in itens), [x["ativo"] for x in itens if not x["regional"]]
    por_reg = resumo(itens, lambda x: (x["ano"], x["regional"]))
    por_reg_tipo = resumo(itens, lambda x: (x["ano"], x["regional"], x["tipo"]))
    por_ano = resumo(itens, lambda x: x["ano"])
    for ano in ANOS:                              # as regionais somam o ano
        assert sum(por_reg[(ano, r)]["falhas"] for r in REGIONAIS if (ano, r) in por_reg) == por_ano[ano]["falhas"]
    tmp = tempo(itens)
    grava(itens, por_reg, por_reg_tipo, por_ano, tmp)
    with open(JSON, "w", encoding="utf-8") as fh:
        json.dump({"posicao": HOJE.strftime("%Y-%m-%d"), "prazo": {**PRAZO, "Sem classificação": PRAZO_SEM},
                   "por_ano": {str(k): v for k, v in por_ano.items()},
                   "por_regional": {f"{k[0]} {k[1]}": v for k, v in por_reg.items()},
                   "tempo": {"por_ano": {str(k): v for k, v in tmp["por_ano"].items()},
                             "por_tipo": {f"{k[0]} {k[1]}": v for k, v in tmp["por_tipo"].items()},
                             "janela": {f"{k[0]} em até {k[1]} dias": v for k, v in tmp["janela"].items()},
                             "ano_da_troca": {str(k): v for k, v in tmp["ano_da_troca"].items()},
                             "mes_da_troca": {f"{k[0]}-{k[1]:02d}": v for k, v in tmp["mes_da_troca"].items()}},
                   "itens": itens}, fh, ensure_ascii=False, indent=1, default=str)
    for ano in ANOS:
        for reg in REGIONAIS + ("Total",):
            s = por_ano[ano] if reg == "Total" else por_reg.get((ano, reg))
            if s:
                print(f"{ano} {reg:6s} falhas {s['falhas']:2d} · trocadas {s['trocadas']:2d} · no prazo "
                      f"{s['no_prazo']:2d} · SLA {s['sla'] if s['sla'] is None else round(100 * s['sla'])}% "
                      f"(pela SS {s['sla_pela_ss'] if s['sla_pela_ss'] is None else round(100 * s['sla_pela_ss'])}%) · "
                      f"aberto estourado {s['aberto_estourado']:2d} · aberto no prazo {s['aberto_no_prazo']} · "
                      f"sem troca {s['sem_troca']} · sem data {s['sem_data']} · mediana {s['mediana_dias']}")
    for ano in ANOS:
        s, j = tmp["por_ano"][ano], tmp["janela"]
        print(f"{ano} tempo: trocadas {s['trocadas']} (mediana {s['mediana_trocadas']}) · abertas {s['abertas']} "
              f"(paradas há {s['mediana_abertas']}) · mediana mínima {s['mediana_minima']} · janela " +
              " · ".join(f"{n}d {j[(ano, n)]['trocadas']}/{j[(ano, n)]['base']}" for n in JANELAS))
    print(Counter(x["situacao"] for x in itens), Counter(str(x["regra"]) for x in itens if x["regra"]))
    print(Counter(x["fonte_criticidade"] for x in itens), Counter(x["fonte_regional"] for x in itens))
    return itens


if __name__ == "__main__":
    main()

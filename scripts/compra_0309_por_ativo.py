"""
A segunda compra de religador (03/09/2026), peça por peça, nas abas Estoque e Gestão da COEP 4.

Pedido do gestor (25/09): «atualize a segunda compra na aba Gestão e na aba Estoque, das outras
criticidades — essa foi a segunda compra de RL». A compra veio numa foto, com a requisição
Web Supply **29756055** de 03/09/2026:

| PMA   | Código | Peça                     | Qtd | Unitário      |
|-------|--------|--------------------------|-----|---------------|
| 38744 | 690005 | tanque de religador 34,5 | 17  | R$ 16.130,41  |
| 38745 | 690001 | tanque de religador 13,8 | 2   | R$ 11.990,27  |
| 38746 | 690916 | controle de religador 13,8 | 2 | R$ 29.851,96  |
| 38747 | 692263 | controle de religador 34,5 | 4 | R$ 39.741,39  |

**A régua é a mesma da primeira compra** (`compra_2907_por_ativo.py`), e tudo sai da própria
planilha: o que cada ativo já recebeu da primeira compra é lido da `Tabela7`; a segunda completa
o que falta. Muito Alta e Alta primeiro, depois a ordem da coluna Índice; só recebe quem a compra
conserta por inteiro; não recebe status «Em logistica (N1>N3)» nem «Realizado».

**Depois, a peça que sobrou vai para quem ainda não tem PMA e precisa dela** (gestor, 25/09: «se
tem ativo pra usar que já não tem PMA associado use, e que precisam dessa peça também; se não tem
ativo pra usar realmente deixe como reserva»), mesmo que ainda falte outra peça do conserto — é o
caso dos RL Completo de 34,5 que ficam com o tanque esperando o controle. O que ninguém usa fica
como «Reserva».

**Preço: o da primeira compra, por enquanto** (gestor, 25/09: «utilize o preço da compra anterior
por enquanto»). O unitário de cada código sai das linhas 36–75 da própria `Tabela7`, e o total é
qtd × unitário, como na primeira compra. O que a requisição trouxe fica registrado em `COMPRA2`.

**Status: «Em aprovação corporativa»** nas 25 linhas novas e no quadro «Status PMA» (gestor, 25/09:
«o status atual dela é que está em aprovação corporativo»). **Previsão de chegada em branco**:
«ainda não temos exatamente a previsão de chegada».

Edição direta no XML (como nas anteriores): a `Tabela7` cresce de A35:M75 para A35:M100, o quadro
«Status PMA» ganha as linhas 13 a 16, as fórmulas de saldo passam a olhar até a linha 100, e a
coluna M da Gestão recebe os PMAs novos.

Base: `data/raw/GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4_ENTREGUE_2409.xlsx` (o que
`compra_2907_por_ativo.py` gerou e foi entregue em 24/09). Saída:
`dist/GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4.xlsx`.

Rodar: python3 scripts/compra_0309_por_ativo.py
"""

import datetime as dt
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from compra_2907_por_ativo import (JA_TEM, MA_ALTA, RESERVA, Textos, _t, cel, excel_data,  # noqa: E402
                                   linha_xml, todas_as_celulas)

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.path.join(RAIZ, "data", "raw", "GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4_ENTREGUE_2409.xlsx")
SAIDA = os.path.join(RAIZ, "dist", "GESTAO_EQUIPAMENTOS_ESPECIAIS_COEP_4.xlsx")
ABA_GESTAO, ABA_ESTOQUE = "xl/worksheets/sheet3.xml", "xl/worksheets/sheet5.xml"
TABELA = "xl/tables/table4.xml"                      # a Tabela7 (conferido pelo name= no XML)

DATA_REQ = dt.date(2026, 9, 3)
STATUS2 = "Em aprovação corporativa"                  # gestor, 25/09; previsão ainda não existe
REQ_WEB = 29756055
# (PMA, código, descrição, qtd, unitário e total DA REQUISIÇÃO — não vão para a planilha, que usa o
# preço da 1ª compra por ordem do gestor). A foto corta em 45 caracteres o nome do 690001 e do
# 690916 («…12,5K», «…115V»); vai o nome de catálogo que a 1ª compra já usa para o mesmo código,
# senão um filtro ou dinâmica por Descrição separa a mesma peça em duas (achado da verificação, 25/09).
COMPRA2 = [
    ("38744", 690005, "RELIGADOR AUTO EXT 3F LIN VAC 38KV 630A 12,5KA 127VCA", 17, 16130.41, 274216.97),
    ("38745", 690001, "RELIGADOR AUTO EXT 3F LIN VAC 15KV 630A 12,5KA 127VCA", 2, 11990.27, 23980.54),
    ("38746", 690916, "CONTROLE P/ RELIGADOR LINHA DISTR 15,0KV 115VCA", 2, 29851.96, 59703.92),
    ("38747", 692263, "CONTROLE P/ RELIGADOR LINHA DISTR 36,2KV 115VCA", 4, 39741.39, 158965.56),
]
# peça de cada PMA, das duas compras
PECA = {"38291": "controle RT", "38292": "célula 400", "38294": "tanque 13,8", "38295": "tanque 34,5",
        "38296": "controle 13,8", "38297": "controle 34,5", "38744": "tanque 34,5",
        "38745": "tanque 13,8", "38746": "controle 13,8", "38747": "controle 34,5"}
# peças de cada defeito (a mesma BOM da primeira compra, em peça em vez de PMA)
BOM = {"RL Completo 13,8": ["tanque 13,8", "controle 13,8"],
       "RL Completo 34,5": ["tanque 34,5", "controle 34,5"],
       "Tanque 13,8": ["tanque 13,8"], "Tanque 34,5": ["tanque 34,5"],
       "Controle 34,5": ["controle 34,5"],
       "RT Completo 400 34,5": ["célula 400"] * 3 + ["controle RT"],
       "Célula 400 34,5": ["célula 400"], "Controle 200 34,5": ["controle RT"],
       "Controle 200 13,8": ["controle RT"], "Controle 167 13,8": ["controle RT"]}
PRIMEIRA = {"38291", "38292", "38294", "38295", "38296", "38297"}
ULTIMA = 100                                          # 75 + 25 peças novas


def ler(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb["Gestão"]
    col = {_t(c.value): i for i, c in enumerate(ws[1], 1)}
    fim = int(re.sub(r"\D", "", ws.tables["Table1"].ref.split(":")[1]))
    g = {}
    for r in range(2, fim + 1):
        a = _t(ws.cell(row=r, column=col["Ativo"]).value)
        if a:
            v = lambda n: ws.cell(row=r, column=col[n]).value
            g[a] = dict(linha=r, ativo=a, crit=_t(v("Criticidade")), status=_t(v("Status")),
                        defeito=_t(v("Defeito")), indice=v("Índice"), pma=_t(v("PMA")))
    we = wb["Estoque"]
    assert we.tables["Tabela7"].ref == "A35:M75", we.tables["Tabela7"].ref
    ja = [dict(linha=r, pma=_t(we.cell(row=r, column=1).value), ativo=_t(we.cell(row=r, column=11).value))
          for r in range(36, 76)]
    precos = defaultdict(set)
    for r in range(36, 76):
        precos[we.cell(row=r, column=2).value].add(we.cell(row=r, column=6).value)
    preco_1 = {}
    for cod, v in precos.items():
        assert len(v) == 1, (cod, v)                    # um preço só por código na 1ª compra
        preco_1[cod] = v.pop()
    return g, ja, preco_1


def divide(g, ja):
    recebido = defaultdict(Counter)
    for u in ja:
        if u["ativo"] and u["ativo"] != RESERVA:
            recebido[u["ativo"]][PECA[u["pma"]]] += 1
    estoque = {pma: q for pma, _, _, q, _, _ in COMPRA2}
    pma_da_peca = {PECA[p]: p for p in estoque}
    lista = {p: [] for p in estoque}
    novo = defaultdict(Counter)
    fila = sorted(g.values(), key=lambda x: (0 if x["crit"] in MA_ALTA else 1, x["indice"]))
    sem = {}
    for x in fila:
        a = x["ativo"]
        if x["status"] in JA_TEM or not BOM.get(x["defeito"]):
            continue
        falta = Counter(BOM[x["defeito"]]) - recebido[a]
        if not falta:
            continue
        if any(p not in pma_da_peca for p in falta):          # célula ou controle de RT: não há
            continue
        if all(estoque[pma_da_peca[p]] >= q for p, q in falta.items()):
            for p, q in falta.items():
                pma = pma_da_peca[p]
                estoque[pma] -= q
                lista[pma] += [a] * q
                novo[a][pma] += q
        else:
            sem[a] = sorted(p for p, q in falta.items() if estoque[pma_da_peca[p]] < q)
    # a peça que sobrou vai para quem ainda NÃO tem PMA e precisa dela, na mesma ordem
    parcial = {}
    for x in fila:
        a = x["ativo"]
        if x["status"] in JA_TEM or not BOM.get(x["defeito"]):
            continue
        if recebido[a] or novo[a] or (x["pma"] and x["pma"] != "Sem PMA"):
            continue                                            # já tem PMA associado
        for p, q in Counter(BOM[x["defeito"]]).items():
            pma = pma_da_peca.get(p)
            if pma and estoque[pma]:
                k = min(q, estoque[pma])
                estoque[pma] -= k
                lista[pma] += [a] * k
                novo[a][pma] += k
        if novo[a]:
            ainda = Counter(BOM[x["defeito"]]) - Counter({PECA[p]: q for p, q in novo[a].items()})
            parcial[a] = sorted(ainda)
            sem.pop(a, None)
    for pma, q in estoque.items():
        lista[pma] += [RESERVA] * q
    return lista, {a: c for a, c in novo.items() if c}, sem, recebido, parcial


def main():
    g, ja, preco_1 = ler(BASE)
    lista, novo, sem, recebido, parcial = divide(g, ja)
    for pma, cod, _, q, unit, tot in COMPRA2:            # a foto fecha: qtd × unitário = total
        assert round(q * unit, 2) == tot, pma
        assert len(lista[pma]) == q, pma
        assert cod in preco_1, cod                      # todo código da 2ª compra existe na 1ª
    for a, c in novo.items():
        tudo = recebido[a] + Counter({PECA[p]: q for p, q in c.items()})
        if a in parcial:                                # peça que sobrou: parte do conserto, sem PMA antes
            assert tudo < Counter(BOM[g[a]["defeito"]]) and not recebido[a] and g[a]["pma"] == "Sem PMA", a
        else:                                           # quem recebe primeiro fica com o conserto inteiro
            assert tudo == Counter(BOM[g[a]["defeito"]]), (a, tudo)
        assert g[a]["status"] not in JA_TEM

    # coluna M da Gestão: os PMAs das duas compras
    pma_novo = {}
    for a, c in novo.items():
        antes = g[a]["pma"]
        velhos = set(antes.split("+")) if re.fullmatch(r"[\d+]+", antes or "") else set()
        pma_novo[a] = "+".join(sorted(velhos | set(c)))

    with zipfile.ZipFile(BASE) as z:
        infos = z.infolist()
        partes = {i.filename: z.read(i.filename) for i in infos}
    assert 'name="Tabela7"' in partes[TABELA].decode("utf-8")
    textos = Textos(partes["xl/sharedStrings.xml"].decode("utf-8"))
    styles = partes["xl/styles.xml"].decode("utf-8")
    xfs = re.findall(r"<xf [^>]*?(?:/>|>.*?</xf>)", re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", styles, re.S).group(1), re.S)
    estilo_data = next(i for i, xf in enumerate(xfs)
                       if all(k in xf for k in ('numFmtId="14"', 'fillId="12"', 'borderId="19"')))

    # ---- aba Estoque ----
    x = partes[ABA_ESTOQUE].decode("utf-8")
    linhas = []
    n = 76
    for pma, cod, desc, q, _, _ in COMPRA2:
        unit = preco_1[cod]                             # preço da 1ª compra, por enquanto
        tot = round(q * unit, 2)
        for ativo in lista[pma]:
            c = (cel(f"A{n}", int(pma)) + cel(f"B{n}", cod) + cel(f"C{n}", desc, textos=textos)
                 + cel(f"D{n}", "UN", textos=textos) + cel(f"E{n}", q) + cel(f"F{n}", unit, s=38)
                 + cel(f"G{n}", tot, s=38) + cel(f"H{n}", "Material para Manutenção", textos=textos)
                 + cel(f"I{n}", excel_data(DATA_REQ), s=12) + cel(f"J{n}", REQ_WEB)
                 + cel(f"K{n}", ativo, textos=textos) + cel(f"L{n}", STATUS2, textos=textos)
                 + f'<c r="M{n}" s="12"/>')      # previsão vazia, já no formato de data
            linhas.append(f'<row r="{n}" spans="1:13" x14ac:dyDescent="0.25">{c}</row>')
            n += 1
    assert n - 1 == ULTIMA
    assert linha_xml(x, 76) is None
    x = x.replace("</sheetData>", "".join(linhas) + "</sheetData>", 1)
    x = x.replace('<dimension ref="A1:M75"/>', f'<dimension ref="A1:M{ULTIMA}"/>', 1)
    # saldo: as fórmulas das linhas 7-12 passam a olhar até a 100
    for r in range(7, 13):
        velha = linha_xml(x, r)
        nova = velha.replace("$A$36:$A$75", f"$A$36:$A${ULTIMA}").replace("$K$36:$K$75", f"$K$36:$K${ULTIMA}")
        assert nova != velha
        x = x.replace(velha, nova, 1)
    # quadro «Status PMA»: uma linha por PMA novo (13 a 16)
    blocos = []
    for i, (pma, _, _, q, _, _) in enumerate(COMPRA2):
        r = 13 + i
        com = sum(1 for a in lista[pma] if a != RESERVA)
        rp, rk = f"$A$36:$A${ULTIMA}", f"$K$36:$K${ULTIMA}"
        c = (cel(f"A{r}", int(pma), s=141) + cel(f"B{r}", excel_data(DATA_REQ), s=estilo_data)
             + cel(f"C{r}", REQ_WEB, s=141) + cel(f"D{r}", STATUS2, s=141, textos=textos)
             + f'<c r="E{r}" s="{estilo_data}"/>'
             + cel(f"F{r}", q, s=141, formula=f"COUNTIF({rp},A{r})")
             + cel(f"G{r}", com, s=141, formula=f'COUNTIFS({rp},A{r},{rk},"<>",{rk},"<>{RESERVA}")')
             + cel(f"H{r}", q - com, s=141, formula=f"F{r}-G{r}"))
        assert linha_xml(x, r) is None
        blocos.append(f'<row r="{r}" spans="1:8" x14ac:dyDescent="0.25">{c}</row>')
    l12 = linha_xml(x, 12)
    x = x.replace(l12, l12 + "".join(blocos), 1)
    reserva_1 = sum(1 for u in ja if u["ativo"] == RESERVA)
    saldo = reserva_1 + sum(lista[p].count(RESERVA) for p in lista)
    l1 = linha_xml(x, 1)
    b1 = re.search(r'<c r="B1"[^>]*>.*?</c>', l1, re.S).group(0)
    x = x.replace(l1, l1.replace(b1, cel("B1", saldo, formula="SUM(H7:H16)")), 1)
    partes[ABA_ESTOQUE] = x.encode("utf-8")

    # ---- a Tabela7 cresce ----
    t = partes[TABELA].decode("utf-8")
    assert t.count('ref="A35:M75"') == 2
    partes[TABELA] = t.replace('ref="A35:M75"', f'ref="A35:M{ULTIMA}"').encode("utf-8")

    # ---- aba Gestão, coluna M ----
    y = partes[ABA_GESTAO].decode("utf-8")
    for a, v in pma_novo.items():
        ref = f"M{g[a]['linha']}"
        m = re.search(rf'<c r="{ref}"([^>]*?)(?:/>|>(.*?)</c>)', y, re.S)
        attrs = re.sub(r'\s*t="[^"]*"', "", m.group(1))
        y = y.replace(m.group(0), f'<c r="{ref}"{attrs} t="s"><v>{textos.de(v)}</v></c>', 1)
    partes[ABA_GESTAO] = y.encode("utf-8")

    refs = sum(len(re.findall(r'<c [^>]*t="s"', b.decode("utf-8"))) for nm, b in partes.items()
               if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", nm))
    partes["xl/sharedStrings.xml"] = textos.grava(refs).encode("utf-8")
    with zipfile.ZipFile(SAIDA, "w", zipfile.ZIP_DEFLATED) as z:
        for i in infos:
            z.writestr(i, partes[i.filename], compress_type=zipfile.ZIP_DEFLATED)

    # ---- conferência do arquivo gravado ----
    with zipfile.ZipFile(BASE) as a_, zipfile.ZipFile(SAIDA) as b_:
        assert a_.namelist() == b_.namelist()
        mexidas = {nm for nm in a_.namelist() if a_.read(nm) != b_.read(nm)}
    assert mexidas == {ABA_ESTOQUE, ABA_GESTAO, TABELA, "xl/sharedStrings.xml"}, mexidas
    antes, depois = todas_as_celulas(BASE), todas_as_celulas(SAIDA)
    dif = {k for k in set(antes) | set(depois) if antes.get(k) != depois.get(k)}
    permitido = ({("Estoque", f"{c}{r}") for c in "ABCDEFGHIJKLM" for r in range(76, ULTIMA + 1)}
                 | {("Estoque", f"{c}{r}") for c in "ABCDEFGH" for r in range(13, 17)}
                 | {("Estoque", f"{c}{r}") for c in "FG" for r in range(7, 13)}
                 | {("Estoque", "B1")}
                 | {("Gestão", f"M{g[a]['linha']}") for a in pma_novo})
    assert dif <= permitido, sorted(dif - permitido)
    wv = openpyxl.load_workbook(SAIDA, data_only=True)
    e, gg = wv["Estoque"], wv["Gestão"]
    assert e.tables["Tabela7"].ref == f"A35:M{ULTIMA}"
    # a primeira compra (linhas 36-75) não mudou: nenhuma célula dela está em `permitido`
    tudo = defaultdict(Counter)
    for r in range(36, ULTIMA + 1):
        a = _t(e.cell(row=r, column=11).value)
        if a != RESERVA:
            tudo[a][PECA[_t(e.cell(row=r, column=1).value)]] += 1
    for a, c in tudo.items():
        assert c <= Counter(BOM[g[a]["defeito"]]), (a, c)          # ninguém recebe peça a mais
        pmas = {_t(e.cell(row=r, column=1).value) for r in range(36, ULTIMA + 1)
                if _t(e.cell(row=r, column=11).value) == a}
        assert _t(gg.cell(row=g[a]["linha"], column=13).value) == "+".join(sorted(pmas)), a
    for i, (pma, _, _, q, _, _) in enumerate(COMPRA2):
        r = 13 + i
        assert e.cell(row=r, column=6).value == q
        assert e.cell(row=r, column=7).value + e.cell(row=r, column=8).value == q
    assert e["B1"].value == saldo
    for r in range(76, ULTIMA + 1):                      # status novo, previsão em branco
        assert e.cell(row=r, column=12).value == STATUS2 and e.cell(row=r, column=13).value is None
    desc = defaultdict(set)
    for r in range(36, ULTIMA + 1):
        desc[e.cell(row=r, column=2).value].add(e.cell(row=r, column=3).value)
    assert all(len(v) == 1 for v in desc.values()), desc
    for r in range(13, 17):
        assert e.cell(row=r, column=4).value == STATUS2 and e.cell(row=r, column=5).value is None
    for r in range(76, ULTIMA + 1):                      # preço = o da 1ª compra para o mesmo código
        cod, qt = e.cell(row=r, column=2).value, e.cell(row=r, column=5).value
        assert e.cell(row=r, column=6).value == preco_1[cod], r
        assert e.cell(row=r, column=7).value == round(qt * preco_1[cod], 2), r

    # ---- relatório ----
    print(f"gravado {SAIDA}")
    print(f"partes alteradas: {sorted(mexidas)} · células alteradas: {len(dif)}")
    for pma, cod, desc, q, unit, tot in COMPRA2:
        print(f"\n{pma} {PECA[pma]} ({q} un.)")
        for k, a in enumerate(lista[pma], 1):
            x_ = g.get(a, {})
            print(f"   {k:2d} {a:<11} {x_.get('crit', ''):<13} {x_.get('indice', '') or '':>4} {x_.get('defeito', '')}")
    print("\nGestão, coluna M:")
    for a, v in sorted(pma_novo.items(), key=lambda kv: g[kv[0]]["indice"]):
        print(f"   M{g[a]['linha']:<3} {a} {g[a]['crit']:<13} {g[a]['pma']:<10} -> {v}")
    print("\nRecebem a peça que sobrou (não tinham PMA) e ainda esperam:")
    for a, ps in parcial.items():
        print(f"   {a} {g[a]['crit']:<13} Índice {g[a]['indice']:>3} {g[a]['defeito']:<22} falta {', '.join(ps)}")
    print("\nPreço usado (1ª compra) x requisição:")
    for pma, cod, _, q, unit, tot in COMPRA2:
        print(f"   {pma} {cod}: {preco_1[cod]:>10.2f} x {q} = {round(q * preco_1[cod], 2):>11.2f}   (requisição {unit:.2f} = {tot:.2f})")
    print("\nContinuam sem peça (faltou):")
    for a, ps in sem.items():
        print(f"   {a} {g[a]['crit']:<13} Índice {g[a]['indice']:>3} {g[a]['defeito']:<22} faltou {', '.join(ps)}")
    print(f"\nSaldo (Reserva) nas duas compras: {saldo}")


if __name__ == "__main__":
    main()
